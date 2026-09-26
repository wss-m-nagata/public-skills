# -*- coding: utf-8 -*-
"""DFD（データの流れ）とデータ量の見積もりを、1冊のExcelとして生成する。

内容はデータ定義ファイル（YAML）が正であり、Excelを直接編集しない。
Excelが直接編集された場合は、diff で差分を出してYAMLへ取り込む。

使い方:
    python build_data_flow_volume.py build <データ定義.yaml>
        YAMLと同じフォルダに、DFDの図（PNG）とExcelを出力する。

    python build_data_flow_volume.py diff <比較元.xlsx> <比較先.xlsx>
        2つのExcelの値の違いを「シート名!セル番地」で一覧にする。
        修正箇所の報告や、直接編集されたExcelの取り込みに使う。

    python build_data_flow_volume.py preview <Excel> <出力フォルダ>
        全シートをPDFに書き出し、はみ出しやページ分かれを目で確かめられるようにする。
        Microsoft Excel（Windows、pywin32が必要）か LibreOffice（soffice）を使う。
"""

import argparse
import math
import platform
import shutil
import subprocess
import tempfile
import sys
from pathlib import Path

import openpyxl
import yaml
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image as PILImage
from PIL import ImageDraw, ImageFont

# 描画の倍率（高解像度で描き、Excel上で縮小して表示する）
SCALE = 2
# 図に使う日本語フォントの候補（通常, 太字）。上から順に、見つかったものを使う。
# YAMLの font: {regular: ..., bold: ...} で指定した場合はそちらを優先する
FONT_CANDIDATES = [
    ("C:/Windows/Fonts/meiryo.ttc", "C:/Windows/Fonts/meiryob.ttc"),  # Windows
    ("C:/Windows/Fonts/YuGothM.ttc", "C:/Windows/Fonts/YuGothB.ttc"),  # Windows
    (
        "/System/Library/Fonts/ヒラギノ角ゴシック W3.ttc",
        "/System/Library/Fonts/ヒラギノ角ゴシック W6.ttc",
    ),  # macOS
    (
        "/System/Library/Fonts/Hiragino Sans GB.ttc",
        "/System/Library/Fonts/Hiragino Sans GB.ttc",
    ),  # macOS
    (
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc",
    ),  # Linux
    (
        "/usr/share/fonts/noto-cjk/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/noto-cjk/NotoSansCJK-Bold.ttc",
    ),  # Linux
    (
        "/usr/share/fonts/google-noto-cjk/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/google-noto-cjk/NotoSansCJK-Bold.ttc",
    ),  # Linux
]
# 実際に使うフォント（build の最初に _resolve_fonts で決める）
FONTS = {"regular": None, "bold": None}
# Excelのセルのフォント（YAMLの excel_font で変えられる。無いフォント名はExcelが代わりのフォントで表示する）
XL_FONT = "Meiryo UI"

# 図の配色（種類 → (塗り, 線, 文字)）
COLORS = {
    "ext": ("#FFF4E0", "#E08A00", "#5A3A00"),  # 外部（画面・他のシステム）
    "proc": ("#E6F0FF", "#2F6FD6", "#0B2E6B"),  # 処理
    "store": ("#EAF7EA", "#2E8B57", "#11401F"),  # テーブル
    "big": ("#BFE6BF", "#1E6B3E", "#0B2E15"),  # 件数の多いテーブル
    "file": ("#F3ECFF", "#6A5ACD", "#2A1F66"),  # ファイル
}
# 線の種類（bold＝件数の多い流れ, solid＝通常, dash＝読むだけ, both＝双方向）
FLOW_STYLES = ("bold", "solid", "dash", "both")
# 凡例の説明（YAMLの legend_labels で種類ごとに上書きできる）
LEGEND_LABELS = {
    "ext": "画面・他のシステム",
    "proc": "処理",
    "store": "テーブル",
    "big": "テーブル（件数の多いもの）",
    "file": "ファイル",
    "bold": "件数の多い流れ",
    "solid": "データの流れ",
    "dash": "読むだけ",
    "both": "読み書きの両方",
}
# 凡例1行の高さ（倍率を掛ける前の値）
LEGEND_ROW_H = 56

# 表の見出し色と罫線
HEADER_FILL = "2F6FD6"
# 合計行の塗り色（凡例にも同じ色を出す）
TOTAL_FILL = "DDDDDD"
BORDER_COLOR = "B0B7C3"
# 表の行の高さを文字量から決めるときの、1行あたりの高さと上下の余白（ポイント）
LINE_HEIGHT = 14
ROW_PADDING = 10
MIN_ROW_HEIGHT = 30


# ---------------------------------------------------------------------------
# 図（DFD）
# ---------------------------------------------------------------------------
def _resolve_fonts(spec_font):
    """図に使う日本語フォントを決める。

    Args:
        spec_font: YAMLの font 指定（{regular, bold}）。無ければNone。

    Raises:
        ValueError: 使えるフォントが見つからない場合（英字フォントで描くと日本語が豆腐になるため止める）。
    """
    if spec_font:
        regular = spec_font.get("regular")
        bold = spec_font.get("bold") or regular
        if not regular or not Path(regular).exists() or not Path(bold).exists():
            raise ValueError(
                f"YAMLの font に指定したフォントが見つかりません: {spec_font}"
            )
        FONTS.update(regular=regular, bold=bold)
        return
    for regular, bold in FONT_CANDIDATES:
        if Path(regular).exists():
            # 太字が無い環境では通常のフォントで代用する（見た目の差だけで、読めなくはならない）
            FONTS.update(regular=regular, bold=bold if Path(bold).exists() else regular)
            return
    raise ValueError(
        "日本語フォントが見つかりません。YAMLに font: {regular: <パス>, bold: <パス>} を指定してください"
        f"（{platform.system()}）"
    )


def _font(size, bold=False):
    """描画用のフォントを返す。

    Args:
        size: 文字の大きさ（倍率を掛ける前の値）。
        bold: 太字にするかどうか。

    Returns:
        ImageFont.FreeTypeFont: フォント。
    """
    return ImageFont.truetype(FONTS["bold" if bold else "regular"], size * SCALE)


def _s(v):
    """座標・長さへ描画の倍率を掛ける。"""
    return int(v * SCALE)


def _draw_text(draw, cx, cy, text, font, color, bg=None):
    """複数行の文字を中心揃えで描く。bgを指定すると文字の後ろを塗る（線の上に文言を置くとき）。"""
    box = draw.multiline_textbbox(
        (0, 0), text, font=font, align="center", spacing=_s(4)
    )
    w, h = box[2] - box[0], box[3] - box[1]
    x, y = _s(cx) - w / 2, _s(cy) - h / 2
    if bg:
        pad = _s(4)
        draw.rounded_rectangle(
            (x - pad, y - pad, x + w + pad, y + h + pad), radius=_s(4), fill=bg
        )
    draw.multiline_text(
        (x - box[0], y - box[1]),
        text,
        font=font,
        fill=color,
        align="center",
        spacing=_s(4),
    )


def _draw_node(draw, node):
    """図の要素を1つ描く（処理＝円、テーブル＝円柱、ファイル＝平行四辺形、外部＝四角）。"""
    kind, cx, cy, w, h, text = (
        node["kind"],
        node["x"],
        node["y"],
        node["w"],
        node["h"],
        node["text"],
    )
    fill, line, color = COLORS[kind]
    x0, y0, x1, y1 = _s(cx - w / 2), _s(cy - h / 2), _s(cx + w / 2), _s(cy + h / 2)
    width = _s(3) if kind == "big" else _s(2)
    if kind == "proc":
        draw.ellipse((x0, y0, x1, y1), fill=fill, outline=line, width=width)
    elif kind in ("store", "big"):
        # 円柱。側面を塗ってから上下の楕円を重ねる
        e = _s(14)
        draw.rectangle((x0, y0 + e, x1, y1 - e), fill=fill)
        draw.ellipse((x0, y1 - 2 * e, x1, y1), fill=fill, outline=line, width=width)
        draw.rectangle((x0 + width, y0 + e, x1 - width, y1 - e), fill=fill)
        draw.line((x0, y0 + e, x0, y1 - e), fill=line, width=width)
        draw.line((x1, y0 + e, x1, y1 - e), fill=line, width=width)
        draw.ellipse((x0, y0, x1, y0 + 2 * e), fill=fill, outline=line, width=width)
        cy += 6
    elif kind == "file":
        k = _s(20)
        draw.polygon(
            [(x0 + k, y0), (x1, y0), (x1 - k, y1), (x0, y1)],
            fill=fill,
            outline=line,
            width=width,
        )
    else:
        draw.rounded_rectangle(
            (x0, y0, x1, y1), radius=_s(10), fill=fill, outline=line, width=width
        )
    _draw_text(draw, cx, cy, text, _font(15, bold=kind == "big"), color)


def _edge_point(node, tx, ty):
    """要素の中心から(tx, ty)へ向かう線が、要素の外周と交わる点を返す。"""
    cx, cy, w, h = node["x"], node["y"], node["w"], node["h"]
    dx, dy = tx - cx, ty - cy
    if node["kind"] == "proc":
        # 円は半径で切る
        r = w / 2
        d = math.hypot(dx, dy)
        return cx + dx * r / d, cy + dy * r / d
    # 四角形は縦横の近いほうの辺で切る
    sx = (w / 2) / abs(dx) if dx else float("inf")
    sy = (h / 2) / abs(dy) if dy else float("inf")
    t = min(sx, sy)
    return cx + dx * t, cy + dy * t


def _draw_arrow_head(draw, x0, y0, x1, y1, color):
    """(x0,y0)→(x1,y1)の向きで、(x1,y1)に矢じりを描く（座標は倍率を掛けた後の値）。"""
    ang = math.atan2(y1 - y0, x1 - x0)
    size = _s(12)
    pts = [
        (x1, y1),
        (x1 - size * math.cos(ang - 0.4), y1 - size * math.sin(ang - 0.4)),
        (x1 - size * math.cos(ang + 0.4), y1 - size * math.sin(ang + 0.4)),
    ]
    draw.polygon(pts, fill=color)


def _draw_flow(draw, nodes, flow):
    """データの流れ（矢印）を1本描き、文言を置く位置を返す。

    offsetを指定すると進行方向の右手側へ平行にずらす。逆向きの線は反対側へずれるため、
    同じ2要素の間を行きと帰りの2本で結んでも重ならない。
    """
    a, b = nodes[flow["from"]], nodes[flow["to"]]
    style = flow.get("style", "solid")
    sx, sy = _edge_point(a, b["x"], b["y"])
    ex, ey = _edge_point(b, a["x"], a["y"])
    offset = flow.get("offset", 0)
    if offset:
        length = math.hypot(ex - sx, ey - sy)
        nx, ny = -(ey - sy) / length * offset, (ex - sx) / length * offset
        sx, sy, ex, ey = sx + nx, sy + ny, ex + nx, ey + ny

    color = "#333333" if style == "bold" else "#666666"
    width = _s(4) if style == "bold" else _s(2)
    p0, p1 = (_s(sx), _s(sy)), (_s(ex), _s(ey))
    if style == "dash":
        # 破線。一定間隔で線分を描く
        n = int(math.hypot(p1[0] - p0[0], p1[1] - p0[1]) // _s(12))
        for i in range(0, n, 2):
            q0 = (p0[0] + (p1[0] - p0[0]) * i / n, p0[1] + (p1[1] - p0[1]) * i / n)
            q1 = (
                p0[0] + (p1[0] - p0[0]) * (i + 1) / n,
                p0[1] + (p1[1] - p0[1]) * (i + 1) / n,
            )
            draw.line((q0, q1), fill=color, width=width)
    else:
        draw.line((p0, p1), fill=color, width=width)
    _draw_arrow_head(draw, *p0, *p1, color)
    if style == "both":
        _draw_arrow_head(draw, *p1, *p0, color)
    pos = flow.get("pos", 0.5)
    return sx + (ex - sx) * pos, sy + (ey - sy) * pos


def _legend_items(dfd):
    """凡例に載せる項目を、図で実際に使っている要素の種類・線の種類から決める。

    使っていない種類まで載せると、読み手がその要素を図の中で探してしまうため、使っているものだけにする。

    Returns:
        list: (種類, 説明) のリスト。種類は COLORS のキーか FLOW_STYLES の値。
    """
    labels = {**LEGEND_LABELS, **dfd.get("legend_labels", {})}
    used_kinds = {node["kind"] for node in dfd["nodes"].values()}
    used_styles = {flow.get("style", "solid") for flow in dfd["flows"]}
    order = list(COLORS) + list(FLOW_STYLES)
    return [
        (key, labels[key]) for key in order if key in used_kinds or key in used_styles
    ]


def _draw_legend(draw, width, items):
    """凡例（色・形の見本と説明）を、図の上端に横一列で描く。収まらなければ折り返す。

    Args:
        draw: 描画先。
        width: 図の幅（倍率を掛ける前の値）。
        items: _legend_items の結果。

    Returns:
        int: 凡例が使った高さ（倍率を掛ける前の値）。
    """
    font = _font(16)
    x, y, row_h = 30, 20, LEGEND_ROW_H
    _draw_text(draw, x + 20, y + row_h / 2, "凡例", _font(16, bold=True), "#333333")
    x += 70
    for key, text in items:
        text_w = draw.textlength(text, font=font) / SCALE
        item_w = 70 + 10 + text_w + 34
        if x + item_w > width - 20:
            x, y = 100, y + row_h
        cy = y + row_h / 2
        if key in COLORS:
            # 図の要素と同じ描き方で、小さな見本を描く
            w, h = (38, 38) if key == "proc" else (70, 36)
            _draw_node(
                draw, {"kind": key, "x": x + 35, "y": cy, "w": w, "h": h, "text": ""}
            )
        else:
            # 線の見本は、見本用の2点を結ぶ矢印として描く
            pts = {
                "a": {"kind": "ext", "x": x, "y": cy, "w": 0.01, "h": 0.01},
                "b": {"kind": "ext", "x": x + 70, "y": cy, "w": 0.01, "h": 0.01},
            }
            _draw_flow(draw, pts, {"from": "a", "to": "b", "style": key})
        draw.text((_s(x + 80), _s(cy)), text, font=font, fill="#333333", anchor="lm")
        x += item_w
    return y + row_h + 10


def build_png(dfd, out_png):
    """DFDをPNGとして描く。上端に凡例を置き、その下に図を描く。

    Args:
        dfd: YAMLのdfdシート定義（canvas・bands・nodes・flows、任意で legend_labels）。
        out_png: 出力パス。

    Returns:
        tuple: 画像の大きさ（幅, 高さ）。
    """
    if "legend" in dfd:
        raise ValueError(
            "dfdシートの legend は廃止しました。凡例は図から自動で作るため、legend の行を削除してください"
            "（説明を変えたいときは legend_labels を使う）"
        )
    canvas_w, canvas_h = dfd["canvas"]["width"], dfd["canvas"]["height"]
    nodes = dfd["nodes"]
    for flow in dfd["flows"]:
        if flow["from"] not in nodes or flow["to"] not in nodes:
            raise ValueError(f"flowsが存在しない要素を指しています: {flow}")
        if flow.get("style", "solid") not in FLOW_STYLES:
            raise ValueError(f"flowsのstyleは {FLOW_STYLES} のいずれかです: {flow}")
        # 「1,000件」を引用符で囲み忘れると、{ } の中で「,」が区切りと見なされ数値や別キーになる
        if "label" in flow and not isinstance(flow["label"], str):
            raise ValueError(
                f"flowsのlabelが文字列ではありません（「,」を含む値は引用符で囲む）: {flow}"
            )
    for key, node in nodes.items():
        if node.get("kind") not in COLORS or not isinstance(node.get("text"), str):
            raise ValueError(
                f"nodesの{key}のkindかtextが不正です（textは文字列、kindは {tuple(COLORS)}）: {node}"
            )

    # 凡例の高さを先に測る（折り返しの行数で変わるため、仮の画像に描いて求める）
    items = _legend_items(dfd)
    probe = ImageDraw.Draw(PILImage.new("RGB", (_s(canvas_w), _s(200)), "#FFFFFF"))
    legend_h = _draw_legend(probe, canvas_w, items)

    img = PILImage.new("RGB", (_s(canvas_w), _s(canvas_h + legend_h)), "#FFFFFF")
    _draw_legend(ImageDraw.Draw(img), canvas_w, items)

    # 図は凡例の下に描く（座標はYAMLのまま使えるよう、別の画像に描いてから貼る）
    diagram = PILImage.new("RGB", (_s(canvas_w), _s(canvas_h)), "#FFFFFF")
    draw = ImageDraw.Draw(diagram)

    # 列の帯と見出し
    for band in dfd["bands"]:
        cx, w = band["x"], band["width"]
        draw.rounded_rectangle(
            (_s(cx - w / 2), _s(20), _s(cx + w / 2), _s(canvas_h - 20)),
            radius=_s(12),
            fill="#F7F8FA",
            outline="#D5D9E0",
            width=_s(1),
        )
        _draw_text(draw, cx, 50, band["title"], _font(16, bold=True), "#333333")

    # 線 → 要素 → 線の文言の順に描く（文言が線や要素に隠れないようにする）
    labels = []
    for flow in dfd["flows"]:
        lx, ly = _draw_flow(draw, nodes, flow)
        if flow.get("label"):
            labels.append((lx, ly, flow["label"]))
    for node in nodes.values():
        _draw_node(draw, node)
    for lx, ly, label in labels:
        _draw_text(draw, lx, ly, label, _font(12), "#222222", bg="#FFFFFF")

    img.paste(diagram, (0, _s(legend_h)))
    img.save(out_png)
    return img.size


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def _display_width(text):
    """全角を2、半角を1として文字列の表示幅を返す。"""
    return sum(2 if ord(ch) > 0x2E7F else 1 for ch in text)


def _needed_lines(value, col_width):
    """セルの値を列幅で折り返したときの行数を見積もる。"""
    if value is None:
        return 1
    per_line = max(col_width * 0.95, 1)
    return sum(
        max(1, math.ceil(_display_width(part) / per_line))
        for part in str(value).split("\n")
    )


def _title(ws, text, sub=None):
    """シートの表題（と補足の1行）を書く。"""
    ws.column_dimensions["A"].width = 2
    ws["B2"] = text
    ws["B2"].font = Font(name=XL_FONT, bold=True, size=16, color="0B2E6B")
    if sub:
        ws["B3"] = sub
        ws["B3"].font = Font(name=XL_FONT, size=10, color="555555")
    ws.sheet_view.showGridLines = False


def _write_table(ws, sheet):
    """表のシートを書く（見出し・行・合計行）。行の高さは文字量から決める。

    Args:
        ws: 書き込み先シート。
        sheet: YAMLのtableシート定義（header・widths・rows・category_column・category_fill・total）。
    """
    header, widths, rows = sheet["header"], sheet["widths"], sheet["rows"]
    if len(header) != len(widths):
        raise ValueError(f"シート「{sheet['name']}」のheaderとwidthsの数が一致しません")
    # 「1,000件」を引用符で囲み忘れると [ ] の中で「,」が区切りと見なされ、列が1つ増える。合計行も同じく確かめる
    checks = [(f"{i}行目", row) for i, row in enumerate(rows, start=1)]
    if sheet.get("total"):
        checks.append(("合計行", sheet["total"]))
    for where, row in checks:
        if len(row) != len(header):
            raise ValueError(
                f"シート「{sheet['name']}」の{where}の列数がheaderと一致しません"
                f"（「,」を含む値は引用符で囲む）: {row}"
            )

    thin = Side(style="thin", color=BORDER_COLOR)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    top = 5

    # 見出し行
    for c, (text, width) in enumerate(zip(header, widths), start=2):
        ws.column_dimensions[ws.cell(top, c).column_letter].width = width
        cell = ws.cell(top, c, text)
        cell.font = Font(name=XL_FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        cell.border = border
    ws.row_dimensions[top].height = 30

    # 本文（分類の列があれば、分類ごとに行を塗る）と合計行
    category_column = sheet.get("category_column")
    category_fill = sheet.get("category_fill", {})

    # 行を色で塗り分けるときは、表のすぐ上（4行目）に色の見本つきの凡例を置く。
    # 色の意味が書かれていないと、読み手は推測するしかないため。見出しの青のような飾りの色は載せない
    swatches = list(category_fill.items())
    if sheet.get("total"):
        swatches.append(("合計", TOTAL_FILL))
    if swatches:
        head = ws.cell(top - 1, 2, "凡例")
        head.font = Font(name=XL_FONT, bold=True, size=10, color="333333")
        head.alignment = Alignment(horizontal="center", vertical="center")
        for c, (name, color) in enumerate(swatches, start=3):
            cell = ws.cell(top - 1, c, name)
            cell.font = Font(name=XL_FONT, size=10)
            cell.fill = PatternFill("solid", fgColor=color)
            cell.alignment = Alignment(
                horizontal="center", vertical="center", shrink_to_fit=True
            )
            cell.border = border
        ws.row_dimensions[top - 1].height = 22
    body = [(row, False) for row in rows]
    if sheet.get("total"):
        body.append((sheet["total"], True))
    for r, (row, is_total) in enumerate(body, start=top + 1):
        fill = TOTAL_FILL if is_total else None
        if not is_total and category_column is not None:
            fill = category_fill.get(row[category_column])
        for c, value in enumerate(row, start=2):
            cell = ws.cell(r, c, value)
            cell.font = Font(name=XL_FONT, size=10, bold=is_total)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = border
            if fill:
                cell.fill = PatternFill("solid", fgColor=fill)
        lines = max(_needed_lines(v, w) for v, w in zip(row, widths))
        ws.row_dimensions[r].height = max(
            MIN_ROW_HEIGHT, lines * LINE_HEIGHT + ROW_PADDING
        )
    ws.freeze_panes = ws.cell(top + 1, 2)


def _write_dfd_sheet(ws, sheet, png_path, png_size):
    """DFDのシートを書く（要点の欄・凡例・図）。"""
    points = sheet.get("points", [])
    for c in range(2, 2 + len(points) * 4 + 1):
        ws.column_dimensions[ws.cell(1, c).column_letter].width = 9
    for point in points:
        if not isinstance(point.get("label"), str) or not isinstance(
            point.get("value"), str
        ):
            raise ValueError(
                f"pointsのlabel・valueは文字列にしてください（「,」を含む値は引用符で囲む）: {point}"
            )
    for i, point in enumerate(points):
        col = 2 + i * 4
        head = ws.cell(5, col, point["label"])
        head.font = Font(name=XL_FONT, bold=True, size=10, color="FFFFFF")
        head.fill = PatternFill("solid", fgColor="1E6B3E")
        head.alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells(start_row=5, start_column=col, end_row=5, end_column=col + 2)
        body = ws.cell(6, col, point["value"])
        body.font = Font(name=XL_FONT, bold=True, size=11, color="0B2E15")
        body.fill = PatternFill("solid", fgColor="E2F2E2")
        body.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        ws.merge_cells(start_row=6, start_column=col, end_row=6, end_column=col + 2)
    if points:
        lines = max(_needed_lines(p["value"], 27) for p in points)
        ws.row_dimensions[6].height = max(MIN_ROW_HEIGHT, lines * 16 + ROW_PADDING)

    # 凡例は図の画像の上端に描いてあるので、セルには書かない
    img = XLImage(str(png_path))
    img.width = sheet.get("image_width", 1300)
    img.height = int(img.width * png_size[1] / png_size[0])
    ws.add_image(img, "B8")


def build(yaml_path):
    """データ定義（YAML）から、DFDの図（PNG）とExcelを生成する。

    Args:
        yaml_path: データ定義ファイルのパス。

    Returns:
        tuple: (Excelのパス, PNGのパス)。
    """
    yaml_path = Path(yaml_path)
    spec = yaml.safe_load(yaml_path.read_text(encoding="utf-8"))
    out_xlsx = yaml_path.parent / spec["output"]["xlsx"]
    out_png = yaml_path.parent / spec["output"]["png"]

    # フォントを決める（図は日本語フォントが必須、Excelのセルは指定があれば差し替える）
    global XL_FONT
    _resolve_fonts(spec.get("font"))
    XL_FONT = spec.get("excel_font", XL_FONT)

    sheets = spec["sheets"]
    dfd_sheets = [s for s in sheets if s["type"] == "dfd"]
    if len(dfd_sheets) != 1:
        raise ValueError("type: dfd のシートはちょうど1枚にしてください")
    png_size = build_png(dfd_sheets[0], out_png)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet in sheets:
        ws = wb.create_sheet(sheet["name"])
        _title(ws, sheet["title"], sheet.get("subtitle"))
        if sheet["type"] == "dfd":
            _write_dfd_sheet(ws, sheet, out_png, png_size)
        elif sheet["type"] == "table":
            _write_table(ws, sheet)
        else:
            raise ValueError(f"シートのtypeは dfd か table です: {sheet['name']}")

    # 印刷はA3横・1シート1ページに収める
    for ws in wb.worksheets:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    wb.save(out_xlsx)
    return out_xlsx, out_png


def diff(before_path, after_path):
    """2つのExcelの値の違いを「シート名!セル番地」の一覧で返す。

    Args:
        before_path: 比較元のExcel。
        after_path: 比較先のExcel。

    Returns:
        list: (場所, 比較元の値, 比較先の値) のリスト。
    """
    before = openpyxl.load_workbook(before_path)
    after = openpyxl.load_workbook(after_path)
    result = []
    for name in dict.fromkeys(before.sheetnames + after.sheetnames):
        if name not in before.sheetnames or name not in after.sheetnames:
            result.append(
                (
                    name,
                    "シートあり" if name in before.sheetnames else "シートなし",
                    "シートあり" if name in after.sheetnames else "シートなし",
                )
            )
            continue
        a, b = before[name], after[name]
        for r in range(1, max(a.max_row, b.max_row) + 1):
            for c in range(1, max(a.max_column, b.max_column) + 1):
                va, vb = a.cell(r, c).value, b.cell(r, c).value
                if va != vb:
                    result.append((f"{name}!{a.cell(r, c).coordinate}", va, vb))
    return result


def preview(xlsx_path, out_dir):
    """Excelの全シートをPDFに書き出す（印刷時のはみ出し・ページ分かれを目で確かめるため）。

    Microsoft Excel（Windows、pywin32）があればそれを使い、無ければ LibreOffice（soffice）を使う。
    Excelは実際の印刷設定どおりに出力できるため、使えるなら優先する。

    Args:
        xlsx_path: 対象のExcel。
        out_dir: PDFの出力フォルダ。

    Returns:
        list: 出力したPDFのパス。

    Raises:
        RuntimeError: ExcelもLibreOfficeも使えない場合。
    """
    xlsx_path, out_dir = Path(xlsx_path).resolve(), Path(out_dir).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)

    # Microsoft Excel（シートごとに1つのPDF）
    try:
        import win32com.client  # noqa: PLC0415 （Windows以外では入っていないため、ここで読み込む）
    except ImportError:
        win32com = None
    if win32com is not None:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        outputs = []
        # Excelはパスが長い（約218文字超）ファイルを開けないため、短い一時フォルダで開いて書き出してから移す
        with tempfile.TemporaryDirectory(prefix="dfv") as tmp:
            src = Path(tmp) / "book.xlsx"
            shutil.copyfile(xlsx_path, src)
            try:
                wb = excel.Workbooks.Open(str(src), ReadOnly=True)
                try:
                    for i, ws in enumerate(wb.Worksheets, start=1):
                        tmp_pdf = Path(tmp) / f"{i}.pdf"
                        ws.ExportAsFixedFormat(0, str(tmp_pdf))
                        pdf = out_dir / f"{xlsx_path.stem}_{ws.Name}.pdf"
                        shutil.move(str(tmp_pdf), pdf)
                        outputs.append(pdf)
                finally:
                    wb.Close(SaveChanges=False)
            finally:
                excel.Quit()
        return outputs

    # LibreOffice（ブック全体を1つのPDF）
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        subprocess.run(
            [
                soffice,
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(out_dir),
                str(xlsx_path),
            ],
            check=True,
            capture_output=True,
        )
        return [out_dir / f"{xlsx_path.stem}.pdf"]
    raise RuntimeError(
        "Microsoft Excel（pywin32）も LibreOffice（soffice）も見つからないため、PDFに書き出せません"
    )


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    sub = parser.add_subparsers(dest="command", required=True)
    p_build = sub.add_parser("build")
    p_build.add_argument("yaml_path")
    p_diff = sub.add_parser("diff")
    p_diff.add_argument("before")
    p_diff.add_argument("after")
    p_preview = sub.add_parser("preview")
    p_preview.add_argument("xlsx_path")
    p_preview.add_argument("out_dir")
    args = parser.parse_args()

    if args.command == "build":
        out_xlsx, out_png = build(args.yaml_path)
        print(f"出力: {out_xlsx}")
        print(f"出力: {out_png}")
        return

    if args.command == "preview":
        for pdf in preview(args.xlsx_path, args.out_dir):
            print(f"出力: {pdf}")
        return

    rows = diff(args.before, args.after)
    if not rows:
        print("値の違いはありません")
        return
    for where, va, vb in rows:
        print(f"{where}: {va!r} → {vb!r}")


if __name__ == "__main__":
    main()
