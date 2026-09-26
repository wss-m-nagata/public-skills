# -*- coding: utf-8 -*-
"""
案件（機能）ごとに分かれたNablarch様式のテーブル定義書を、1冊のテーブル定義書と
テーブル一覧へ統合するスクリプト。

中間データ（YAML。1テーブル=1ファイル）を挟んで、次の3段階で処理する。

1. extract : 案件ごとのテーブル定義書(xlsx)からテーブルシートを読み出し、
             優先度に従ってテーブルごとに1つの定義を選び、中間データへ書き出す。
2. build   : 中間データから、統合版のテーブル定義書とテーブル一覧を生成する。
             生成後に統合版を読み戻し、中間データと一致することを検証する。
3. check   : 既存の統合版テーブル定義書が中間データと一致するかだけを検証する。

優先度は、案件のベース関係（どの案件の定義を土台に変更したか）を数値で表したもの。
数値が大きいほど新しい定義として採用する。
同じテーブルが同じ優先度の複数の冊にあり、内容が異なる場合は、どちらを採るか
機械的に決められないためエラーにする（内容が同一なら重複として1つにまとめる）。

使い方:
    python merge_table_books.py extract <中間データ出力ディレクトリ>
        --book <優先度> <テーブル定義書xlsx> [--book ...]
        --author <担当者> --date <YYYY-MM-DD> [--force]

    python merge_table_books.py build <中間データディレクトリ>
        --base-def <書式の土台にするテーブル定義書xlsx>
        --base-list <書式の土台にするテーブル一覧xlsx>
        --out-def <出力テーブル定義書xlsx> --out-list <出力テーブル一覧xlsx>

    python merge_table_books.py check <中間データディレクトリ> <テーブル定義書xlsx>

土台にする冊は、データ型プルダウン（データシート）をPostgreSQL型へ置き換え済みの
既存の冊を指定する（テンプレート原本はOracle型のため）。
土台の表紙・変更履歴・目次・データの各シートの書式を引き継ぎ、テーブルシートは
土台の先頭テーブルシートを複製して作り直す。
"""

import sys
import argparse
import datetime
from pathlib import Path

import openpyxl
import yaml

sys.path.insert(0, str(Path(__file__).parent))
from xlsx_common import copy_data_validations, save_with_shapes, ensure_utf8_stdio

# テーブルシートの列の意味（キー → 列）。行11〜26が1行1カラムの記入欄である
COLUMN_CELLS = {
    "no": "A",  # No.
    "logical": "B",  # 論理名称
    "physical": "G",  # 物理名称
    "domain": "L",  # ドメイン名
    "type": "Q",  # データ型
    "length": "T",  # 桁数
    "pk": "V",  # PK（主キー内の順序）
    "required": "W",  # 必須（○／×）
    "definition": "AE",  # 項目定義
    "default": "AO",  # 初期値
    "encrypted": "AQ",  # 暗号化対象（○／×）
    "remarks": "AS",  # 備考
}
# INDEX欄（インデックス番号1〜7）の列
INDEX_COLUMNS = ["X", "Y", "Z", "AA", "AB", "AC", "AD"]
# カラム記入欄の開始行・終了行
COLUMN_FIRST_ROW = 11
COLUMN_LAST_ROW = 26
# INDEX欄の見出し行（U＝ユニークインデックスの印を書く行）
INDEX_FLAG_ROW = 10

# テーブル一覧シート（シート名「1」）の列と記入欄の範囲
LIST_SHEET = "1"
LIST_FIRST_ROW = 10
LIST_LAST_ROW = 38

# 変更履歴シートの列と記入欄の範囲
HISTORY_CELLS = {
    "version": "B",  # 版数
    "date": "D",  # 変更日
    "kind": "G",  # 区分（新規／変更）
    "place": "J",  # 変更箇所（項番等）
    "content": "Q",  # 変更内容
    "author": "AF",  # 担当者
}
HISTORY_FIRST_ROW = 8
HISTORY_LAST_ROW = 33

# 目次シートの記入欄の範囲
TOC_FIRST_ROW = 7
TOC_LAST_ROW = 49

# 中間データのファイル名
BOOK_META_FILE = "_book.yaml"
TABLES_DIR = "tables"


def _table_sheets(wb):
    """テーブルシート（目次とデータの間にあるシート）を順に返す。

    Args:
        wb: 対象のワークブック。

    Returns:
        list: テーブルシートのリスト。
    """
    names = wb.sheetnames
    if "目次" not in names or "データ" not in names:
        raise ValueError(
            f"テーブル定義書の構成ではありません（目次・データシートが無い）: {names}"
        )
    first = names.index("目次") + 1
    last = names.index("データ")
    return [wb[name] for name in names[first:last]]


def _read_table_sheet(ws):
    """テーブルシート1枚を中間データ（dict）へ読み出す。

    Args:
        ws: テーブルシート。

    Returns:
        dict: シート名・論理/物理テーブル名・説明・INDEX見出し・カラム一覧。
    """
    # テーブルの見出し部分
    table = {
        "sheet": ws.title,
        "logical_name": ws["F5"].value,
        "physical_name": ws["W5"].value,
        "description": ws["F6"].value,
        "index_flags": [ws[f"{col}{INDEX_FLAG_ROW}"].value for col in INDEX_COLUMNS],
        "columns": [],
    }
    if not table["physical_name"]:
        raise ValueError(f"物理テーブル名（W5）が空です: シート「{ws.title}」")

    # カラム記入欄。論理名称〜備考のいずれかに値がある行をカラムとみなす
    for row in range(COLUMN_FIRST_ROW, COLUMN_LAST_ROW + 1):
        column = {key: ws[f"{col}{row}"].value for key, col in COLUMN_CELLS.items()}
        column["index"] = [ws[f"{col}{row}"].value for col in INDEX_COLUMNS]
        values = [v for k, v in column.items() if k not in ("no", "index")]
        if any(v is not None for v in values) or any(
            v is not None for v in column["index"]
        ):
            table["columns"].append(column)
    return table


def _subsystem_from_filename(path):
    """ファイル名「テーブル定義書_<ID>_<名称>.xlsx」からサブシステムID・名称を取り出す。

    Args:
        path: テーブル定義書のパス。

    Returns:
        tuple: (サブシステムID, サブシステム名)。
    """
    parts = Path(path).stem.split("_", 2)
    if len(parts) != 3 or parts[0] != "テーブル定義書":
        raise ValueError(
            f"ファイル名が「テーブル定義書_<ID>_<名称>.xlsx」ではありません: {path}"
        )
    return parts[1], parts[2]


def extract(out_dir, books, author, date, force):
    """案件ごとのテーブル定義書を読み、テーブルごとに採用する定義を中間データへ書き出す。

    Args:
        out_dir: 中間データの出力ディレクトリ。
        books: (優先度, xlsxパス) のリスト。
        author: 変更履歴に書く担当者。
        date: 変更履歴に書く日付（datetime.date）。
        force: 既存の中間データを上書きするかどうか。

    Returns:
        dict: 採用結果（テーブル → 採用元・不採用元）。
    """
    out_dir = Path(out_dir)
    meta_path = out_dir / BOOK_META_FILE
    if meta_path.exists() and not force:
        raise ValueError(
            f"中間データが既にあります（上書きする場合は --force）: {out_dir}"
        )
    # 取り込み直しでも変更履歴は引き継ぐ（版の追加は _book.yaml の history へ手で書く）
    history = None
    if meta_path.exists():
        history = yaml.safe_load(meta_path.read_text(encoding="utf-8"))["history"]

    # 全冊が同じサブシステムであることを確認する
    subsystems = {_subsystem_from_filename(path) for _, path in books}
    if len(subsystems) != 1:
        raise ValueError(
            f"サブシステムの異なる冊が混在しています: {sorted(subsystems)}"
        )
    subsystem_id, subsystem_name = subsystems.pop()

    # 優先度の低い順に読み、テーブルごとに最も優先度の高い定義を残す
    adopted = {}  # 物理テーブル名 → (優先度, 出典, テーブル定義)
    rejected = {}  # 物理テーブル名 → 不採用になった出典のリスト
    order = []  # テーブルの並び順（最初に現れた順）
    for priority, path in sorted(books, key=lambda b: b[0]):
        wb = openpyxl.load_workbook(path)
        for ws in _table_sheets(wb):
            table = _read_table_sheet(ws)
            name = table["physical_name"]
            source = Path(path).as_posix()
            if name not in adopted:
                order.append(name)
            elif adopted[name][0] == priority:
                # 同じ優先度での重複は、内容が同一の場合に限り許す
                if adopted[name][2] != table:
                    raise ValueError(
                        f"同じ優先度({priority})の冊に、内容の異なる {name} があります: "
                        f"{adopted[name][1]} / {source}"
                    )
                continue
            else:
                rejected.setdefault(name, []).append(adopted[name][1])
            adopted[name] = (priority, source, table)

    # 中間データを書き出す（1テーブル=1ファイル）
    tables_dir = out_dir / TABLES_DIR
    tables_dir.mkdir(parents=True, exist_ok=True)
    for old in tables_dir.glob("*.yaml"):
        old.unlink()
    for name in order:
        _, source, table = adopted[name]
        data = {"source": source, **table}
        _write_yaml(tables_dir / f"{name}.yaml", data)

    meta = {
        "subsystem_id": subsystem_id,
        "subsystem_name": subsystem_name,
        "tables": order,
        "sources": [
            {"priority": p, "path": Path(path).as_posix()}
            for p, path in sorted(books, key=lambda b: b[0])
        ],
        "history": history
        or [
            {
                "version": "第1.0版",
                "date": date,
                "kind": "新規",
                "place": "-",
                "content": "案件ごとのテーブル定義書を統合して新規作成",
                "author": author,
            }
        ],
    }
    _write_yaml(out_dir / BOOK_META_FILE, meta)
    return {
        name: {"adopted": adopted[name][1], "rejected": rejected.get(name, [])}
        for name in order
    }


def _write_yaml(path, data):
    """dictをUTF-8・LF・キー順維持でYAMLへ書き出す。

    Args:
        path: 出力パス。
        data: 書き出す内容。
    """
    text = yaml.safe_dump(data, allow_unicode=True, sort_keys=False, width=1000)
    Path(path).write_bytes(text.encode("utf-8"))


def load_intermediate(src_dir):
    """中間データを読み込む。

    Args:
        src_dir: 中間データディレクトリ。

    Returns:
        tuple: (ブック情報dict, テーブル定義dictのリスト（並び順どおり）)。
    """
    src_dir = Path(src_dir)
    meta = yaml.safe_load((src_dir / BOOK_META_FILE).read_text(encoding="utf-8"))
    tables = []
    for name in meta["tables"]:
        data = yaml.safe_load(
            (src_dir / TABLES_DIR / f"{name}.yaml").read_text(encoding="utf-8")
        )
        data.pop("source")
        tables.append(data)
    return meta, tables


def _to_datetime(value):
    """YAMLの日付をExcelへ書くためにdatetimeへそろえる。

    Args:
        value: date または datetime。

    Returns:
        datetime.datetime: 変換後の値。
    """
    if isinstance(value, datetime.datetime):
        return value
    return datetime.datetime(value.year, value.month, value.day)


def _write_common_sheets(wb, meta):
    """表紙・変更履歴を中間データのブック情報で書き直す。

    Args:
        wb: 対象のワークブック。
        meta: ブック情報。
    """
    history = meta["history"]
    if len(history) > HISTORY_LAST_ROW - HISTORY_FIRST_ROW + 1:
        raise ValueError(
            f"変更履歴が記入欄（{HISTORY_LAST_ROW - HISTORY_FIRST_ROW + 1}行）を超えています"
        )

    # 変更履歴。記入欄を一度空にしてから、履歴を先頭から書く
    ws = wb["変更履歴"]
    ws["E3"] = meta["subsystem_name"]
    for i, row in enumerate(range(HISTORY_FIRST_ROW, HISTORY_LAST_ROW + 1)):
        entry = history[i] if i < len(history) else None
        ws[f"A{row}"] = i + 1 if entry else None
        for key, col in HISTORY_CELLS.items():
            value = entry[key] if entry else None
            ws[f"{col}{row}"] = (
                _to_datetime(value) if key == "date" and value else value
            )

    # 表紙。サブシステム名と最新の版数
    ws = wb["表紙"]
    ws["J16"] = meta["subsystem_name"]
    ws["J19"] = history[-1]["version"]


def _write_table_sheet(ws, table):
    """テーブルシート1枚へ中間データのテーブル定義を書く（記入欄は先に空にする）。

    Args:
        ws: 書き込み先のテーブルシート（土台シートの複製）。
        table: テーブル定義。
    """
    columns = table["columns"]
    if len(columns) > COLUMN_LAST_ROW - COLUMN_FIRST_ROW + 1:
        raise ValueError(
            f"{table['physical_name']} のカラム数({len(columns)})が記入欄"
            f"（{COLUMN_LAST_ROW - COLUMN_FIRST_ROW + 1}行）を超えています。"
            "先に extend_table_rows.py で土台シートの行を増やしてください"
        )

    # 見出し部分
    ws["F5"] = table["logical_name"]
    ws["W5"] = table["physical_name"]
    ws["F6"] = table["description"]
    for col, flag in zip(INDEX_COLUMNS, table["index_flags"]):
        ws[f"{col}{INDEX_FLAG_ROW}"] = flag

    # カラム記入欄。空行も含めて全行を書き直し、土台シートの値を残さない
    for i, row in enumerate(range(COLUMN_FIRST_ROW, COLUMN_LAST_ROW + 1)):
        column = columns[i] if i < len(columns) else None
        for key, col in COLUMN_CELLS.items():
            ws[f"{col}{row}"] = column[key] if column else None
        for j, col in enumerate(INDEX_COLUMNS):
            ws[f"{col}{row}"] = column["index"][j] if column else None


def _write_toc(ws, entries):
    """目次シートの記入欄を「n. 名称」の並びで書き直す。

    Args:
        ws: 目次シート。
        entries: 目次に並べる名称のリスト。
    """
    if len(entries) > TOC_LAST_ROW - TOC_FIRST_ROW + 1:
        raise ValueError(f"目次の項目数({len(entries)})が記入欄を超えています")
    for i, row in enumerate(range(TOC_FIRST_ROW, TOC_LAST_ROW + 1)):
        ws[f"B{row}"] = f"{i + 1}. {entries[i]}" if i < len(entries) else None


def build_definition(meta, tables, base_def, out_def):
    """中間データから統合版のテーブル定義書を生成する。

    Args:
        meta: ブック情報。
        tables: テーブル定義のリスト。
        base_def: 書式の土台にするテーブル定義書のパス。
        out_def: 出力パス。
    """
    wb = openpyxl.load_workbook(base_def)
    base_sheets = _table_sheets(wb)
    prototype = base_sheets[0]
    # 印刷範囲はシート複製で引き継がれないため、土台の範囲（シート名を除いた部分）を控える
    print_range = prototype.print_area.split("!")[-1] if prototype.print_area else None

    # 土台の先頭テーブルシートを複製して、テーブルごとのシートを作る
    new_sheets = []
    for table in tables:
        ws = wb.copy_worksheet(prototype)
        copy_data_validations(prototype, ws)
        _write_table_sheet(ws, table)
        new_sheets.append((ws, table["sheet"]))

    # 土台のテーブルシートを消してから、複製したシートに正式な名前を付ける
    for ws in base_sheets:
        wb.remove(ws)
    for ws, title in new_sheets:
        ws.title = title
        if print_range:
            ws.print_area = print_range.replace("$", "")

    # 複製したシートをデータシートの直前へ、中間データの並び順で置く
    for ws, _ in new_sheets:
        offset = wb.sheetnames.index("データ") - wb.sheetnames.index(ws.title)
        wb.move_sheet(ws.title, offset=offset)

    _write_toc(wb["目次"], [table["logical_name"] for table in tables])
    _write_common_sheets(wb, meta)
    save_with_shapes(wb, base_def, out_def)


def build_list(meta, tables, base_list, out_list):
    """中間データから統合版のテーブル一覧を生成する。

    Args:
        meta: ブック情報。
        tables: テーブル定義のリスト。
        base_list: 書式の土台にするテーブル一覧のパス。
        out_list: 出力パス。
    """
    if len(tables) > LIST_LAST_ROW - LIST_FIRST_ROW + 1:
        raise ValueError(f"テーブル数({len(tables)})が一覧の記入欄を超えています")

    wb = openpyxl.load_workbook(base_list)
    ws = wb[LIST_SHEET]

    # 一覧の記入欄。備考は案件ごとの経緯を書いていた欄のため、統合版では空にする
    for i, row in enumerate(range(LIST_FIRST_ROW, LIST_LAST_ROW + 1)):
        table = tables[i] if i < len(tables) else None
        ws[f"C{row}"] = i + 1 if table else None
        ws[f"D{row}"] = meta["subsystem_id"] if table else None
        ws[f"E{row}"] = meta["subsystem_name"] if table else None
        ws[f"J{row}"] = table["logical_name"] if table else None
        ws[f"S{row}"] = table["physical_name"] if table else None
        ws[f"AB{row}"] = None

    _write_common_sheets(wb, meta)
    save_with_shapes(wb, base_list, out_list)


def check(src_dir, def_path):
    """テーブル定義書を読み戻し、中間データと一致するかを検証する。

    Args:
        src_dir: 中間データディレクトリ。
        def_path: 検証するテーブル定義書のパス。

    Returns:
        list: 差異の説明（一致していれば空）。
    """
    _, expected = load_intermediate(src_dir)
    wb = openpyxl.load_workbook(def_path)
    actual = [_read_table_sheet(ws) for ws in _table_sheets(wb)]

    problems = []
    if [t["physical_name"] for t in actual] != [t["physical_name"] for t in expected]:
        problems.append("テーブルの並びが一致しません")
    for exp, act in zip(expected, actual):
        if exp != act:
            problems.append(f"{exp['physical_name']} の内容が一致しません")
    return problems


def main():
    ensure_utf8_stdio()
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    sub = parser.add_subparsers(dest="command", required=True)

    p_extract = sub.add_parser("extract")
    p_extract.add_argument("out_dir")
    p_extract.add_argument(
        "--book", nargs=2, action="append", required=True, metavar=("PRIORITY", "XLSX")
    )
    p_extract.add_argument("--author", required=True)
    p_extract.add_argument("--date", required=True, type=datetime.date.fromisoformat)
    p_extract.add_argument("--force", action="store_true")

    p_build = sub.add_parser("build")
    p_build.add_argument("src_dir")
    p_build.add_argument("--base-def", required=True)
    p_build.add_argument("--base-list", required=True)
    p_build.add_argument("--out-def", required=True)
    p_build.add_argument("--out-list", required=True)

    p_check = sub.add_parser("check")
    p_check.add_argument("src_dir")
    p_check.add_argument("def_path")

    args = parser.parse_args()

    if args.command == "extract":
        books = [(int(priority), path) for priority, path in args.book]
        result = extract(args.out_dir, books, args.author, args.date, args.force)
        print(f"=== 抽出結果: {args.out_dir}（{len(result)}テーブル） ===")
        for name, info in result.items():
            print(f"{name}: 採用 {info['adopted']}")
            for src in info["rejected"]:
                print(f"    不採用 {src}")
        return

    if args.command == "build":
        meta, tables = load_intermediate(args.src_dir)
        build_definition(meta, tables, args.base_def, args.out_def)
        build_list(meta, tables, args.base_list, args.out_list)
        print(f"=== 生成結果（{len(tables)}テーブル） ===")
        print(f"テーブル定義書: {args.out_def}")
        print(f"テーブル一覧  : {args.out_list}")

    # build の後、または check 単独で、読み戻し検証を行う
    def_path = args.out_def if args.command == "build" else args.def_path
    problems = check(args.src_dir, def_path)
    if problems:
        print("=== 検証NG ===")
        for p in problems:
            print(p)
        sys.exit(1)
    print("=== 検証OK（テーブル定義書の内容は中間データと一致） ===")


if __name__ == "__main__":
    main()
