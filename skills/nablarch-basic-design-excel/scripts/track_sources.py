# -*- coding: utf-8 -*-
"""
作業の元にした設計書が、作業中に誰かに直されていないかを確かめる。

【背景】設計書（Excel・draw.io）は、AIが作業している間に、人が開いて直すことがある。
それに気づかずに、作業前に読んだ内容をもとに書き出したり、ファイルサーバーへ上げたりすると、人が直した内容が消える。
そこで、作業を始めるときに元にしたファイルを記録し、書き出す前・上げる前に、今のファイルと比べる。

使い方:
    # 作業を始めるとき：元にしたファイルを記録する（作業フォルダの sources/ に写しを残す）
    python track_sources.py record <作業フォルダ> <ファイル> [<ファイル> ...]

    # 書き出す前・上げる前：記録したときから変わっていないかを比べる
    python track_sources.py check <作業フォルダ> [<ファイル> ...]

    record は、同じファイルをもう一度記録すると、記録を今の状態に置き換える。
    自分で書き出した後や、人が直した内容を取り込んだ後は、record し直して比べる基準を今の状態にする。
    check でファイルを指定しなければ、記録したすべてのファイルを比べる。
    ファイルサーバー上のファイル（\\\\サーバー\\共有\\...）も、そのまま指定できる。

判定のしかた:
    1. ファイルのハッシュ値（SHA-256）が記録と同じなら「変わっていない」。
    2. 違うときは、中身を比べる。
       - Excel（.xlsx）：シートの構成、セルの値・数式、書式（結合セル・背景色・文字色・罫線・行の高さ・列の幅・入力規則）。
       - draw.io（.drawio）：ページと図形（文字・スタイル・位置・線のつながり）。保存した時刻やアプリの情報は比べない。
       - それ以外：ハッシュ値だけで判定する。
    3. 中身が同じなら「保存し直しただけ（中身は同じ）」として、変わっていないものと扱う。
       中身が違えば「直されている」として、どこが変わったかを表示する。

終了コード:
    0 = すべて変わっていない（保存し直しただけを含む）
    1 = 直されているファイル、または見つからないファイルがある（作業を止めてユーザーに確かめること）
"""

import argparse
import base64
import datetime
import hashlib
import json
import shutil
import sys
import urllib.parse
import xml.etree.ElementTree as ET
import zlib
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent))
from xlsx_common import ensure_utf8_stdio, snapshot_sheet_structure

# 記録ファイルと、写しを置くフォルダの名前
MANIFEST = "sources.json"
COPY_DIR = "sources"
# 行の高さ・列の幅の違いを無視する幅。Excelは保存のたびに、ピクセル単位に丸め直す（例：24.3 → 24.4）
SIZE_TOLERANCE = 0.3
# 1つのファイルについて表示する変更の上限（多すぎると読めないため）
MAX_REPORT = 30


def _sha256(path):
    """ファイルのSHA-256を返す。"""
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _load_manifest(work_dir):
    """記録ファイルを読む。無ければ空の記録を返す。"""
    path = work_dir / MANIFEST
    if not path.exists():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def _save_manifest(work_dir, manifest):
    """記録ファイルを書く。"""
    (work_dir / MANIFEST).write_text(
        json.dumps(manifest, ensure_ascii=False, indent=1),
        encoding="utf-8",
        newline="\n",
    )


def record(work_dir, files):
    """ファイルの写しを残し、ハッシュ値・大きさ・更新日時を記録する。

    Args:
        work_dir (Path): 作業フォルダ。
        files (list[str]): 記録するファイル。
    """
    manifest = _load_manifest(work_dir)
    copy_dir = work_dir / COPY_DIR
    copy_dir.mkdir(parents=True, exist_ok=True)
    for f in files:
        src = Path(f)
        if not src.is_file():
            print(f"NG: ファイルがありません: {f}")
            sys.exit(1)
        key = (
            str(src.resolve()) if not str(src).startswith(("\\\\", "//")) else str(src)
        )
        digest = _sha256(src)
        # 前の写しは消し、今の写しに置き換える
        old = manifest.get(key)
        if old and (work_dir / old["copy"]).exists():
            (work_dir / old["copy"]).unlink()
        copy_name = f"{COPY_DIR}/{digest[:12]}_{src.name}"
        shutil.copy2(src, work_dir / copy_name)
        stat = src.stat()
        manifest[key] = {
            "sha256": digest,
            "size": stat.st_size,
            "mtime": datetime.datetime.fromtimestamp(stat.st_mtime).isoformat(
                timespec="seconds"
            ),
            "recorded_at": datetime.datetime.now().isoformat(timespec="seconds"),
            "copy": copy_name,
        }
        print(f"記録した: {src.name}（更新日時 {manifest[key]['mtime']}）")
    _save_manifest(work_dir, manifest)


def _xlsx_snapshot(path):
    """Excelの中身（シートの並び・セルの値と数式・書式）を比べられる形で取り出す。"""
    wb = openpyxl.load_workbook(path)
    sheets = {}
    for ws in wb.worksheets:
        values = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    values[cell.coordinate] = str(cell.value)
        structure = snapshot_sheet_structure(wb, ws)
        # 列の幅は、Excelが保存のたびに「どの列にまとめて設定するか」を変えるため、1列ずつに展開して比べる
        structure["col_widths"] = _col_widths(ws)
        sheets[ws.title] = {"values": values, "structure": structure}
    return wb.sheetnames, sheets


def _col_widths(ws):
    """列の幅を、まとめて設定された範囲（min〜max）も1列ずつに展開して返す。"""
    widths = {}
    for key, dim in ws.column_dimensions.items():
        if dim.width is None:
            continue
        low = dim.min or openpyxl.utils.column_index_from_string(key)
        high = dim.max or low
        for idx in range(low, min(high, 200) + 1):
            widths[openpyxl.utils.get_column_letter(idx)] = dim.width
    return widths


def _diff_sizes(name, before, after):
    """行の高さ・列の幅の違いを、丸めの差（SIZE_TOLERANCE以下）を除いて返す。"""
    out = []
    for k in sorted(set(before) | set(after), key=str):
        a, b = before.get(k), after.get(k)
        if a is not None and b is not None and abs(a - b) <= SIZE_TOLERANCE:
            continue
        if a != b:
            out.append(f"{name} {k}: {a!r} → {b!r}")
    return out


def _diff_dict(name, before, after):
    """2つの辞書の違いを、キーごとの説明にして返す。"""
    out = []
    for k in sorted(set(before) | set(after), key=str):
        if before.get(k) != after.get(k):
            out.append(f"{name} {k}: {before.get(k)!r} → {after.get(k)!r}")
    return out


def _diff_xlsx(old_path, new_path):
    """Excelの中身の違いを返す（空なら中身は同じ）。"""
    old_names, old = _xlsx_snapshot(old_path)
    new_names, new = _xlsx_snapshot(new_path)
    diffs = []
    if old_names != new_names:
        diffs.append(f"シートの構成: {old_names} → {new_names}")
    for title in [t for t in new_names if t in old]:
        o, n = old[title], new[title]
        diffs += [f"[{title}] " + d for d in _diff_dict("値", o["values"], n["values"])]
        for key in o["structure"]:
            a, b = o["structure"][key], n["structure"][key]
            if a == b:
                continue
            if key in ("row_heights", "col_widths"):
                diffs += [f"[{title}] " + d for d in _diff_sizes(key, a, b)]
            elif isinstance(a, dict) and isinstance(b, dict):
                diffs += [f"[{title}] " + d for d in _diff_dict(key, a, b)]
            else:
                diffs.append(f"[{title}] {key} が変わった")
    return diffs


def _drawio_model(diagram):
    """ページの図形（mxGraphModel）を返す。

    draw.ioは、設定によってページの中身を圧縮して保存する（diagram要素の文字が、
    deflate圧縮してbase64にしたもの）。その場合は元に戻してから読む。
    """
    model = diagram.find("mxGraphModel")
    if model is not None:
        return model
    text = (diagram.text or "").strip()
    if not text:
        return ET.Element("mxGraphModel")
    raw = zlib.decompress(base64.b64decode(text), -15).decode("utf-8")
    return ET.fromstring(urllib.parse.unquote(raw))


def _drawio_snapshot(path):
    """draw.ioの中身（ページごとの図形）を、保存の情報を除いて取り出す。"""
    root = ET.parse(path).getroot()
    pages = {}
    for d in root.findall("diagram"):
        cells = {}
        for c in _drawio_model(d).iter("mxCell"):
            geo = c.find("mxGeometry")
            cells[c.get("id")] = {
                "value": c.get("value"),
                "style": c.get("style"),
                "source": c.get("source"),
                "target": c.get("target"),
                "parent": c.get("parent"),
                "geometry": (
                    None if geo is None else ET.tostring(geo, encoding="unicode")
                ),
            }
        pages[d.get("name")] = cells
    return [d.get("name") for d in root.findall("diagram")], pages


def _diff_drawio(old_path, new_path):
    """draw.ioの中身の違いを返す（空なら中身は同じ）。"""
    old_names, old = _drawio_snapshot(old_path)
    new_names, new = _drawio_snapshot(new_path)
    diffs = []
    if old_names != new_names:
        diffs.append(f"ページの構成: {old_names} → {new_names}")
    for page in [p for p in new_names if p in old]:
        o, n = old[page], new[page]
        for cid in sorted(set(o) | set(n), key=str):
            if cid not in n:
                diffs.append(f"[{page}] 図形 {cid} が消えた")
            elif cid not in o:
                diffs.append(
                    f"[{page}] 図形 {cid} が増えた: {str(n[cid]['value'])[:60]}"
                )
            else:
                changed = [k for k in o[cid] if o[cid][k] != n[cid][k]]
                if changed:
                    diffs.append(
                        f"[{page}] 図形 {cid} の {'・'.join(changed)} が変わった"
                    )
    return diffs


def check(work_dir, files):
    """記録したときから変わっていないかを比べ、結果を表示する。

    Returns:
        bool: すべて変わっていなければ True。
    """
    manifest = _load_manifest(work_dir)
    if not manifest:
        print(
            f"NG: 記録がありません（先に record してください）: {work_dir / MANIFEST}"
        )
        return False
    keys = list(manifest)
    if files:
        wanted = {
            str(Path(f).resolve()) if not f.startswith(("\\\\", "//")) else f
            for f in files
        }
        keys = [k for k in keys if k in wanted]
        missing = wanted - set(keys)
        for m in sorted(missing):
            print(f"NG: 記録されていないファイルです: {m}")
        if missing:
            return False

    ok = True
    for key in keys:
        entry = manifest[key]
        path = Path(key)
        name = path.name
        if not path.is_file():
            print(f"NG 見つからない: {name}")
            ok = False
            continue
        if _sha256(path) == entry["sha256"]:
            print(f"OK 変わっていない: {name}")
            continue
        copy = work_dir / entry["copy"]
        suffix = path.suffix.lower()
        if suffix == ".xlsx":
            diffs = _diff_xlsx(copy, path)
        elif suffix == ".drawio":
            diffs = _diff_drawio(copy, path)
        else:
            diffs = [
                "ファイルの中身が変わった（このファイルの種類は、中身の比較に対応していない）"
            ]
        mtime = datetime.datetime.fromtimestamp(path.stat().st_mtime).isoformat(
            timespec="seconds"
        )
        if not diffs:
            print(
                f"OK 保存し直しただけ（中身は同じ）: {name}（更新日時 {entry['mtime']} → {mtime}）"
            )
            continue
        ok = False
        print(f"NG 直されている: {name}（更新日時 {entry['mtime']} → {mtime}）")
        for d in diffs[:MAX_REPORT]:
            print(f"    - {d}")
        if len(diffs) > MAX_REPORT:
            print(f"    …ほか {len(diffs) - MAX_REPORT} 件")
    return ok


def main():
    ensure_utf8_stdio()
    parser = argparse.ArgumentParser(
        description="作業の元にした設計書が、作業中に直されていないかを確かめる"
    )
    sub = parser.add_subparsers(dest="command", required=True)
    p_rec = sub.add_parser("record")
    p_rec.add_argument("work_dir")
    p_rec.add_argument("files", nargs="+")
    p_chk = sub.add_parser("check")
    p_chk.add_argument("work_dir")
    p_chk.add_argument("files", nargs="*")
    args = parser.parse_args()

    work_dir = Path(args.work_dir)
    work_dir.mkdir(parents=True, exist_ok=True)
    if args.command == "record":
        record(work_dir, args.files)
    else:
        sys.exit(0 if check(work_dir, args.files) else 1)


if __name__ == "__main__":
    main()
