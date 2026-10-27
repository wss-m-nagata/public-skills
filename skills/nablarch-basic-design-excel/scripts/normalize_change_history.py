# -*- coding: utf-8 -*-
"""
「変更履歴」シートの、9〜33行目の結合と、10〜33行目の書式を、そろえる。

【背景】変更履歴は、版を重ねると、14行目以降も使う。
古いテンプレートでは、8〜13行目だけセルが結合されていて（B:C・D:F・G:I・J:P・Q:AE）、14行目以降は結合が外れている。
そのまま行を足すと、版数・変更日・区分・変更箇所・変更内容が、結合されていないセルに入り、表示が崩れる。
このスクリプトは、行を足す前に、次のようにそろえる。

  - 結合：9〜33行目を、8行目と同じ結合にする。
  - 書式（罫線・配置・フォント）：10〜33行目を、9行目と同じ書式にする。

書式を8行目から写さないのは、8行目の上の罫線が、ヘッダーの下線（二重線）になっているためである。
8行目の書式を写すと、各行の上が二重線になる（実際に起きた）。9行目は、細い1本線の通常の行である。

すでにそろっているシート（新しいテンプレート）では、何も変えずに終わる（出力は元のファイルと同じ内容）。

使い方:
    python normalize_change_history.py <入力xlsx> <出力xlsx>

変更履歴の行を足す前（apply_mapping.py の前）に流す。
流したあとのファイルを、構造の比べる基準にする（結合が増えるため、流す前のファイルとは比べない）。

終了コード:
    0 = 正常終了（変えたとき・変える必要が無いとき、どちらも）
"""

import sys
from copy import copy
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent))
from xlsx_common import ensure_utf8_stdio, save_with_shapes  # noqa: E402

# 変更履歴シートの、8行目（初版の行。結合の見本）、9行目（書式の見本）、そろえる範囲の最終行
SHEET_NAME = "変更履歴"
MERGE_ROW = 8
STYLE_ROW = 9
LAST_ROW = 33


def normalize(input_path, output_path):
    """9〜33行目の結合と、10〜33行目の書式をそろえて保存する。

    Returns:
        bool: 変えたとき True。変える必要が無く、入力をそのまま出力に写したとき False。
    """
    wb = openpyxl.load_workbook(input_path)
    ws = wb[SHEET_NAME]

    # 8行目の結合の並び（列の範囲）
    pattern = sorted(
        (m.min_col, m.max_col)
        for m in ws.merged_cells.ranges
        if m.min_row == MERGE_ROW and m.max_row == MERGE_ROW
    )
    existing = {(m.min_row, m.min_col, m.max_col) for m in ws.merged_cells.ranges}

    changed = False
    for row in range(MERGE_ROW + 1, LAST_ROW + 1):
        # 結合が足りない行だけを直す。すでにそろっている行（新しいテンプレートなど）は、書式も含めて、何も変えない
        missing = [
            (first, last)
            for first, last in pattern
            if (row, first, last) not in existing
        ]
        if not missing:
            continue

        # 書式を、9行目にそろえる（9行目自身は見本のため写さない）。
        # 結合されているセルの左上以外は、書式を持たないため飛ばす
        if row > STYLE_ROW:
            for column in range(1, ws.max_column + 1):
                source = ws.cell(row=STYLE_ROW, column=column)
                target = ws.cell(row=row, column=column)
                if type(target).__name__ != "MergedCell":
                    target._style = copy(source._style)
        # 結合を、8行目にそろえる
        for first, last in missing:
            ws.merge_cells(
                start_row=row, start_column=first, end_row=row, end_column=last
            )
        changed = True

    if changed:
        save_with_shapes(wb, input_path, output_path)
    else:
        Path(output_path).write_bytes(Path(input_path).read_bytes())
    return changed


def main():
    ensure_utf8_stdio()
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(2)

    input_path, output_path = sys.argv[1:3]
    changed = normalize(input_path, output_path)
    print(
        "変更履歴の9〜33行目の結合と、10〜33行目の書式をそろえました。"
        if changed
        else "すでにそろっています（変更なし）。"
    )
    print(f"=== 出力: {output_path} ===")


if __name__ == "__main__":
    main()
