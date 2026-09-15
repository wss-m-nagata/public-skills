# -*- coding: utf-8 -*-
"""シートの途中へ行を挿入して、表を広げるスクリプト。

`extend_table_rows.py` との使い分け:
- `extend_table_rows.py` … 表の下が空き領域のとき。行は挿入せず、最終行の書式を下へ複製する。
- 本スクリプト … 表の下に別の表やブロックがあるとき。行を挿入して下の内容を押し下げる。

外部インタフェース設計書(JSON)の『【レコード名】』シートのように、1シートへ
ブロックが縦に3つ並ぶレイアウトでは、`extend_table_rows.py` は使えない。
下のブロックの領域を上書きし、既存の結合セルと衝突して
`MergedCell object attribute 'value' is read-only` で失敗する。

openpyxl の `insert_rows` は結合セル・行高・入力規則を動かさないため、本スクリプトで補う。

使い方（スクリプトとして）:
    python insert_rows.py <元xlsxパス> <シート名> <挿入位置の行番号> <追加する行数>         <書式の見本にする行番号> <出力xlsxパス>

呼び出し側で使う場合は `insert_rows_keep_layout(ws, at_row, count, style_row)` を使う。

例(『【レコード名】』シートの17行目の直前に、16行目の書式で3行挿入する場合):
    python insert_rows.py 外部インタフェース設計書_xxx.xlsx "【レコード名】" 17 3 16 出力.xlsx
"""

from copy import copy

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import CellRange


def insert_rows_keep_layout(ws, at_row, count, style_row, max_col=40):
    """at_row の直前に count 行を挿入し、style_row の書式を複製する。

    at_row 以降にある結合セル・行高・入力規則も同じだけ下へずらす。
    """
    if count <= 0:
        return

    # 1) 結合セルをいったん全て解除して記録する
    merges = [str(m) for m in list(ws.merged_cells.ranges)]
    for m in merges:
        ws.unmerge_cells(m)

    # 2) 入力規則の適用範囲を記録して外す
    validations = []
    for dv in list(ws.data_validations.dataValidation):
        validations.append((dv, [str(r) for r in dv.sqref.ranges]))
        ws.data_validations.dataValidation.remove(dv)

    # 3) 行高を記録する
    heights = {
        r: d.height for r, d in ws.row_dimensions.items() if d.height is not None
    }

    ws.insert_rows(at_row, count)

    # 4) 結合セルを戻す（挿入位置より下はずらす）
    for m in merges:
        rng = CellRange(m)
        if rng.min_row >= at_row:
            rng.shift(row_shift=count)
        ws.merge_cells(str(rng))

    # 5) 入力規則を戻す
    for dv, ranges in validations:
        shifted = []
        for r in ranges:
            rng = CellRange(r)
            if rng.min_row >= at_row:
                rng.shift(row_shift=count)
            shifted.append(str(rng))
        dv.sqref = " ".join(shifted)
        ws.add_data_validation(dv)

    # 6) 行高を戻す
    for r in sorted(heights, reverse=True):
        if r >= at_row:
            ws.row_dimensions[r + count].height = heights[r]
    for r in range(at_row, at_row + count):
        if style_row in heights:
            ws.row_dimensions[r].height = heights[style_row]

    # 7) 挿入した行へ、見本の行の書式を複製する
    src = style_row + count if style_row >= at_row else style_row
    for r in range(at_row, at_row + count):
        for c in range(1, max_col + 1):
            s = ws.cell(row=src, column=c)
            d = ws.cell(row=r, column=c)
            d._style = copy(s._style)
        # 見本の行にある1行分の結合セルを、挿入した行にも作る
        for m in [
            x
            for x in list(ws.merged_cells.ranges)
            if x.min_row == src and x.max_row == src
        ]:
            ws.merge_cells(
                start_row=r, start_column=m.min_col, end_row=r, end_column=m.max_col
            )


def clear_rows(ws, rows, max_col=40):
    """値・罫線・塗り・結合を消して、使っていない領域を空白にする。"""
    from openpyxl.styles import Border, PatternFill

    for rng in [
        m
        for m in list(ws.merged_cells.ranges)
        if rows[0] <= m.min_row and m.max_row <= rows[-1]
    ]:
        ws.unmerge_cells(str(rng))
    for r in rows:
        for c in range(1, max_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.value = None
            cell.border = Border()
            cell.fill = PatternFill(fill_type=None)


def col(idx):
    return get_column_letter(idx)


if __name__ == "__main__":
    import sys

    import openpyxl

    sys.path.insert(0, str(__import__("pathlib").Path(__file__).parent))
    from xlsx_common import save_with_shapes, ensure_utf8_stdio

    ensure_utf8_stdio()
    src, sheet, at_row, count, style_row, out = sys.argv[1:7]
    wb = openpyxl.load_workbook(src)
    insert_rows_keep_layout(wb[sheet], int(at_row), int(count), int(style_row))
    save_with_shapes(wb, src, out)
    print("inserted %s rows at %s of %s -> %s" % (count, at_row, sheet, out))
