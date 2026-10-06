# -*- coding: utf-8 -*-
"""
「変更履歴」シートの行の高さを、書いた内容の行数に合わせる。

【背景】変更履歴シートは行の高さが15で固定されており、変更箇所(J〜P列)・変更内容(Q〜AE列)は
結合セルのため、セル内で改行してもExcelは行の高さを自動で広げない。そのため2行目以降が隠れる。
このスクリプトは、表示される行数を数え、行の高さを「15×行数」に設定する(1行=15、2行=30、3行=45…)。

処理内容:
    - 対象は「変更履歴」シートの8行目以降(データの行)。
    - セル内の改行をLFにそろえる(CRLF・CRが混じっていればLFに直す。Excelのセル内改行はLFのため)。
    - 変更箇所(J列)と変更内容(Q列)の「折り返して全体を表示」をオンにし、上詰めにする。
    - 行数を数える。行数は、セル内改行で分けた各行について、列の幅に収まらずExcelが自動で
      折り返す分も含めて数える(1行に入る文字数は、実際の表示で測った値から見積もる)。
    - 行の高さを「15×行数」にする(J列とQ列のうち、行数の多いほうに合わせる)。

1行に入る文字数の見積もり:
    2026-10-05、WebサービスAPI一覧・外部インタフェース設計書(JSON)の変更履歴をExcelでPDFにして測ったところ、
    変更内容(Q〜AE列、列幅4.83×15列＝合計約72)の1行に入るのは、表示幅(半角1・全角2)でおよそ71〜74だった
    (列幅の合計とほぼ同じ)。英数字は単語の区切りで折り返されて1行が短くなることがあるため、0.95倍として見積もる。
    列幅は、複数の列にまとめて設定されている場合(openpyxlでは先頭の列だけに min・max 付きで入る)もあるため、
    その範囲も含めて読む。

使い方:
    python fit_change_history_rows.py <入力xlsxパス> <出力xlsxパス>

実行後は、構造の検証で行の高さの変化を許す指定を付ける:
    python verify_layout.py <入力xlsxパス> <出力xlsxパス> --allow-history-row-heights
"""

import math
import sys
import unicodedata
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment

sys.path.insert(0, str(Path(__file__).parent))
from xlsx_common import ensure_utf8_stdio, save_with_shapes

# 対象のシート名
SHEET_NAME = "変更履歴"
# データの先頭行(7行目は見出し)
FIRST_DATA_ROW = 8
# 1行あたりの行の高さ(テンプレートの既定値)
LINE_HEIGHT = 15.0
# 行数を数える列(変更箇所・変更内容)と、その結合範囲の列
TARGET_COLUMNS = {"J": ("J", "P"), "Q": ("Q", "AE")}
# 列幅の合計に対する、1行に入る表示幅の割合(実際の表示で測った値。上のdocstring参照)
CHARS_PER_WIDTH = 0.95


def _display_width(text):
    """文字列の表示幅を、半角1・全角2として数える。

    Args:
        text (str): 対象の文字列。

    Returns:
        int: 表示幅。
    """
    return sum(
        2 if unicodedata.east_asian_width(ch) in ("F", "W", "A") else 1 for ch in text
    )


def _column_width(ws, idx):
    """列の幅を返す。

    列幅は、複数の列にまとめて設定されている場合がある(openpyxlでは、範囲の先頭の列の
    column_dimensions に min・max 付きで入る)。そのため、範囲に idx を含む設定を探す。
    column_dimensions[letter] と書くと、設定が無い列にも既定値の設定が作られて列幅が変わるため使わない。

    Args:
        ws: 対象のワークシート。
        idx (int): 列番号(1始まり)。

    Returns:
        float: 列幅(Excelの列幅の単位)。設定が無ければシートの既定の列幅。
    """
    for dim in ws.column_dimensions.values():
        low = dim.min or openpyxl.utils.column_index_from_string(dim.index)
        high = dim.max or low
        if low <= idx <= high and dim.width:
            return dim.width
    return ws.sheet_format.defaultColWidth or 8.43


def _merged_width(ws, first_col, last_col):
    """結合範囲の列幅の合計を返す。

    Args:
        ws: 対象のワークシート。
        first_col (str): 結合範囲の先頭の列。
        last_col (str): 結合範囲の末尾の列。

    Returns:
        float: 列幅の合計(Excelの列幅の単位)。
    """
    start = openpyxl.utils.column_index_from_string(first_col)
    end = openpyxl.utils.column_index_from_string(last_col)
    return sum(_column_width(ws, idx) for idx in range(start, end + 1))


def fit_rows(ws):
    """変更履歴シートの改行をLFにそろえ、行の高さを表示される行数に合わせる。

    Args:
        ws: 「変更履歴」シート。

    Returns:
        list[str]: 行の高さを変えた行の説明。
    """
    changed = []
    # 1行に入る表示幅(列ごと)
    capacities = {
        col: _merged_width(ws, *span) * CHARS_PER_WIDTH
        for col, span in TARGET_COLUMNS.items()
    }

    for row in range(FIRST_DATA_ROW, ws.max_row + 1):
        line_counts = []
        for col in TARGET_COLUMNS:
            cell = ws[f"{col}{row}"]
            if not isinstance(cell.value, str):
                continue

            # 改行をLFにそろえる
            text = cell.value.replace("\r\n", "\n").replace("\r", "\n")
            if text != cell.value:
                cell.value = text

            # 折り返しをオンにし、上詰めにする(既存の配置の他の設定は引き継ぐ)
            align = cell.alignment
            cell.alignment = Alignment(
                horizontal=align.horizontal,
                vertical="top",
                wrap_text=True,
                indent=align.indent,
                shrink_to_fit=align.shrink_to_fit,
                text_rotation=align.text_rotation,
            )

            # 表示される行数を数える(列の幅に収まらず自動で折り返す分も含める)
            lines = text.split("\n")
            line_counts.append(
                sum(
                    max(1, math.ceil(_display_width(line) / capacities[col]))
                    for line in lines
                )
            )

        if not line_counts:
            continue

        # 行の高さを「15×行数」にする
        lines_needed = max(line_counts)
        height = LINE_HEIGHT * lines_needed
        current = ws.row_dimensions[row].height
        if current != height:
            ws.row_dimensions[row].height = height
            changed.append(f"{row}行目: {current} -> {height}（{lines_needed}行）")

    return changed


def main():
    ensure_utf8_stdio()
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(2)

    input_path, output_path = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(input_path)
    if SHEET_NAME not in wb.sheetnames:
        print(f"NG: 「{SHEET_NAME}」シートがありません。")
        sys.exit(1)

    changed = fit_rows(wb[SHEET_NAME])
    save_with_shapes(wb, input_path, output_path)

    print(f"OK: 行の高さを変えた行 {len(changed)} 件")
    for line in changed:
        print(f"  - {line}")


if __name__ == "__main__":
    main()
