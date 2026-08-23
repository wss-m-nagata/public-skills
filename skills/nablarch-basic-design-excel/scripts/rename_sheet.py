# -*- coding: utf-8 -*-
"""
既存のシートを複製せず、名前だけを変更するスクリプト。

「1インスタンス＝1シート」型テンプレートで、テンプレートのプレースホルダー名
（例：「2. 取引ID（取引名）」「【レコード名】」）を実際の値
（例：「2. E10101（評価依頼受付）」「3.1.顧客一覧」）に置き換える際に使う。
プレースホルダーの正確な文字列（角括弧の有無・括弧の全角半角）はテンプレートごとに異なるため、
対象の記入要領に記載された表記をそのまま使うこと（目次に書かれた角括弧表記とは異なる場合がある）。

duplicate_sheet.pyとの違い:
- duplicate_sheet.pyは既存シートを複製して新しいシートを追加する（元のシートは残る）。
  テーブル定義書の『論理テーブル名』シートのように、同じテンプレート内で複数インスタンスを
  作る場合に使う。
- rename_sheet.pyは既存シートの名前を変更するだけで、複製は行わない。
  システム機能設計書のように「1ファイル＝1インスタンス」で、唯一のシートの名前だけを
  実際の値に置き換えたい場合に使う。

使い方:
    python rename_sheet.py <元xlsxパス> <現在のシート名> <新しいシート名> <出力xlsxパス>

例:
    python rename_sheet.py システム機能設計書.xlsx "2. 取引ID（取引名）" "2. E10101（評価依頼受付）" 出力.xlsx
"""
import sys
import argparse
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent))
from xlsx_common import save_with_shapes, ensure_utf8_stdio

MAX_SHEET_NAME_LENGTH = 31


def rename_sheet(xlsx_path, old_name, new_name, output_path):
    wb = openpyxl.load_workbook(xlsx_path)
    if old_name not in wb.sheetnames:
        raise ValueError(f"シートが存在しません: {old_name}（現在のシート構成: {wb.sheetnames}）")
    if new_name != old_name and new_name in wb.sheetnames:
        raise ValueError(f"シート名が既に存在します(重複はできません): {new_name}")
    if len(new_name) > MAX_SHEET_NAME_LENGTH:
        raise ValueError(
            f"シート名がExcelの上限(31文字)を超えています(現在{len(new_name)}文字): {new_name}\n"
            f"「Webサービス」等の共通部分を削るか、取引名を要約して短縮すること。"
        )

    ws = wb[old_name]
    try:
        ws.title = new_name
    except ValueError as e:
        raise ValueError(
            f"シート名『{new_name}』を設定できません({e})。"
            f"シート名には `[ ] : * ? / \\` を使用できない。"
        )

    save_with_shapes(wb, xlsx_path, output_path)
    return {"old_name": old_name, "new_name": new_name, "sheet_order": wb.sheetnames}


def main():
    ensure_utf8_stdio()
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("xlsx_path")
    parser.add_argument("old_name")
    parser.add_argument("new_name")
    parser.add_argument("output_path")
    args = parser.parse_args()

    result = rename_sheet(args.xlsx_path, args.old_name, args.new_name, args.output_path)
    print(f"=== シートリネーム結果: {args.output_path} ===")
    print(f"『{result['old_name']}』 -> 『{result['new_name']}』")
    print(f"シート構成: {result['sheet_order']}")
    print("\n※目次シートにこのシート名への参照がある場合は、手動での更新も忘れずに。")


if __name__ == "__main__":
    main()
