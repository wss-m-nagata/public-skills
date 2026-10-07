# -*- coding: utf-8 -*-
"""
指定したシート・行の範囲で、文字が隠れないように行の高さを広げる。

【背景】WebサービスAPI一覧の「処理概要」やテーブル定義書の「項目定義」「備考」は結合セルのため、
長い文を書いてもExcelは行の高さを自動で広げない。そのため、2行目以降が隠れる。
fit_change_history_rows.py は「変更履歴」シート専用のため、それと同じ数え方を、任意のシートと行に使えるようにした。

処理内容:
    - 対象の行にある文字列のセル（結合セルは左上のセル）ごとに、表示される行数を数える。
      行数は、セル内改行で分けた各行について、セル（結合範囲）の幅に収まらずExcelが自動で折り返す分も含めて数える。
      1行に入る表示幅は、「列幅の合計×0.9」（半角1・全角2）で見積もる（fit_change_history_rows.py の0.95より少し控えめ）。
      英数字の続き（テーブル名・URLなど）は、Excelが途中で折り返さず丸ごと次の行へ送る。
      そのため、この折り返し方をまねて行数を数える（単純に幅で割ると、行数が足りない例が2026-10-07にあった）。
    - その行で最も行数の多いセルに合わせて、行の高さを「1行の高さ×行数」にする。
      1行の高さは、--line-height の指定、無ければ「文字の大きさ×1.35」とシートの既定の行の高さの大きいほうを使う。
      （9ポイントの文字なら約12。シートの既定の行の高さ11.25では、4行目以降が少しずつ欠けることを2026-10-07にPDFで確かめた）
    - 行の高さは広げるだけで、狭めない（1行に収まる行は変えない）。
    - 2行以上になるセルは、「折り返して全体を表示」をオンにする。
    - 複数の行にまたがる結合セル（見出しなど）は対象にしない。

使い方:
    python fit_row_heights.py <入力xlsxパス> <出力xlsxパス> --rows <開始行>-<終了行> [--rows ...]
        (--sheet <シート名> [--sheet ...] | --body-sheets) [--line-height <高さ>]

    --sheet       : 対象のシート名。複数指定できる。
    --body-sheets : 「表紙」「変更履歴」「目次」「データ」以外のすべてのシートを対象にする
                    （テーブル定義書の各テーブルのシートなど）。
    --rows        : 対象の行の範囲。複数指定できる。
    --line-height : 1行の高さ（ポイント）。省略時は上の説明のとおり。

テンプレートごとの指定例:
    # WebサービスAPI一覧（シート「1」の一覧の行）
    python fit_row_heights.py 入力.xlsx 出力.xlsx --sheet 1 --rows 18-38

    # テーブル定義書（各テーブルのシートの、テーブル説明の行とカラムの行）
    python fit_row_heights.py 入力.xlsx 出力.xlsx --body-sheets --rows 6-6 --rows 11-200

実行後は、構造の検証で行の高さが広がったことを許す指定を付ける:
    python verify_layout.py <入力xlsxパス> <出力xlsxパス> --allow-row-heights
"""

import argparse
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment

sys.path.insert(0, str(Path(__file__).parent))
from xlsx_common import ensure_utf8_stdio, save_with_shapes
from fit_change_history_rows import _display_width, _merged_width

# 列幅の合計に対する、1行に入る表示幅の割合。
# 変更履歴（fit_change_history_rows.py）は0.95だが、幅の狭い備考欄などで1行足りない例があった（2026-10-07、PDFで確認）ため、少し控えめにする
CHARS_PER_WIDTH = 0.9
# 文字の大きさに対する1行の高さの割合（Excelの既定の行の高さはおおよそ文字の大きさ×1.35）
LINE_HEIGHT_PER_FONT_SIZE = 1.35
# --body-sheets で対象にしない共通のシート
COMMON_SHEETS = ("表紙", "変更履歴", "目次", "データ")


def _parse_rows(text):
    """「開始行-終了行」を (開始行, 終了行) にする。

    Args:
        text (str): 「18-38」のような文字列。

    Returns:
        tuple[int, int]: 開始行と終了行。
    """
    first, last = text.split("-")
    return int(first), int(last)


# Excelが途中で折り返さない、英数字の続き（単語）
WORD = re.compile(r"[A-Za-z0-9_.:/@#%&=+\-]+")


def _count_lines(text, capacity):
    """折り返しをまねて、文字列が何行で表示されるかを数える。

    英数字の続きは1語として扱い、行の残りに入らなければ次の行へ送る（1行より長い語だけは途中で切る）。
    それ以外の文字（日本語・記号）は、どこでも折り返せるものとして1文字ずつ置く。

    Args:
        text (str): 対象の文字列（セル内改行はLF）。
        capacity (float): 1行に入る表示幅。

    Returns:
        int: 表示される行数。
    """
    total = 0
    for part in text.split("\n"):
        lines, used = 1, 0.0
        pos = 0
        while pos < len(part):
            match = WORD.match(part, pos)
            token = match.group(0) if match else part[pos]
            width = _display_width(token)
            if width > capacity:
                # 1行より長い語は、新しい行から始め、途中で切って置く（Excelの表示に合わせる）
                if used > 0:
                    lines, used = lines + 1, 0.0
                for ch in token:
                    if used + _display_width(ch) > capacity:
                        lines, used = lines + 1, 0.0
                    used += _display_width(ch)
            else:
                # 語が行の残りに入らなければ、次の行へ送る
                if used + width > capacity:
                    lines, used = lines + 1, 0.0
                used += width
            pos += len(token)
        total += lines
    return total


def _single_row_spans(ws):
    """1行だけの結合セルについて、左上のセル → (先頭の列, 末尾の列) の対応を返す。

    Args:
        ws: 対象のワークシート。

    Returns:
        tuple[dict, set]: 結合範囲の対応と、複数行にまたがる結合範囲に含まれるセル座標の集合。
    """
    spans = {}
    multi_row = set()
    for rng in ws.merged_cells.ranges:
        if rng.min_row == rng.max_row:
            spans[(rng.min_row, rng.min_col)] = (rng.min_col, rng.max_col)
        else:
            for row in range(rng.min_row, rng.max_row + 1):
                for col in range(rng.min_col, rng.max_col + 1):
                    multi_row.add((row, col))
    return spans, multi_row


def fit_sheet(ws, row_ranges, line_height=None):
    """シートの指定した行の高さを、表示される行数に合わせて広げる。

    Args:
        ws: 対象のワークシート。
        row_ranges (list[tuple[int, int]]): 対象の行の範囲。
        line_height (float | None): 1行の高さ。None なら行ごとに文字の大きさから求める。

    Returns:
        list[str]: 行の高さを変えた行の説明。
    """
    changed = []
    spans, multi_row = _single_row_spans(ws)
    letter = openpyxl.utils.get_column_letter
    default_height = ws.sheet_format.defaultRowHeight or 15.0

    for first, last in row_ranges:
        for row in range(first, min(last, ws.max_row) + 1):
            lines_needed = 1
            font_size = 0
            for cell in ws[row]:
                if not isinstance(cell.value, str) or (row, cell.column) in multi_row:
                    continue
                # セル（結合範囲）の幅から、1行に入る表示幅を求める
                low, high = spans.get((row, cell.column), (cell.column, cell.column))
                capacity = (
                    _merged_width(ws, letter(low), letter(high)) * CHARS_PER_WIDTH
                )
                text = cell.value.replace("\r\n", "\n").replace("\r", "\n")
                lines = _count_lines(text, capacity)
                if lines < 2:
                    continue
                # 折り返しをオンにする（既存の配置の他の設定は引き継ぐ）
                align = cell.alignment
                if not align.wrap_text:
                    cell.alignment = Alignment(
                        horizontal=align.horizontal,
                        vertical=align.vertical,
                        wrap_text=True,
                        indent=align.indent,
                        shrink_to_fit=align.shrink_to_fit,
                        text_rotation=align.text_rotation,
                    )
                lines_needed = max(lines_needed, lines)
                font_size = max(font_size, cell.font.sz or 11)

            if lines_needed < 2:
                continue
            # 行の高さを「1行の高さ×行数」に広げる（狭めはしない）
            current = ws.row_dimensions[row].height
            base = line_height or max(
                default_height, round(font_size * LINE_HEIGHT_PER_FONT_SIZE, 2)
            )
            height = base * lines_needed
            if current is None or height > current:
                ws.row_dimensions[row].height = height
                changed.append(f"{row}行目: {current} -> {height}（{lines_needed}行）")
    return changed


def main():
    ensure_utf8_stdio()
    parser = argparse.ArgumentParser(description="文字が隠れないように行の高さを広げる")
    parser.add_argument("input")
    parser.add_argument("output")
    parser.add_argument("--sheet", action="append", default=[])
    parser.add_argument("--body-sheets", action="store_true")
    parser.add_argument("--rows", action="append", required=True)
    parser.add_argument("--line-height", type=float)
    args = parser.parse_args()
    if not args.sheet and not args.body_sheets:
        parser.error("--sheet か --body-sheets のどちらかを指定してください")

    wb = openpyxl.load_workbook(args.input)
    sheets = list(args.sheet)
    if args.body_sheets:
        sheets += [
            s for s in wb.sheetnames if s not in COMMON_SHEETS and s not in sheets
        ]
    missing = [s for s in sheets if s not in wb.sheetnames]
    if missing:
        print(f"NG: シートがありません: {missing}")
        sys.exit(1)

    row_ranges = [_parse_rows(r) for r in args.rows]
    total = 0
    report = []
    for name in sheets:
        changed = fit_sheet(wb[name], row_ranges, args.line_height)
        total += len(changed)
        report += [f"[{name}] {line}" for line in changed]
    save_with_shapes(wb, args.input, args.output)

    print(f"OK: 行の高さを広げた行 {total} 件")
    for line in report:
        print(f"  - {line}")


if __name__ == "__main__":
    main()
