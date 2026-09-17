# -*- coding: utf-8 -*-
"""
外部インタフェース設計書(JSON)を、電文の定義ファイル(JSON)から組み立てるスクリプト。

このテンプレートは他の様式と違い、座標の計算を間違えやすい。
- 『【レコード名】』シートが3ブロック構成で、1ブロックのデータ行は7行しかない。
- 項目数が7を超えると行の挿入が要り、挿入すると下のブロックとデータ構成サンプルの行がずれる。
- 選択欄はチェックボックスではなくセル文字列(□/☑)であり、書き換える位置が決まっている。

これらを毎回マッピングJSONへ手で起こすと間違えるため、本スクリプトへ寄せる。

使い方:
    python build_external_if_json.py <テンプレートxlsx> <電文定義json> <出力xlsx>

電文定義JSONの形式:
{
  "msg_id": "EV0101R",                         # 電文ID。ファイル名・U8に入る
  "msg_name": "評価依頼受付（回答評価）要求",   # 電文名。表紙・目次・変更履歴に入る
  "io": "入力",                                 # 入力 / 出力
  "peer": "アンケート実施サブシステム",          # 相手先
  "tx": "EV0101／評価依頼受付（回答評価）",      # 入出力取引ID/名称
  "purpose": "…",                              # 目的・概要
  "condition": "…",                            # 作成条件
  "cycle": "随時",                              # 日次/週次/月次/年次/随時/その他
  "cycle_detail": "回答が確定した都度",
  "transfer": "Lambda Invoke（同期）",          # 授受方式
  "charset": "UTF-8",
  "encrypted": "無し",                          # 無し / 有り
  "remarks": ["なし。"],                        # 特記事項。1要素=1行(E28から連番セルへ)。
                                                #   意味のまとまり(句点等)で呼び出し側が分割する。
                                                #   最大4行(E28:E31の罫線内)。文字列1個でも可（後方互換）。
  "header": {"pj": "…", "system": "…", "author": "…", "date": 46279,
             "company": "…", "dept": "…"},      # date はExcelのシリアル値
  "records": [                                  # 先頭がルートレコード。最大3件
    {"name": "評価依頼", "rec_id": "-", "rec_type": "単一",
     "items": [
       {"name": "スタッフ識別子", "id": "staffNo", "type": "半角英数字",
        "size": "7", "required": "○", "multi": "1",
        "note": "", "sample": "\\"1000123\\""}
     ]}
  ]
}

記入の決まりは `記入要領_外部インタフェース設計書_JSON.md` を参照。特に次の2点を守る。
- 構成イメージはレコードが2件以上のときだけ書く。書く場所はレコードを定義する『3.1.〜』シート。
- 該当しない欄は空欄にせず「-」を書く。まだ決まっていない欄は「T.B.D（後ほど記載）」と書き分ける。

例(空テンプレートに、電文定義ファイルEV0101R.jsonの内容を反映する場合):
    python build_external_if_json.py "外部インタフェース設計書_(電文ID)_(電文名)_(JSON).xlsx" EV0101R.json 出力.xlsx
"""

import argparse
import json
import os
import shutil
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent))
from apply_mapping import apply_mapping  # noqa: E402
from insert_rows import insert_rows_keep_layout  # noqa: E402
from xlsx_common import save_with_shapes, ensure_utf8_stdio  # noqa: E402

TEMPLATE_RECORD_SHEET = "【レコード名】"
SPEC_SHEET = "1. 外部インタフェース仕様"
STRUCT_SHEET = "2. レコード構成"

# 『【レコード名】』のブロック（レコード名行, レコードID行, データ先頭行）と1ブロックの行数
BLOCKS = [(7, 8, 10), (18, 19, 20), (28, 29, 30)]
BLOCK_ROWS = 7

# 『1. 外部インタフェース仕様』シートの特記事項欄。E28を先頭に、罫線がE31までしか無いため最大4行
REMARKS_MAX_LINES = 4
SAMPLE_LABEL_ROW = 39  # データ構成サンプル／データ構成イメージの見出し行
IMAGE_COL = "Y"  # データ構成イメージ欄の左端の列

STRUCT_FIRST_ROW = 9
# 『2. レコード構成』の列（結合セルの左上）
STRUCT_COLS = {
    "name": "B",
    "rec_type": "G",
    "identify": "J",
    "length": "Q",
    "repeat": "S",
    "unit": "V",
    "sort_key": "AD",
    "sort_order": "AH",
}
# 『【レコード名】』の列
ITEM_COLS = {
    "no": "A",
    "name": "B",
    "id": "G",
    "domain": "L",
    "required": "Q",
    "type": "R",
    "size": "V",
    "multi": "X",
    "default": "Z",
    "format": "AB",
    "note": "AE",
}
# 選択欄（□を☑へ書き換える位置）
IO_CELL = {"入力": "E7", "出力": "I7"}
ENCRYPT_CELL = {"無し": "E22", "有り": "E23"}
CYCLE_CELL = {
    "日次": ("E25", "H25"),
    "週次": ("R25", "U25"),
    "月次": ("E26", "H26"),
    "年次": ("R26", "U26"),
    "随時": ("E27", "H27"),
    "その他": ("R27", "U27"),
}

# 該当しない欄の表し方。記入要領がJSON電文の長さを「-」としており、様式全体でこれに揃える。
# 「未定」は T.B.D と書き分ける。
NA = "-"


def layout(records):
    """項目数から、行を広げたあとの各ブロックの位置と、全体のずれ幅を求める。"""
    needs = [max(0, len(r["items"]) - BLOCK_ROWS) for r in records]
    blocks, offset = [], 0
    for i in range(len(records)):
        name_row, id_row, data_row = BLOCKS[i]
        blocks.append((name_row + offset, id_row + offset, data_row + offset))
        offset += needs[i]
    return blocks, needs, offset


def record_tree(records):
    """レコード間の入れ子関係。(深さ, 表示文字列) の並びで返す。"""
    out = [(0, records[0]["name"])]
    for rec in records[1:]:
        out.append(
            (1, "└ %s（%s：%s）" % (rec["name"], rec["rec_id"], rec["rec_type"]))
        )
    return out


def json_sample(records):
    """実データの例。配列は1要素だけ書く。各項目の sample を使う。"""

    def value(it):
        if it.get("sample"):
            return it["sample"]
        if it["type"] == "真偽値":
            return "true"
        if it["type"] in ("符号無数値", "符号付数値"):
            return "0"
        return '"..."'

    child = {r["rec_id"]: r for r in records[1:]}
    root = records[0]
    lines = ["{"]
    for n, it in enumerate(root["items"]):
        tail = "," if n < len(root["items"]) - 1 else ""
        if it["id"] in child:
            c = child[it["id"]]
            lines.append('  "%s": [' % it["id"])
            lines.append("    {")
            for k, ci in enumerate(c["items"]):
                lines.append(
                    '      "%s": %s%s'
                    % (ci["id"], value(ci), "," if k < len(c["items"]) - 1 else "")
                )
            lines.append("    }")
            lines.append("  ]%s" % tail)
        else:
            lines.append('  "%s": %s%s' % (it["id"], value(it), tail))
    lines.append("}")
    return lines


def prepare(template, out_path, sheet_name, records):
    """シート名を変え、ブロックを広げ、使わないブロックを空にする。"""
    from insert_rows import clear_rows

    wb = openpyxl.load_workbook(template)
    wb[TEMPLATE_RECORD_SHEET].title = sheet_name
    ws = wb[sheet_name]

    _, needs, total = layout(records)
    # 下のブロックから広げる。先に上を広げると下の行位置がずれるため
    for i in range(len(records) - 1, -1, -1):
        if needs[i] > 0:
            data_row = BLOCKS[i][2]
            insert_rows_keep_layout(ws, data_row + BLOCK_ROWS, needs[i], data_row)

    if len(records) < len(BLOCKS):
        start = BLOCKS[len(records)][0] + total
        end = BLOCKS[-1][2] + BLOCK_ROWS - 1 + total
        clear_rows(ws, range(start, end + 1))

    save_with_shapes(wb, template, out_path)


def build_mapping(d, sheet_name):
    records = d["records"]
    blocks, _, total = layout(records)
    h = d["header"]

    spec = {
        IO_CELL[d["io"]]: "☑" + d["io"],
        "U7": d["peer"],
        "E8": d["tx"],
        "U8": d["msg_id"],
        "E9": d["purpose"],
        "E14": d["condition"],
        "E21": d["transfer"],
        ENCRYPT_CELL[d["encrypted"]]: "☑" + d["encrypted"],
        "E24": d["charset"],
    }
    cyc_cell, cyc_detail = CYCLE_CELL[d["cycle"]]
    spec[cyc_cell] = "☑" + d["cycle"]
    spec[cyc_detail] = d["cycle_detail"]

    # 特記事項(E28)は、テンプレートの罫線がE28〜E31の4行分しか無い。
    # 文字列(従来形式)ならE28の1行、配列(1要素=1行)ならE28から連番セルへ展開する。
    # 意味のまとまり(句点等)での分割は呼び出し側(電文定義JSON)の責務とし、
    # ここでは機械的な文字列処理(句点split等)を行わない。
    remarks = d["remarks"]
    if isinstance(remarks, str):
        remarks = [remarks]
    if len(remarks) > REMARKS_MAX_LINES:
        raise ValueError(
            "remarksは最大%d行までです（テンプレートのE28:E%dの罫線内）。"
            "現在%d行あります。電文定義JSON側で行数を減らしてください。"
            % (REMARKS_MAX_LINES, 27 + REMARKS_MAX_LINES, len(remarks))
        )
    for i, line in enumerate(remarks):
        spec["E%d" % (28 + i)] = line

    # レコード構成はルートレコード1行だけ。該当しない欄は「-」で表す
    root = records[0]
    struct_vals = {
        "name": root["name"],
        "rec_type": root["rec_type"],
        "identify": root.get("identify") or NA,
        # JSON電文では長さ（Byte）を使わない（記入要領）
        "length": NA,
        "repeat": root.get("repeat") or NA,
        "unit": root.get("unit") or NA,
        "sort_key": root.get("sort_key") or NA,
        "sort_order": root.get("sort_order") or NA,
    }
    struct = {
        "%s%d" % (STRUCT_COLS[k], STRUCT_FIRST_ROW): v for k, v in struct_vals.items()
    }

    cells = {"A5": sheet_name}
    for i, rec in enumerate(records):
        name_row, id_row, data_row = blocks[i]
        cells["D%d" % name_row] = rec["name"]
        cells["D%d" % id_row] = rec["rec_id"]
        for j, it in enumerate(rec["items"]):
            r = data_row + j
            vals = {
                "no": j + 1,
                "name": it["name"],
                "id": it["id"],
                # ドメイン定義書が無い場合は項目名をそのまま使う
                "domain": it.get("domain") or it["name"],
                "required": it["required"],
                "type": it["type"],
                "size": it.get("size", ""),
                "multi": it.get("multi", "1"),
            }
            for k, v in vals.items():
                if v != "":
                    cells["%s%d" % (ITEM_COLS[k], r)] = v
            if it.get("note"):
                cells["%s%d" % (ITEM_COLS["note"], r)] = it["note"]

    # データ構成サンプルは必ず書く。1行につき1セル（セル内改行は表示されないため）
    first = SAMPLE_LABEL_ROW + total + 1
    for k, line in enumerate(json_sample(records)):
        cells["A%d" % (first + k)] = line
    # 構成イメージはレコードが2件以上のときだけ。階層は列をずらして表す
    if len(records) > 1:
        for k, (depth, text) in enumerate(record_tree(records)):
            cells["%s%d" % (chr(ord(IMAGE_COL) + depth), first + k)] = text

    return {
        "sheets": {
            "変更履歴": {
                "E1": h["pj"],
                "E2": h["system"],
                "E3": d["msg_name"],
                "A8": 1,
                "B8": "第1.0版",
                "D8": h["date"],
                "G8": "新規",
                "J8": "-",
                "Q8": "新規作成",
                "AF8": h["author"],
            },
            "表紙": {
                "J12": h["pj"],
                "J16": d["msg_name"],
                "J28": h["company"],
                "J30": h["dept"],
            },
            "目次": {"B7": "1. " + d["msg_name"]},
            SPEC_SHEET: spec,
            STRUCT_SHEET: struct,
            sheet_name: cells,
        }
    }


def ungray_struct_row(path, out_path):
    """レコード構成で記入した行のグレーアウト（記入不要の表現）を解除する。"""
    from openpyxl.styles import PatternFill
    from openpyxl.utils import column_index_from_string

    wb = openpyxl.load_workbook(path)
    ws = wb[STRUCT_SHEET]
    for c in range(column_index_from_string("B"), column_index_from_string("AI") + 1):
        ws.cell(row=STRUCT_FIRST_ROW, column=c).fill = PatternFill(fill_type=None)
    save_with_shapes(wb, path, out_path)


# テンプレート側の仕様上、罫線設定範囲(bbox)の外にあるが正当な書き込み先のセル。
# 表紙のJ28(会社名)・J30(部門名)、目次のB7(電文名見出し)は、他のテンプレートと違い
# bboxの外側に配置された静的テキストセルであり、SKILL.mdの「ローカル実行の手順」にも
# 同様の記載がある(表紙のJ12・J16・J28・J30は数式ではなく別途セル指定が必要)。
# これ以外の座標をここに追加してはいけない -- 追加するとJSON定義や座標計算の誤りを
# 見逃すようになる(allow_outside_borderを全体へ一律適用していたことによる過検出漏れ)。
OUTSIDE_BORDER_EXCEPTIONS = {
    "表紙": {"J28", "J30"},
    "目次": {"B7"},
}


def _split_mapping_by_bbox_exception(mapping):
    """mappingを、bboxの外への書き込みを許可してよいセルとそれ以外に分ける。"""
    strict_sheets = {}
    loose_sheets = {}
    for sheet_name, cell_map in mapping.get("sheets", {}).items():
        exceptions = OUTSIDE_BORDER_EXCEPTIONS.get(sheet_name, set())
        strict_cells = {k: v for k, v in cell_map.items() if k not in exceptions}
        loose_cells = {k: v for k, v in cell_map.items() if k in exceptions}
        if strict_cells:
            strict_sheets[sheet_name] = strict_cells
        if loose_cells:
            loose_sheets[sheet_name] = loose_cells
    return {"sheets": strict_sheets}, {"sheets": loose_sheets}


def build(template, definition_path, output):
    d = json.load(open(definition_path, encoding="utf-8"))
    records = d["records"]
    if len(records) > len(BLOCKS):
        raise ValueError(
            "1シートに置けるレコードは%d件までです: %d件" % (len(BLOCKS), len(records))
        )

    sheet_name = "3.1. %s" % records[0]["name"]
    work = output + ".work.xlsx"
    prepare(template, work, sheet_name, records)

    mapping = build_mapping(d, sheet_name)
    strict_mapping, loose_mapping = _split_mapping_by_bbox_exception(mapping)

    # 1回目: OUTSIDE_BORDER_EXCEPTIONS以外は罫線設定範囲を厳格に検証する
    # (JSON定義や座標計算の誤りで範囲外へ書こうとした場合はここでエラーになる)。
    tmp1 = output + ".tmp1.xlsx"
    report = apply_mapping(work, strict_mapping, tmp1, allow_outside_border=False)
    if report["errors"]:
        for e in report["errors"]:
            print("ERROR:", e)
        raise SystemExit(1)

    # 2回目: 既知の正当な例外セルのみ、罫線設定範囲の外への書き込みを許可する。
    tmp = output + ".tmp.xlsx"
    report2 = apply_mapping(tmp1, loose_mapping, tmp, allow_outside_border=True)
    if report2["errors"]:
        for e in report2["errors"]:
            print("ERROR:", e)
        raise SystemExit(1)
    report["warnings"].extend(report2["warnings"])
    report["written"].extend(report2["written"])

    ungray_struct_row(tmp, output)
    for p in (work, tmp1, tmp):
        if os.path.exists(p):
            os.remove(p)
    return output


def main():
    ensure_utf8_stdio()
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    ap.add_argument("template")
    ap.add_argument("definition")
    ap.add_argument("output")
    a = ap.parse_args()
    print("created:", build(a.template, a.definition, a.output))


if __name__ == "__main__":
    main()
