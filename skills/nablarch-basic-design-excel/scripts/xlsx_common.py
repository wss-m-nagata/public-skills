# -*- coding: utf-8 -*-
"""
apply_mapping.py / verify_layout.py / extract_structure.py / duplicate_sheet.py /
extend_table_rows.py で共有するExcel構造解析・操作の共通処理。

解析系の関数(color_str, resolve_list_values, get_validations_map,
get_bordered_bbox, snapshot_sheet_structure)はファイルを書き換えない。
操作系の関数(copy_data_validations, extend_validation_range)は
呼び出し側が明示的にworkbookを保存するまでは書き換えを確定させない。
"""
import os
import sys
import tempfile
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.utils.cell import coordinate_from_string
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.cell_range import CellRange, MultiCellRange

sys.path.insert(0, str(Path(__file__).parent))


def ensure_utf8_stdio():
    """
    標準出力・標準エラー出力の文字エンコーディングをUTF-8に固定する。
    Windows環境ではコンソールの既定コードページ(cp932等)により、print()が日本語を含む
    メッセージでUnicodeEncodeErrorを起こしたり、文字化けしたりすることがあるため、
    各スクリプトのmain()の先頭で呼び出す。reconfigure()が無い環境や失敗する環境では
    何もしない(Mac/Linuxでは既定でUTF-8のため通常は変化なし)。
    """
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            try:
                stream.reconfigure(encoding="utf-8")
            except Exception:
                pass


def color_str(color):
    """openpyxlのColorオブジェクトを比較・表示可能な文字列に変換する"""
    if color is None:
        return None
    ctype = getattr(color, "type", None)
    if ctype == "rgb":
        rgb = color.rgb
        if isinstance(rgb, str) and rgb not in (None, "00000000"):
            return rgb
        return None
    if ctype == "theme":
        return f"theme{color.theme}(tint={round(color.tint, 3)})"
    if ctype == "indexed":
        return f"indexed{color.indexed}"
    return None


def resolve_list_values(wb, formula1, pending=None):
    """
    データ入力規則(type='list')のformula1を実際の選択肢のリストに解決する。

    pending: {シート名: {セル座標: 新しい値}} を渡すと、選択肢の参照先セルがその中に
    含まれている場合はファイル上の現在値ではなくpendingの値を優先する。
    apply_mapping.pyで、プルダウンの選択肢自体(「データ」シートの型一覧等)と、それに
    依存する値を同じ呼び出し内で書き換える場合に、更新前の古い選択肢と比較してしまう
    誤検知を防ぐために使う。
    """
    if formula1 is None:
        return None
    f = formula1.strip()
    if f.startswith('"') and f.endswith('"'):
        return [v for v in f[1:-1].split(",") if v != ""]

    ref = f.lstrip("=")
    defined = wb.defined_names.get(ref) if hasattr(wb.defined_names, "get") else None
    if defined is not None:
        ref = defined.attr_text

    if "!" in ref:
        sheet_part, cell_part = ref.rsplit("!", 1)
        sheet_name = sheet_part.strip("'")
        if sheet_name in wb.sheetnames:
            try:
                pending_for_sheet = (pending or {}).get(sheet_name, {})
                cells = wb[sheet_name][cell_part.replace("$", "")]
                values = []
                for row in cells:
                    row_iter = row if isinstance(row, tuple) else (row,)
                    for c in row_iter:
                        value = pending_for_sheet.get(c.coordinate, c.value)
                        if value is not None:
                            values.append(str(value).strip())
                return values
            except Exception:
                return None
    return None


def get_validations_map(wb, ws, pending=None):
    """
    入力規則(プルダウン)を { セル座標: [選択肢...] } の形にフラット化する。
    pendingの意味はresolve_list_values()を参照。
    """
    result = {}
    for dv in ws.data_validations.dataValidation:
        if dv.type != "list":
            continue
        values = resolve_list_values(wb, dv.formula1, pending)
        if values is None:
            continue
        for rng in dv.sqref.ranges:
            for row in ws.iter_rows(min_row=rng.min_row, max_row=rng.max_row,
                                     min_col=rng.min_col, max_col=rng.max_col):
                for cell in row:
                    result[cell.coordinate] = values
    return result


def get_merge_anchor(ws, coordinate):
    """
    指定セルが結合セルの一部である場合、書き込み可能な左上のアンカーセル座標を返す。
    結合セルでなければそのまま座標を返す。
    """
    col_str, row = coordinate_from_string(coordinate)
    col = column_index_from_string(col_str)
    for merged_range in ws.merged_cells.ranges:
        if (merged_range.min_row <= row <= merged_range.max_row and
                merged_range.min_col <= col <= merged_range.max_col):
            return f"{get_column_letter(merged_range.min_col)}{merged_range.min_row}"
    return coordinate


def get_bordered_bbox(ws, max_scan_row=300, max_scan_col=60):
    """
    罫線が設定されている最大の行・列を返す。
    「テンプレートとして罫線が引かれ、デザインされた領域」の外形とみなし、
    その領域を超える書き込みを検知するための基準値として使う。
    """
    max_row, max_col = 0, 0
    for r in range(1, min(ws.max_row, max_scan_row) + 1):
        for c in range(1, min(ws.max_column, max_scan_col) + 1):
            cell = ws.cell(row=r, column=c)
            b = cell.border
            if b and any([b.top and b.top.style, b.bottom and b.bottom.style,
                          b.left and b.left.style, b.right and b.right.style]):
                max_row = max(max_row, r)
                max_col = max(max_col, c)
    return max_row, max_col


def snapshot_sheet_structure(wb, ws):
    """
    verify_layout.py での前後比較用に、1シート分の構造情報をまとめて取得する。
    値(value)は含めない -- あくまで構造(レイアウト)のスナップショット。
    """
    max_row = min(ws.max_row, 300)
    max_col = min(ws.max_column, 60)

    fills = {}
    fonts = {}
    borders = set()
    formulas = {}
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            cell = ws.cell(row=r, column=c)
            fill = cell.fill
            if fill and fill.fgColor and fill.patternType is not None:
                col = color_str(fill.fgColor)
                if col:
                    fills[cell.coordinate] = col
            font = cell.font
            if font and font.color:
                col = color_str(font.color)
                if col:
                    fonts[cell.coordinate] = col
            b = cell.border
            if b and any([b.top and b.top.style, b.bottom and b.bottom.style,
                          b.left and b.left.style, b.right and b.right.style]):
                borders.add(cell.coordinate)
            if cell.data_type == "f":
                formulas[cell.coordinate] = cell.value

    return {
        "dimensions": (ws.max_row, ws.max_column),
        "merged_cells": sorted(str(m) for m in ws.merged_cells.ranges),
        "row_heights": {r: d.height for r, d in ws.row_dimensions.items() if d.height is not None},
        "col_widths": {k: d.width for k, d in ws.column_dimensions.items() if d.width is not None},
        "fills": fills,
        "fonts": fonts,
        "borders": borders,
        "formulas": formulas,
        "validations": [
            (str(dv.sqref), dv.type, dv.formula1)
            for dv in ws.data_validations.dataValidation
        ],
        "sheet_protection": bool(ws.protection.sheet),
    }


def copy_data_validations(src_ws, dst_ws):
    """
    openpyxlのcopy_worksheet()は入力規則(データ検証/プルダウン)を複製しない
    (実機検証で確認済み: 2件->0件)。
    そのため、シート複製後にこの関数で明示的にコピーする必要がある。
    座標(sqref)は元シートと同じレイアウトである前提でそのまま引き継ぐ。
    """
    count = 0
    for dv in src_ws.data_validations.dataValidation:
        new_dv = DataValidation(
            type=dv.type,
            formula1=dv.formula1,
            formula2=dv.formula2,
            operator=dv.operator,
            allow_blank=dv.allow_blank,
            showDropDown=dv.showDropDown,
            showErrorMessage=dv.showErrorMessage,
            showInputMessage=dv.showInputMessage,
            errorTitle=dv.errorTitle,
            error=dv.error,
            promptTitle=dv.promptTitle,
            prompt=dv.prompt,
        )
        new_dv.sqref = MultiCellRange(str(dv.sqref))
        dst_ws.add_data_validation(new_dv)
        count += 1
    return count


def extend_validation_ranges(ws, source_row, new_rows, min_col, max_col):
    """
    source_row行に設定されている入力規則(プルダウン)の適用範囲を、
    new_rows(追加した行番号のリスト)まで拡張する。
    既存のDataValidationオブジェクトのsqrefに範囲を追加する形で行う
    (新しいDataValidationは作らない。選択肢の定義はsource_row側のものをそのまま使う)。
    """
    extended = []
    for dv in ws.data_validations.dataValidation:
        for rng in list(dv.sqref.ranges):
            if (rng.min_row <= source_row <= rng.max_row
                    and rng.min_col >= min_col and rng.max_col <= max_col):
                for r in new_rows:
                    new_range = CellRange(
                        min_col=rng.min_col, min_row=r,
                        max_col=rng.max_col, max_row=r,
                    )
                    dv.sqref.ranges.add(new_range)
                    extended.append((str(dv.formula1), str(new_range)))
    return extended


def find_next_populated_row(ws, after_row, min_col, max_col, max_scan_row=1000):
    """
    after_row より後で、指定列範囲のいずれかのセルに値が入っている最初の行番号を返す。
    見つからなければNoneを返す。

    extend_table_rows.py・insert_row_block.pyで、複製先が後続のセクション見出し等の
    既存内容を上書きしてしまうことを防ぐための安全確認に使う。
    """
    max_row = min(ws.max_row, max_scan_row)
    for r in range(after_row + 1, max_row + 1):
        for c in range(min_col, max_col + 1):
            if ws.cell(row=r, column=c).value is not None:
                return r
    return None


def get_single_row_merges(ws, row, min_col, max_col):
    """指定した行の中に収まっている(複数行にまたがらない)結合セル範囲を取得する"""
    result = []
    for m in ws.merged_cells.ranges:
        if m.min_row == row and m.max_row == row and m.min_col >= min_col and m.max_col <= max_col:
            result.append((m.min_col, m.max_col))
    return result


def save_with_shapes(wb, original_path, output_path, text_replacements=None):
    """
    openpyxlのworkbookを保存し、その直後にrestore_shapes.pyで図形(テキストボックス等)を
    元ファイルから復元する。

    【背景】openpyxlは読み込み→保存の過程で、セルコメント以外の図形(DrawingML形式の
    テキストボックス等)を破棄してしまう(実機検証で確認済み)。Nablarch標準テンプレートの
    表紙シートには「関係者外秘」等を表示するテキストボックスが配置されているため、
    このまま保存すると気付かれないままレイアウトが劣化する。

    apply_mapping.py / duplicate_sheet.py / extend_table_rows.py は、
    wb.save(output_path) を直接呼ぶ代わりに必ずこの関数を使うこと。

    original_path: 今回の処理の入力に使った元xlsxファイルのパス(図形が残っているファイル)
    text_replacements: 表紙のテキストボックス内プレースホルダーを実値に置き換えたい場合に指定する。
        例: {"プロジェクト名": "サンプルプロジェクト", "サブシステム名": "顧客管理システム"}
        省略した場合、プレースホルダーは元のまま(`[プロジェクト名]`等)復元される。
    """
    import restore_shapes

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        tmp_path = tmp.name
    try:
        wb.save(tmp_path)
        report = restore_shapes.restore_shapes(original_path, tmp_path, output_path, text_replacements)
    finally:
        try:
            os.remove(tmp_path)
        except OSError:
            pass
    return report
