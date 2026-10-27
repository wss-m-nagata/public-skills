"""単体テスト仕様書に、試験データのシートを足す。

build_unit_test_spec.py で作った単体テスト仕様書（.xlsx）の末尾に、試験データの定義（JSON）から
「試験データ」のシートを作って足す。リクエスト単体・取引単体のケースは、「テストデータシート名」と
「データNo」の列で、このシートのデータNoを指す。

データは、テーブルごとに、カラム名とレコード（実際に入れる値）で書く。
概要（「主な値」など）で済ませず、仕様書を見て、試験できることが確かめられる値を書く。
カラムが多すぎるテーブルは、試験に関係しないカラムを割愛してよい（割愛したことを note に書く）。

シートの形（左から右へ、階段状に書く）:
    A列：シート名（A1 だけ）
    B列：「共通データ」「個別データ」の見出し
    C列：使わない
    D列：個別データのデータNo（D01 など）の見出しだけを書く
    E列：共通データも個別データも、ここからテーブルを書く
    A〜D列は見出し用で、幅は3にする（見出しは右の空きセルへはみ出して表示する）。
    データをE列から書くのは、Shift＋矢印キーで、データのまとまりごとに選びやすくするため

定義（JSON）の形:
    {
      "sheet_name": "試験データ",                          # 省略時は「試験データ」
      "common": {                                          # 試験では作らず、今ある前提のデータ（省略可）
        "note": "値が変わったら、ここも直す",              # 見出しの下の説明（省略可）
        "tables": [TABLE, …]
      },
      "datasets": [                                        # 試験で作るデータのまとまり
        {
          "no": "D01",
          "name": "注文の一式",
          "cases": "注文API_リクエスト単体 1-1-1",          # このデータを使うケース
          "notes": ["作る処理：…", "後片付け：…"],         # 表の上に書く説明（省略可）
          "tables": [TABLE, …]                             # 入れる順に書く
        }
      ]
    }

    TABLE = {
      "table": "shop.order_header（注文ヘッダー）",         # テーブル名（和名を付けてよい）
      "columns": ["order_id", "customer_id", …],           # カラム名
      "rows": [[900000001, 900000001, …], …],              # レコード。null は None、真偽値は true / false
      "note": "割愛したカラムなどの補足"                    # 省略可
    }

使い方:
    python scripts/add_test_data_sheet.py <仕様書.xlsx> <試験データの定義.json> <出力.xlsx>
"""

import argparse
import json
import unicodedata
from copy import copy

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side

# 列の幅の上限（Excel の幅の単位）。これより長い値は折り返す
MAX_COLUMN_WIDTH = 46
# テーブルを書き始める列（共通データも個別データも、E列）
COMMON_FIRST_COL = 5
DATASET_FIRST_COL = 5
# 個別データの見出し（D01：〜）を書く列
DATASET_HEADING_COL = 4
# A〜D列（見出し用）の幅
HEADING_COLUMN_WIDTH = 3
# 罫線
THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _text(value):
    """セルに書く文字列にする。null は NULL、真偽値は true / false、数値は文字列にする。"""
    if value is None:
        return "NULL"
    if isinstance(value, bool):
        return "true" if value else "false"
    return str(value)


def _display_width(text):
    """値の、いちばん長い行の表示幅を返す（全角は2、半角は1と数える）。"""
    return max(
        sum(
            2 if unicodedata.east_asian_width(ch) in ("F", "W", "A") else 1
            for ch in line
        )
        for line in _text(text).split("\n")
    )


def _header_style(wb):
    """リクエスト単体などのシートの見出し（ケースNo のセル）から、見出しの書式を借りる。"""
    for ws in wb.worksheets:
        for row in ws.iter_rows(max_row=20):
            for cell in row:
                if cell.value == "ケースNo":
                    return copy(cell.font), copy(cell.fill)
    return Font(bold=True), None


def _body_font(wb):
    """見出しの次の行の、本文のフォントを借りる（フォント名とサイズをそろえる）。"""
    for ws in wb.worksheets:
        for row in ws.iter_rows(max_row=20):
            for cell in row:
                if cell.value == "ケースNo":
                    return copy(ws.cell(row=cell.row + 1, column=cell.column + 9).font)
    return Font()


def _write_table(
    ws, row, table, header_font, header_fill, body_font, title_font, first_col
):
    """テーブル名・補足・カラム名・レコードを書き、次に書ける行を返す。"""
    columns = table["columns"]

    # テーブル名と補足
    ws.cell(row=row, column=first_col, value=table["table"]).font = title_font
    row += 1
    if table.get("note"):
        ws.cell(row=row, column=first_col, value=f"※{table['note']}").font = copy(
            body_font
        )
        row += 1

    # カラム名
    for index, name in enumerate(columns):
        cell = ws.cell(row=row, column=first_col + index, value=name)
        cell.font = header_font
        if header_fill is not None:
            cell.fill = header_fill
        cell.border = BORDER
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    row += 1

    # レコード
    for values in table["rows"]:
        if len(values) != len(columns):
            raise ValueError(f"列の数がカラム名と合いません: {table['table']} {values}")
        lines = 1
        for index, value in enumerate(values):
            col = first_col + index
            cell = ws.cell(row=row, column=col, value=_text(value))
            cell.font = copy(body_font)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            width = ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width
            wrapped = -(-_display_width(value) // max(int(width) - 1, 1))
            lines = max(lines, wrapped, _text(value).count("\n") + 1)
        ws.row_dimensions[row].height = 15 * lines + 3
        row += 1
    return row + 1


def add_sheet(spec_path, definition_path, output_path):
    """仕様書に試験データのシートを足して保存する。"""
    definition = json.load(open(definition_path, encoding="utf-8"))
    wb = openpyxl.load_workbook(spec_path)

    # 同じ名前のシートがあれば作り直す
    name = definition.get("sheet_name", "試験データ")
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)

    header_font, header_fill = _header_style(wb)
    body_font = _body_font(wb)
    title_font = Font(name=body_font.name, size=body_font.size, bold=True)

    # 列の幅。A〜D列は見出し用に細くし、E列からデータの値のいちばん長いものに合わせる
    widths = {}
    common = definition.get("common")
    for first_col, tables in (
        (COMMON_FIRST_COL, common["tables"] if common else []),
        (DATASET_FIRST_COL, [t for d in definition["datasets"] for t in d["tables"]]),
    ):
        for table in tables:
            for values in [table["columns"], *table["rows"]]:
                for index, value in enumerate(values):
                    col = first_col + index
                    widths[col] = max(widths.get(col, 0), _display_width(value) + 2)
    for col, width in widths.items():
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = min(
            max(width, 12), MAX_COLUMN_WIDTH
        )
    for letter in ("A", "B", "C", "D"):
        ws.column_dimensions[letter].width = HEADING_COLUMN_WIDTH

    # シート名（A1）
    row = 1
    ws.cell(row=row, column=1, value=name).font = title_font
    row += 2

    # 共通データ（B列に見出し、E列からテーブル）
    if common:
        ws.cell(row=row, column=2, value="共通データ").font = title_font
        row += 1
        if common.get("note"):
            ws.cell(row=row, column=COMMON_FIRST_COL, value=common["note"]).font = copy(
                body_font
            )
            row += 1
        row += 1
        for table in common["tables"]:
            row = _write_table(
                ws,
                row,
                table,
                header_font,
                header_fill,
                body_font,
                title_font,
                COMMON_FIRST_COL,
            )

    # 個別データ（B列に見出し、D列にデータNo、E列から使うケースとテーブル）
    ws.cell(row=row, column=2, value="個別データ").font = title_font
    row += 2
    for dataset in definition["datasets"]:
        ws.cell(
            row=row,
            column=DATASET_HEADING_COL,
            value=f"{dataset['no']}：{dataset['name']}",
        ).font = title_font
        row += 1
        if dataset.get("cases"):
            ws.cell(
                row=row,
                column=DATASET_FIRST_COL,
                value=f"使うケース：{dataset['cases']}",
            ).font = copy(body_font)
            row += 1
        for note in dataset.get("notes", []):
            ws.cell(row=row, column=DATASET_FIRST_COL, value=note).font = copy(
                body_font
            )
            row += 1
        row += 1
        for table in dataset["tables"]:
            row = _write_table(
                ws,
                row,
                table,
                header_font,
                header_fill,
                body_font,
                title_font,
                DATASET_FIRST_COL,
            )

    wb.save(output_path)


def main():
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("spec")
    parser.add_argument("definition")
    parser.add_argument("output")
    args = parser.parse_args()
    add_sheet(args.spec, args.definition, args.output)


if __name__ == "__main__":
    main()
