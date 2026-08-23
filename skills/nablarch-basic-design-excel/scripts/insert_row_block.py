# -*- coding: utf-8 -*-
"""
「見出し行＋表ヘッダ＋データ行＋補足行」等、複数行から成るブロックを、後続の内容
（後ろの見出しや別の表）を下に押し出しながら複数個複製するスクリプト。
Excelの「行の挿入」に相当する操作を、結合セル・入力規則・行高を壊さずに行う。

extend_table_rows.pyとの違い:
- extend_table_rows.pyは、単純な1行の繰り返し(一覧表の最終行)を下方向に追加する。
  後ろに何も無い(その表がシートの末尾、または罫線の無い余白が続く)ことを前提にしている。
- insert_row_block.pyは、複数行から成るブロック(例:「2.5.1 [テーブル名]」という見出し＋
  取得項目表＋取得条件、という一塊)を、後ろに別のセクション（2.6・2.7等）が既に存在する
  状態でも、その後続部分を丸ごと下にシフトしながら複製する。

【注意】安全のため、このスクリプトは「一覧系テンプレートの本体シート」「1インスタンス1シート型
テンプレートの詳細シート（『2. 取引ID（取引名）』等）」でのみ使うこと。
**表紙・変更履歴・目次シートには絶対に使わない**（これらのシートには他シートを参照する数式や
図形があり、行シフトで壊れる可能性が高い）。

【注意】**対象列範囲の外側（多くの場合は左側）に章見出し（例:「2.5. 入力データ定義」）がある場合、
その列を対象列範囲に含め忘れると、見出しだけがシフト・複製されずズレる**
（例:「2.4 入出力一覧」の表本体はD列から始まるが、章見出し自体はC列にあるテンプレートがある）。
このスクリプトは対象列範囲の外に値が入ったセルがあれば実行前に検出してエラーで止める。
見出しの実際の列は、事前に対象シートを直接セル単位で確認するか、`export_sheet_preview.py`で
PDF化して目視するのが確実（記入要領・構造メモに列番号までは書かれていない場合がある）。

使い方:
    python insert_row_block.py <元xlsxパス> <シート名> <ブロック開始行> <ブロック終了行> \\
        <追加ブロック数> <開始列> <終了列> <出力xlsxパス> [--allow-partial-columns]

例（「2.5.1 [テーブル名]」ブロックが38〜47行目の10行で、これと同じ形のブロックを
   さらに2つ追加する場合。元の1ブロック＋新規2ブロック＝合計3ブロックになる）:
    python insert_row_block.py システム機能設計書.xlsx "2. E10201（評価処理）" 38 47 2 B AH 出力.xlsx

新しく増えた行には、複製元ブロックの罫線・背景色・結合セル・行高・入力規則(プルダウン)は
複製されるが、値(セルの文字列)は複製されない(値の反映は`apply_mapping.py`で行う)。
"""
import sys
import argparse
from pathlib import Path
from copy import copy

import openpyxl
from openpyxl.comments import Comment
from openpyxl.utils import column_index_from_string
from openpyxl.worksheet.cell_range import CellRange

sys.path.insert(0, str(Path(__file__).parent))
from xlsx_common import save_with_shapes, ensure_utf8_stdio


def _copy_comment(src_cell, dst_cell):
    """
    Excelのセルコメント(元テンプレートの記入ガイド等)を複製する。
    フォント・罫線等と違い、openpyxlのCellオブジェクトの標準コピーには含まれないため、
    明示的にコピーしないと欠落する。

    **ブロックの複製(num_copies)には使わない**。同じ説明コメントが複数のブロックに
    重複して残ると、後から人が手直しする際にどれが本来の構造か分かりにくくなるため、
    複製先には意図的にコメントを持たせない。行のシフト(既存内容を下にずらすだけで、
    複製ではない)についてのみ使う。
    """
    if src_cell.comment is not None:
        dst_cell.comment = Comment(src_cell.comment.text, src_cell.comment.author or "")


def _copy_cell_style(src_cell, dst_cell):
    """
    ブロック複製(num_copies)用のスタイルコピー。コメントは意図的にコピーしない
    (理由は_copy_comment()のdocstringを参照)。
    """
    dst_cell.font = copy(src_cell.font)
    dst_cell.fill = copy(src_cell.fill)
    dst_cell.border = copy(src_cell.border)
    dst_cell.alignment = copy(src_cell.alignment)
    dst_cell.number_format = src_cell.number_format
    dst_cell.protection = copy(src_cell.protection)


def _copy_cell_style_and_value(src_cell, dst_cell):
    """行のシフト(既存内容の移動)用。コメントもここでは移動させる(複製ではないため)。"""
    dst_cell.value = src_cell.value
    _copy_cell_style(src_cell, dst_cell)
    _copy_comment(src_cell, dst_cell)


def _find_content_outside_column_range(ws, row_start, row_end, min_col, max_col, max_scan_col=60):
    """
    指定した行範囲について、対象列範囲(min_col〜max_col)の外側に値の入ったセルがないか調べる。
    見出し等が対象列範囲に含まれていないと、シフト・複製の対象から漏れて位置がズレるため、
    insert_row_block.py実行前の安全確認に使う。
    """
    found = []
    last_col = min(ws.max_column, max_scan_col)
    for r in range(row_start, row_end + 1):
        for c in range(1, last_col + 1):
            if min_col <= c <= max_col:
                continue
            v = ws.cell(row=r, column=c).value
            if v is not None:
                found.append((ws.cell(row=r, column=c).coordinate, v))
    return found


def insert_row_block(xlsx_path, sheet_name, block_start_row, block_end_row, num_copies,
                      min_col_letter, max_col_letter, output_path, allow_partial_columns=False):
    if block_end_row < block_start_row:
        raise ValueError(f"ブロック終了行({block_end_row})がブロック開始行({block_start_row})より前です。")
    if num_copies <= 0:
        raise ValueError(f"追加ブロック数({num_copies})は1以上を指定してください。")

    wb = openpyxl.load_workbook(xlsx_path)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"シートが存在しません: {sheet_name}")
    if sheet_name in ("表紙", "変更履歴", "目次"):
        raise ValueError(
            f"『{sheet_name}』シートへの使用は禁止されています。"
            f"insert_row_block.pyは本体シート専用です（表紙・変更履歴・目次には絶対に使わないこと）。"
        )
    ws = wb[sheet_name]

    min_col = column_index_from_string(min_col_letter)
    max_col = column_index_from_string(max_col_letter)
    block_height = block_end_row - block_start_row + 1
    shift = block_height * num_copies
    old_max_row = ws.max_row

    if not allow_partial_columns:
        outside = _find_content_outside_column_range(ws, block_start_row, old_max_row, min_col, max_col)
        if outside:
            sample = outside[:10]
            more = f"（他{len(outside) - 10}件）" if len(outside) > 10 else ""
            raise ValueError(
                f"指定した対象列範囲（{min_col_letter}〜{max_col_letter}列）の外側に、"
                f"値が入っているセルがあります: {sample}{more}\n"
                f"章見出し等がこの列範囲に含まれていないと、シフト・複製の対象から漏れて位置がズレます。"
                f"上記セルの列も対象列範囲に含め直してください。"
                f"（意図的に一部の列だけを対象にしたい場合のみ --allow-partial-columns を指定）"
            )

    # --- 1. 結合セルを「ブロックより後ろ(シフト対象)」「ブロック内(複製対象)」に分類する ---
    shift_merges = []
    block_merges = []
    for m in list(ws.merged_cells.ranges):
        if m.min_row > block_end_row:
            shift_merges.append((m.min_row, m.max_row, m.min_col, m.max_col))
        elif m.min_row >= block_start_row and m.max_row <= block_end_row:
            block_merges.append((m.min_row, m.max_row, m.min_col, m.max_col))
        elif m.max_row > block_end_row and m.min_row <= block_end_row:
            raise ValueError(
                f"結合セル{m.coord}がブロックの終了行({block_end_row})をまたいでいます。"
                f"ブロックの開始行・終了行を、結合セルの境界に合わせて指定し直してください。"
            )
        # ブロックより前の結合セルはそのまま(何もしない)

    # --- 2. データ検証(プルダウン)を同様に分類する ---
    shift_validations = []  # (formula1, type, その他プロパティ, [シフト対象range...])
    block_validations = []  # (formula1, type, その他プロパティ, [ブロック内range...])
    for dv in ws.data_validations.dataValidation:
        shift_ranges = []
        block_ranges = []
        for rng in list(dv.sqref.ranges):
            if rng.min_row > block_end_row:
                shift_ranges.append((rng.min_row, rng.max_row, rng.min_col, rng.max_col))
            elif rng.min_row >= block_start_row and rng.max_row <= block_end_row:
                block_ranges.append((rng.min_row, rng.max_row, rng.min_col, rng.max_col))
            elif rng.max_row > block_end_row and rng.min_row <= block_end_row:
                raise ValueError(
                    f"入力規則の範囲{rng.coord}がブロックの終了行({block_end_row})をまたいでいます。"
                )
        if shift_ranges:
            shift_validations.append((dv, shift_ranges))
        if block_ranges:
            block_validations.append((dv, block_ranges))

    # --- 3. シフト対象の結合だけを一旦解除する ---
    # ブロック内(block_merges)は複製元として値・書式を「読む」だけで、複製元自身の位置は
    # 一切書き換えないため解除する必要が無い。解除してしまうと、複製先(delta>0)にしか
    # 結合を作り直さない後段の処理により、複製元自身の結合セルが失われたままになる。
    for (r1, r2, c1, c2) in shift_merges:
        ws.unmerge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)

    # --- 4. ブロックより後ろの内容を下方向にシフトする(下から上に処理して上書き事故を防ぐ) ---
    for r in range(old_max_row, block_end_row, -1):
        new_r = r + shift
        for c in range(min_col, max_col + 1):
            src_cell = ws.cell(row=r, column=c)
            dst_cell = ws.cell(row=new_r, column=c)
            _copy_cell_style_and_value(src_cell, dst_cell)
            src_cell.value = None
            src_cell.comment = None
        old_height = ws.row_dimensions[r].height
        if old_height is not None:
            ws.row_dimensions[new_r].height = old_height

    # --- 5. シフトした結合セル・入力規則を、+shift した位置に再作成する ---
    for (r1, r2, c1, c2) in shift_merges:
        ws.merge_cells(start_row=r1 + shift, start_column=c1, end_row=r2 + shift, end_column=c2)

    for dv, ranges in shift_validations:
        for (r1, r2, c1, c2) in ranges:
            dv.sqref.ranges.add(CellRange(min_col=c1, min_row=r1 + shift, max_col=c2, max_row=r2 + shift))

    # --- 6. ブロックを num_copies 回、複製する(値はコピーしない。構造のみ) ---
    for i in range(num_copies):
        delta = block_height * (i + 1)
        for local_r in range(block_start_row, block_end_row + 1):
            new_r = local_r + delta
            for c in range(min_col, max_col + 1):
                src_cell = ws.cell(row=local_r, column=c)
                dst_cell = ws.cell(row=new_r, column=c)
                _copy_cell_style(src_cell, dst_cell)
            h = ws.row_dimensions[local_r].height
            if h is not None:
                ws.row_dimensions[new_r].height = h
        for (r1, r2, c1, c2) in block_merges:
            ws.merge_cells(start_row=r1 + delta, start_column=c1, end_row=r2 + delta, end_column=c2)
        for dv, ranges in block_validations:
            for (r1, r2, c1, c2) in ranges:
                dv.sqref.ranges.add(CellRange(min_col=c1, min_row=r1 + delta, max_col=c2, max_row=r2 + delta))

    save_with_shapes(wb, xlsx_path, output_path)

    return {
        "sheet": sheet_name,
        "original_block": (block_start_row, block_end_row),
        "num_copies": num_copies,
        "new_blocks": [
            (block_end_row + block_height * i + 1, block_end_row + block_height * (i + 1))
            for i in range(num_copies)
        ],
        "shifted_rows_from": block_end_row + 1,
        "shift_amount": shift,
    }


def main():
    ensure_utf8_stdio()
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("xlsx_path")
    parser.add_argument("sheet_name")
    parser.add_argument("block_start_row", type=int)
    parser.add_argument("block_end_row", type=int)
    parser.add_argument("num_copies", type=int)
    parser.add_argument("min_col")
    parser.add_argument("max_col")
    parser.add_argument("output_path")
    parser.add_argument("--allow-partial-columns", action="store_true")
    args = parser.parse_args()

    result = insert_row_block(
        args.xlsx_path, args.sheet_name, args.block_start_row, args.block_end_row,
        args.num_copies, args.min_col, args.max_col, args.output_path,
        args.allow_partial_columns,
    )
    print(f"=== ブロック挿入結果: {args.output_path} ===")
    print(f"元ブロック: {result['original_block'][0]}〜{result['original_block'][1]}行目（変更なし）")
    print(f"追加したブロック: {result['new_blocks']}")
    print(f"{result['shifted_rows_from']}行目以降の既存内容を{result['shift_amount']}行下にシフトしました。")
    print("\n※新しいブロックにはまだ値が入っていません。apply_mapping.pyで値を反映してください。")
    print("※目次シートに章番号等の参照がある場合は、手動での更新も忘れずに。")


if __name__ == "__main__":
    main()
