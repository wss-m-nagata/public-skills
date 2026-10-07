# -*- coding: utf-8 -*-
"""
単体テスト仕様書(Webサービス／バッチ)を、テスト項目の定義ファイル(JSON)から組み立てるスクリプト。

単体テスト仕様書のテスト仕様のシートには、テスト観点カタログの観点が、あらかじめ1行ずつ並んでいる。
仕様書を作るときは、まずその機能にある分類(大項目〜)を整理し、無い分類の行はシートから消す。
残した分類の観点ごとに「テストケースを書く」か「対象外理由を書く」かのどちらかを埋めることで、
観点の抜けが無いこと（全量をテストしたこと）を示す。対象外の行は灰色にする。
1つの観点に複数のテストケースがある場合は、その観点の下に行を足す必要があり、
足すとケースNo(数式)・罫線・入力規則・印刷範囲がずれるため、手でマッピングJSONを組むと間違える。
そのため、定義ファイルから一括で組み立てる。

使い方:
    python build_unit_test_spec.py <テンプレートxlsx> <テスト項目定義json> <出力xlsx>
        [--allow-unaddressed] [--cases-json <ケース一覧の出力先json>]

  --allow-unaddressed : テストケースも対象外理由も無い観点が残っていても、エラーにしない(書きかけの確認用)。
  --cases-json        : 組み立てた結果の全テストケース(ケースNo・分類・観点・テスト内容など)をJSONに書き出す。

テンプレートは assets/templates/フォーマット/ の
「単体テスト仕様書(Webサービス)_(取引ID)_(取引名).xlsx」または「単体テスト仕様書(バッチ)_(取引ID)_(取引名).xlsx」。
表紙・変更履歴も、定義の "document" から書く。

テスト項目定義JSONの形式:
{
  "document": {                              # 表紙・変更履歴。必須
    "PJ名": "サンプルプロジェクト", "システム名": "サンプルシステム", "サブシステム名": "顧客管理システム",
    "会社名": "…", "部門名": "…",
    "変更履歴": [                             # 1要素=1行(8行目から)。表紙の版数は最後の要素の版数になる
      {"版数": "第1.0版", "変更日": "2026-10-06", "区分": "新規",
       "変更箇所": ["-"], "変更内容": ["新規作成"],   # 1要素=1行(セル内改行で並べる)
       "担当者": "TIS"}
    ]
  },
  "header": {                                # テスト仕様のシートの見出し(全シート共通)。省略可
    "機能名": "顧客管理", "取引名": "顧客登録"
                                             # 作成者・作成日・更新者・更新日は、変更履歴から自動で表示される(書かない)
  },
  "sheets": [
    {
      "template": "クラス単体",               # クラス単体 / リクエスト単体 / 取引単体(バッチのみ)
      "name": "ClientService",                # シート名は「<name>_クラス単体」「<name>_リクエスト単体」になる。
                                              #   取引単体では使わない(シート名は「取引単体」のまま)
      "header": {"テストターゲット名": "ClientService"},   # シートごとの見出し。リクエスト単体では「リクエストID」も書く
      "分類": [                               # この機能にある分類。必須。ここに書いていない分類の行は、シートから消す
        ["バリデーション", "単項目バリデーション", "必須バリデーション"],
        ["DB操作", "登録"]                    # 大項目・中項目・小項目・詳細を上から順に並べたもの(前方一致)
      ],
      "対象外": [                             # 「分類」に書いた分類のうち、まとめて対象外にするもの。省略可
        {"分類": ["DB操作", "登録"], "理由": "リクエスト単体テストで確認するため対象外"}
                                              # 分類は大項目・中項目・小項目・詳細を上から順に並べたもの(前方一致)
      ],
      "観点": [                               # 観点ごとの記入。No は空のテンプレートでのケースNoの前2つ(「詳細-観点」)
        {"No": "1-1", "cases": [
          {"前提": "顧客名を送らない",         # 前提条件。テスト内容の先頭に「【前提】…」として書く。省略可
           "テスト内容": "顧客登録のリクエストを送る。",
           "想定結果": "ステータスコード400が返る。",
           "対応設計書": "システム機能設計書", "対応箇所": "入力データ定義",
           "テストコード": "test_client_service.py::test_顧客名が無いとエラー"}
                                              # ほかに書ける列はテンプレートの見出しと同じ名前(テストデータシート名・データNo・実施者 など)
        ]},
        {"No": "3-5", "対象外理由": "サロゲートペア文字を許容しない設計のため対象外"}
      ],
      "追加観点": [                           # テンプレートに無い、プロジェクト独自の観点。省略可
        {"after": "12-1",                     # この観点(空のテンプレートでの「詳細-観点」)の直後に足す
         "詳細": "事業部",                     # 新しい詳細を始めるときに書く(大項目・中項目・小項目も同様)。省略すると直前の詳細の続き
         "観点": "閲覧範囲の外の事業部を指定した場合、エラーとなること。",
         "cases": [{"テスト内容": "…", "想定結果": "…"}]}
      ]
    }
  ]
}

シートの組み合わせは、単体テスト標準(機能テストの分類と処理方式の組み合わせ)に合わせる。合わない場合はエラーにする。
- リクエスト単体(画面・バッチ単位)は、例外なく作る。
- クラス単体(モジュール単位)は、Webサービスだけ。ロジックが複雑なものに限って作る。バッチでは作らない。
- 取引単体(ユースケース単位)は、バッチで、複数のリクエストから成る取引のときだけ作る。

記入の決まりは `記入要領_単体テスト仕様書.md` を参照する。
"""

import argparse
import json
import re
import sys
from copy import copy
from pathlib import Path

import openpyxl
import datetime

from openpyxl.styles import Border, Color, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula

sys.path.insert(0, str(Path(__file__).parent))
from fit_change_history_rows import (
    _column_width,
    _display_width,
    fit_rows,
)  # noqa: E402
from xlsx_common import (  # noqa: E402
    copy_data_validations,
    ensure_utf8_stdio,
    save_with_shapes,
)

# テンプレートの種類と、テンプレートでのシート名・作成後のシート名の付け方
TEMPLATE_SHEETS = {
    "クラス単体": "【テストターゲット名】_クラス単体",
    "リクエスト単体": "【リクエストID】_リクエスト単体",
    "取引単体": "取引単体",
}

# 見出し行とデータの先頭行(全シート共通)
HEADER_ROW = 10
FIRST_ROW = 11

# 見出し(1〜9行目)に「<項目名>：<値>」の形で書く項目
HEADER_LABELS = (
    "機能名",
    "取引名",
    "テストターゲット名",
    "リクエストID",
)

# 見出しのうち、変更履歴シートから自動で表示する項目(定義には書かない)
HEADER_FROM_HISTORY = ("作成者", "作成日", "更新者", "更新日")

# 対象外の行の塗り(灰色)。Nablarchのテンプレート原本・記入例で、対象外の行に付けている色と同じ
EXCLUDED_FILL = PatternFill(
    fill_type="solid", fgColor=Color(theme=0, tint=-0.499984740745262)
)

# 変更履歴シートの列(結合セルの左上)と、1行目のデータ行・最終行
HISTORY_COLUMNS = {
    "No": "A",
    "版数": "B",
    "変更日": "D",
    "区分": "G",
    "変更箇所": "J",
    "変更内容": "Q",
    "担当者": "AF",
}
HISTORY_FIRST_ROW = 8
HISTORY_LAST_ROW = 33

# 分類の列(上の階層から順)。取引単体は「区分・確認観点1」が分類にあたる
CATEGORY_HEADERS = ("大項目", "中項目", "小項目", "詳細", "区分", "確認観点1")

# 1行あたりの高さ(9ptのフォント)と、列幅1あたりに入る表示幅(半角1・全角2)の目安。
# 2026-10-06、このテンプレート(9pt)をExcelでPDFにして測ったところ、列幅1あたり表示幅1程度が入った。
# 単語の区切りで早めに折り返すことがあるため、少し控えめに見積もる
LINE_HEIGHT = 12.0
DISPLAY_PER_WIDTH = 0.95


class SheetLayout:
    """テスト仕様のシートの列の配置。見出し行の名前から求める。"""

    def __init__(self, ws):
        # 見出しの名前 → 列番号
        self.columns = {c.value: c.column for c in ws[HEADER_ROW] if c.value}
        is_transaction = "確認観点1" in self.columns
        # ケースNoの1つ目を数える列(詳細／確認観点1)
        self.level1 = self.columns["確認観点1" if is_transaction else "詳細"]
        # ケースNoの2つ目を数える列(観点／確認観点2)
        self.level2 = self.columns["確認観点2" if is_transaction else "観点"]
        # ケースNoの3つ目を数える列(テスト内容)
        self.content = self.columns["テスト内容"]
        # 対象外理由の列
        self.reason = self.columns["対象外理由"]
        # 実施結果の列(入力規則 OK/NG を付ける列)
        self.result = self.columns["実施結果"]
        # 分類の列(左から順)
        self.categories = [
            self.columns[h] for h in CATEGORY_HEADERS if h in self.columns
        ]
        # 表の右端の列
        self.last_col = max(self.columns.values())
        # 取引単体かどうか(ケースNoの数え方が少し違う)
        self.is_transaction = is_transaction


def _last_template_row(ws):
    """表の最終行(ケースNoの数式がある最後の行)を返す。"""
    last = FIRST_ROW
    for row in range(FIRST_ROW, ws.max_row + 1):
        if isinstance(ws.cell(row=row, column=1).value, ArrayFormula):
            last = row
    return last


def _catalog_ids(ws, layout, last_row):
    """空のテンプレートの各行に「詳細-観点」の番号を振る(ケースNoの前2つと同じ)。"""
    ids = {}
    level1 = level2 = 0
    for row in range(FIRST_ROW, last_row + 1):
        if ws.cell(row=row, column=layout.level1).value not in (None, ""):
            level1 += 1
            level2 = 0
        if ws.cell(row=row, column=layout.level2).value not in (None, ""):
            level2 += 1
        ids[row] = f"{level1}-{max(level2, 1)}"
    return ids


def _is_category(value):
    """分類の列の値が、分類そのものかどうか。

    大項目の列には、分類の補足の注記(「※ファイルアップロードを含む」など)が下の行に入っていることがある。
    「※」で始まる値は注記であり、新しい分類としては扱わない。
    """
    return value not in (None, "") and not str(value).startswith("※")


def _catalog_paths(ws, layout, last_row):
    """各行の分類(大項目〜詳細を、空欄は上の行から引き継いで並べたもの)を返す。"""
    paths = {}
    current = [""] * len(layout.categories)
    for row in range(FIRST_ROW, last_row + 1):
        for i, col in enumerate(layout.categories):
            value = ws.cell(row=row, column=col).value
            if not _is_category(value):
                continue
            current[i] = str(value)
            # 上の階層が変わったら、下の階層は引き継がない
            for j in range(i + 1, len(current)):
                current[j] = ""
        paths[row] = list(current)
    return paths


def _row_formula(template_text, row):
    """ケースNoの数式(テンプレートの11行目のもの)を、指定した行のものに書き換える。

    $E$11 のような絶対参照の 11 はそのまま残し、E11 のような相対参照の 11 だけを行番号にする。
    """
    return re.sub(
        r"(?<![$A-Z])([A-Z])11(?!\d)", lambda m: f"{m.group(1)}{row}", template_text
    )


def _case_values(case, layout):
    """テストケース1件を、列番号 → 値 の辞書にする。"""
    values = {}
    for key, value in case.items():
        if key == "前提":
            continue
        if key not in layout.columns:
            raise ValueError(
                f"テンプレートに無い列です: {key}（書ける列: {', '.join(layout.columns)}）"
            )
        values[layout.columns[key]] = value
    # 前提条件は、テスト内容の先頭に「【前提】…」として書く(記入要領のとおり)
    if case.get("前提"):
        body = case.get("テスト内容", "")
        values[layout.content] = (
            f"【前提】{case['前提']}\n{body}" if body else f"【前提】{case['前提']}"
        )
    return values


def _plan_rows(sheet_def, ws, layout, last_row, allow_unaddressed):
    """出力する行の並びを組み立てる。

    戻り値は、出力する行ごとの (書式の元にするテンプレートの行, 値の辞書, 行の種類) のリスト。
    行の種類は catalog(テンプレートの観点の行) / extra(同じ観点の2件目以降のケース) / added(追加観点の行)。
    """
    ids = _catalog_ids(ws, layout, last_row)
    paths = _catalog_paths(ws, layout, last_row)
    rows_by_id = {}
    for row, cid in ids.items():
        rows_by_id.setdefault(cid, row)

    # その機能にある分類の行だけを残す(無い分類の行は、シートから消す)
    kept_prefixes = sheet_def.get("分類")
    if not kept_prefixes:
        raise ValueError("「分類」に、この機能にある分類(大項目〜)を書いてください")
    for prefix in kept_prefixes:
        if not any(path[: len(prefix)] == prefix for path in paths.values()):
            raise ValueError(f"分類が、テンプレートのどの観点にも当たりません: {prefix}")
    kept_rows = {
        row
        for row, path in paths.items()
        if any(path[: len(prefix)] == prefix for prefix in kept_prefixes)
    }
    kept_ids = {ids[row] for row in kept_rows}

    # 観点ごとの記入
    entries = {}
    for item in sheet_def.get("観点", []):
        cid = item["No"]
        if cid not in rows_by_id:
            raise ValueError(f"テンプレートに無い観点のNoです: {cid}")
        if cid not in kept_ids:
            raise ValueError(f"観点 {cid} は、「分類」に書いていない分類の観点です")
        if cid in entries:
            raise ValueError(f"同じ観点のNoが2回書かれています: {cid}")
        if bool(item.get("cases")) == bool(item.get("対象外理由")):
            raise ValueError(
                f"観点 {cid} には、cases と 対象外理由 のどちらか一方だけを書いてください"
            )
        entries[cid] = item

    # 分類の単位での対象外(個別の記入が優先)
    for rule in sheet_def.get("対象外", []):
        prefix = rule["分類"]
        matched = False
        for row, path in paths.items():
            if row in kept_rows and path[: len(prefix)] == prefix:
                matched = True
                entries.setdefault(
                    ids[row], {"No": ids[row], "対象外理由": rule["理由"]}
                )
        if not matched:
            raise ValueError(
                f"対象外の分類が、「分類」に書いた分類のどの観点にも当たりません: {prefix}"
            )

    # 追加観点(after の観点の直後に足す)
    added = {}
    for item in sheet_def.get("追加観点", []):
        if item["after"] not in kept_ids:
            raise ValueError(
                f"追加観点の after が、「分類」に書いた分類の観点のNoではありません: {item['after']}"
            )
        if not item.get("cases") and not item.get("対象外理由"):
            raise ValueError(
                f"追加観点「{item.get('観点')}」に、cases も 対象外理由 もありません"
            )
        added.setdefault(item["after"], []).append(item)

    # 記入の無い観点
    missing = [
        cid for cid in dict.fromkeys(ids[r] for r in sorted(kept_rows)) if cid not in entries
    ]
    if missing and not allow_unaddressed:
        raise ValueError(
            "テストケースも対象外理由も無い観点があります（全量を埋めてください）: "
            + ", ".join(missing)
        )

    # テンプレートの行を、同じ観点(番号)のまとまりに分ける
    groups = []
    for row in sorted(kept_rows):
        if groups and groups[-1][0] == ids[row]:
            groups[-1][1].append(row)
        else:
            groups.append((ids[row], [row]))

    plan = []
    for cid, rows in groups:
        entry = entries.get(cid) or {}
        cases = entry.get("cases") or [{}]

        # 1. 観点の行(1件目のケース、または対象外理由)
        values = {}
        if entry.get("対象外理由"):
            values[layout.reason] = entry["対象外理由"]
        else:
            values.update(_case_values(cases[0], layout))
        plan.append((rows[0], values, "catalog"))

        # 2. 同じ番号の続きの行(テンプレートで分類の注記を下の行に持つもの)は、そのまま写す
        for row in rows[1:]:
            plan.append((row, {}, "catalog"))

        # 3. 2件目以降のケース
        for case in cases[1:]:
            plan.append((rows[0], _case_values(case, layout), "extra"))

        # 4. この観点の後ろに足す追加観点
        for item in added.get(cid, []):
            base = {}
            for header in CATEGORY_HEADERS + ("観点", "確認観点2"):
                if header in item and header in layout.columns:
                    base[layout.columns[header]] = item[header]
            if item.get("対象外理由"):
                base[layout.reason] = item["対象外理由"]
                plan.append((rows[0], base, "added"))
                continue
            first = dict(base)
            first.update(_case_values(item["cases"][0], layout))
            plan.append((rows[0], first, "added"))
            for case in item["cases"][1:]:
                plan.append((rows[0], _case_values(case, layout), "extra"))

    return plan, missing


def _estimate_height(ws, values, base_height):
    """セルの表示幅と列の幅から、行の高さの目安を求める(折り返して全部見えるように)。

    列の幅は _column_width で読む(ws.column_dimensions[列] と書くと、まとめて設定された列
    (J:K など)に新しい設定が作られ、列の幅が変わってしまうため)。
    """
    lines = 1
    for col, value in values.items():
        if value in (None, ""):
            continue
        per_line = max(1, int(_column_width(ws, col) * DISPLAY_PER_WIDTH))
        count = 0
        for text_line in str(value).split("\n"):
            count += max(1, -(-_display_width(text_line) // per_line))
        lines = max(lines, count)
    return max(base_height or 0, LINE_HEIGHT * lines + 4)


def _relabel_categories(ws, row, layout, path, previous_path):
    """分類の見出し(大項目〜詳細)を、残した行に合わせて書き直す。

    テンプレートでは、分類の名前は、そのまとまりの最初の行にだけ書いてある。
    最初の行を消すと名前も消えるため、分類が変わる行に名前と上の罫線を書き直し、
    同じ分類の続きの行からは名前と上の罫線を消す。
    """
    for i, col in enumerate(layout.categories):
        cell = ws.cell(row=row, column=col)
        # 分類の補足の注記(「※…」)は、そのまま残す
        if cell.value not in (None, "") and not _is_category(cell.value):
            continue
        border = cell.border
        changed = previous_path is None or path[: i + 1] != previous_path[: i + 1]
        if changed and path[i]:
            cell.value = path[i]
            # 続きの行のセルは、文字色が背景と同じ色(見えない色)になっているため、標準の色に戻す
            font = cell.font
            cell.font = Font(name=font.name, size=font.sz, bold=font.b, italic=font.i, vertAlign=font.vertAlign, underline=font.u, strike=font.strike)
            style = border.left.style or border.right.style or "thin"
            cell.border = Border(left=copy(border.left), right=copy(border.right), top=Side(style=style), bottom=copy(border.bottom))
        else:
            cell.value = None
            cell.border = Border(left=copy(border.left), right=copy(border.right), top=Side(), bottom=copy(border.bottom))


def _without_top(border, columns_without_top, col):
    """続きの行の分類・観点の列は、上の罫線を消す(1つの観点の中の行であることを示す)。"""
    if col not in columns_without_top:
        return copy(border)
    return Border(
        left=copy(border.left),
        right=copy(border.right),
        top=Side(),
        bottom=copy(border.bottom),
    )


def build_sheet(ws, sheet_def, allow_unaddressed):
    """テスト仕様のシート1枚を、定義のとおりに組み立てる。"""
    layout = SheetLayout(ws)
    last_row = _last_template_row(ws)
    template_formula = ws.cell(row=FIRST_ROW, column=1).value.text

    # 1. テンプレートの行の書式・高さ・値を控える
    snapshot = {}
    for row in range(FIRST_ROW, last_row + 1):
        snapshot[row] = {
            "height": ws.row_dimensions[row].height,
            "cells": {
                col: (
                    ws.cell(row=row, column=col).value,
                    ws.cell(row=row, column=col)._style,
                )
                for col in range(2, layout.last_col + 1)
            },
            "a_style": ws.cell(row=row, column=1)._style,
        }
    region_merges = [rng for rng in ws.merged_cells.ranges if rng.min_row >= FIRST_ROW]

    # 2. 出力する行の並びを決める
    paths = _catalog_paths(ws, layout, last_row)
    plan, missing = _plan_rows(sheet_def, ws, layout, last_row, allow_unaddressed)

    # 3. 表の部分を消す
    for rng in region_merges:
        ws.unmerge_cells(str(rng))
    clear_to = max(last_row, FIRST_ROW + len(plan) - 1)
    for row in range(FIRST_ROW, clear_to + 1):
        for col in range(1, layout.last_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.value = None
            cell.style = "Normal"
        ws.row_dimensions[row].height = None

    # 4. 行を書く
    columns_without_top = set(layout.categories) | {layout.level2}
    first_out_row = {}
    previous_path = None
    for offset, (src, values, kind) in enumerate(plan):
        out = FIRST_ROW + offset
        first_out_row.setdefault(src, out)
        src_cells = snapshot[src]["cells"]
        ws.cell(row=out, column=1)._style = copy(snapshot[src]["a_style"])
        for col in range(2, layout.last_col + 1):
            value, style = src_cells[col]
            cell = ws.cell(row=out, column=col)
            cell._style = copy(style)
            if kind == "catalog":
                cell.value = value
            elif kind == "extra" and col in columns_without_top:
                # 同じ観点の2件目以降：分類・観点は空にし、上の罫線を消す
                cell.border = _without_top(cell.border, columns_without_top, col)
            elif kind == "added" and col in columns_without_top:
                # 追加観点：観点は書き、分類は書いた階層だけ上の罫線を残す
                if col not in values:
                    cell.border = _without_top(cell.border, columns_without_top, col)
        if kind == "catalog":
            _relabel_categories(ws, out, layout, paths[src], previous_path)
            previous_path = paths[src]
        for col, value in values.items():
            ws.cell(row=out, column=col).value = value
        # ケースNoの数式
        ws.cell(row=out, column=1).value = ArrayFormula(
            f"A{out}", _row_formula(template_formula, out)
        )
        # 対象外の行は、観点から右を灰色にする(テストした行と見分けられるように)
        if values.get(layout.reason):
            for col in range(layout.level2, layout.last_col + 1):
                ws.cell(row=out, column=col).fill = copy(EXCLUDED_FILL)
        # 行の高さ
        shown = {
            col: ws.cell(row=out, column=col).value
            for col in range(2, layout.last_col + 1)
        }
        ws.row_dimensions[out].height = _estimate_height(
            ws, shown, snapshot[src]["height"] if kind == "catalog" else None
        )

    new_last = FIRST_ROW + len(plan) - 1

    # 表の最終行には、下の罫線を引く(分類の列は、まとまりの最後の行にしか下の罫線が無く、
    # その行を消した場合に表の下が開くため)
    for col in range(1, layout.last_col + 1):
        cell = ws.cell(row=new_last, column=col)
        border = cell.border
        style = border.left.style or border.right.style or "thin"
        cell.border = Border(
            left=copy(border.left), right=copy(border.right), top=copy(border.top), bottom=Side(style=style)
        )

    # 5. 分類の注記の結合セル(「※ファイルアップロードを含む」など)を、新しい行の位置で結合し直す
    for rng in region_merges:
        if rng.min_row not in first_out_row:
            continue
        start = first_out_row[rng.min_row]
        end = first_out_row.get(rng.max_row, start)
        ws.merge_cells(
            start_row=start,
            start_column=rng.min_col,
            end_row=end,
            end_column=rng.max_col,
        )

    # 6. 実施結果の入力規則(OK/NG)を、表の行全体に付け直す
    keep = [
        dv
        for dv in ws.data_validations.dataValidation
        if not any(r.min_row >= FIRST_ROW for r in dv.sqref.ranges)
    ]
    ws.data_validations.dataValidation = keep
    result_letter = openpyxl.utils.get_column_letter(layout.result)
    dv = DataValidation(type="list", formula1='"OK,NG"', allow_blank=True)
    dv.add(f"{result_letter}{FIRST_ROW}:{result_letter}{new_last}")
    ws.add_data_validation(dv)

    # 7. 対象外理由の列を表示する
    # テンプレートの原本では、この列はグループ化して非表示になっている。
    # 観点ごとに「テストした」か「対象外にした理由」のどちらかが見えないと、全量の確認ができないため表示する
    # (グループ化は残すので、Excelで畳むこともできる)。
    # まとめて設定された列に新しい設定を作らないよう、この列を含む既存の設定を探して直す
    for dim in ws.column_dimensions.values():
        low = dim.min or openpyxl.utils.column_index_from_string(dim.index)
        if low <= layout.reason <= (dim.max or low):
            dim.hidden = False

    # 8. 印刷範囲とオートフィルター
    last_letter = openpyxl.utils.get_column_letter(layout.last_col)
    ws.print_area = f"A1:{last_letter}{new_last}"
    if ws.auto_filter.ref:
        ws.auto_filter.ref = f"A{HEADER_ROW}:{last_letter}{HEADER_ROW}"

    return layout, new_last, missing


def _write_document(wb, document):
    """表紙と変更履歴シートを書く。

    PJ名・システム名・サブシステム名は変更履歴シートにだけ書く(ほかのシート・表紙は数式で表示する)。
    表紙の版数は、変更履歴の最後の行の版数にそろえる。
    """
    history_rows = document["変更履歴"]
    if not history_rows:
        raise ValueError("document の変更履歴に、1行以上書いてください")
    if len(history_rows) > HISTORY_LAST_ROW - HISTORY_FIRST_ROW + 1:
        raise ValueError(
            f"変更履歴は {HISTORY_LAST_ROW - HISTORY_FIRST_ROW + 1} 行までです"
        )

    # 1. 変更履歴シート
    history = wb["変更履歴"]
    history["E1"] = document["PJ名"]
    history["E2"] = document["システム名"]
    history["E3"] = document["サブシステム名"]
    for i, item in enumerate(history_rows):
        row = HISTORY_FIRST_ROW + i
        history[f"{HISTORY_COLUMNS['No']}{row}"] = i + 1
        history[f"{HISTORY_COLUMNS['版数']}{row}"] = item["版数"]
        # 変更日は日付型にする(表紙の最終更新日・見出しの作成日などの数式が計算できるように)
        history[f"{HISTORY_COLUMNS['変更日']}{row}"] = datetime.datetime.strptime(
            item["変更日"], "%Y-%m-%d"
        )
        history[f"{HISTORY_COLUMNS['区分']}{row}"] = item["区分"]
        # 変更箇所・変更内容は、1要素=1行のセル内改行で書く(記入要領_共通シート.md)
        history[f"{HISTORY_COLUMNS['変更箇所']}{row}"] = "\n".join(item["変更箇所"])
        history[f"{HISTORY_COLUMNS['変更内容']}{row}"] = "\n".join(item["変更内容"])
        history[f"{HISTORY_COLUMNS['担当者']}{row}"] = item["担当者"]
    fit_rows(history)

    # 2. 表紙(J23・J25・I21 は数式なので書かない)
    cover = wb["表紙"]
    cover["J12"] = document["PJ名"]
    cover["J16"] = document["サブシステム名"]
    cover["J19"] = history_rows[-1]["版数"]
    cover["J28"] = document["会社名"]
    cover["J30"] = document["部門名"]


def _fill_header(ws, header):
    """見出し(1〜9行目)の「<項目名>：」のセルに値を書く。"""
    for key, value in header.items():
        if key in HEADER_FROM_HISTORY:
            raise ValueError(
                f"{key} は変更履歴シートから自動で表示されるため、header には書かない（document の変更履歴に書く）"
            )
        if key not in HEADER_LABELS:
            raise ValueError(
                f"見出しに無い項目です: {key}（書ける項目: {', '.join(HEADER_LABELS)}）"
            )
        target = None
        for row in ws.iter_rows(min_row=1, max_row=HEADER_ROW - 1):
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith(f"{key}："):
                    target = cell
        if target is None:
            raise ValueError(f"シート『{ws.title}』に「{key}：」の欄がありません")
        target.value = f"{key}：{value}"


def compute_case_numbers(ws, layout, last_row):
    """ケースNo(Excelの数式と同じ数え方)と、各行の内容を返す。

    openpyxlで書いたファイルには数式の計算結果が入らないため、確認用に同じ計算をする。
    """
    cases = []
    level1 = level2 = level3 = 0
    current = [""] * len(layout.categories)
    point = ""
    for row in range(FIRST_ROW, last_row + 1):

        def val(col):
            return ws.cell(row=row, column=col).value

        for i, col in enumerate(layout.categories):
            if _is_category(val(col)):
                current[i] = str(val(col))
                for j in range(i + 1, len(current)):
                    current[j] = ""
        if val(layout.level1) not in (None, ""):
            level1 += 1
            level2 = 0
            level3 = 0
        if val(layout.level2) not in (None, ""):
            level2 += 1
            level3 = 0
            point = str(val(layout.level2))
        if val(layout.content) not in (None, ""):
            level3 += 1
        record = {
            "ケースNo": f"{level1}-{max(level2, 1)}-{max(level3, 1)}",
            "分類": [c for c in current if c and c != "-"],
            "観点": point,
        }
        for name, col in layout.columns.items():
            if name in ("ケースNo",) + CATEGORY_HEADERS + ("観点", "確認観点2"):
                continue
            if val(col) not in (None, ""):
                record[name] = val(col)
        cases.append(record)
    return cases


def build(template, definition_path, output, allow_unaddressed=False, cases_json=None):
    definition = json.loads(Path(definition_path).read_text(encoding="utf-8"))
    wb = openpyxl.load_workbook(template)

    # 1. 定義のシートを、テンプレートのシートから複製して作る
    # 単体テスト標準(機能テストの分類と処理方式の組み合わせ)に合わないシートを止める
    kinds = [d["template"] for d in definition["sheets"]]
    is_batch = TEMPLATE_SHEETS["取引単体"] in wb.sheetnames
    if is_batch and "クラス単体" in kinds:
        raise ValueError("バッチでは、モジュール単位(クラス単体)のテストは行わない（単体テスト標準）")
    if "取引単体" in kinds and kinds.count("リクエスト単体") < 2:
        raise ValueError(
            "取引単体(ユースケース単位)は、複数のリクエストから成る取引のときだけ作る（単体テスト標準）"
        )
    if "リクエスト単体" not in kinds:
        raise ValueError("リクエスト単体(画面・バッチ単位)のテストは、例外なく作る（単体テスト標準）")

    used_templates = set()
    built = []
    for sheet_def in definition["sheets"]:
        kind = sheet_def["template"]
        if kind not in TEMPLATE_SHEETS:
            raise ValueError(
                f"template は {', '.join(TEMPLATE_SHEETS)} のどれかです: {kind}"
            )
        source_name = TEMPLATE_SHEETS[kind]
        if source_name not in wb.sheetnames:
            raise ValueError(f"このテンプレートに『{source_name}』シートはありません")
        new_name = source_name if kind == "取引単体" else f"{sheet_def['name']}_{kind}"
        if len(new_name) > 31 or re.search(r"[\[\]:*?/\\]", new_name):
            raise ValueError(
                f"シート名に使えない名前です（31文字以内、[ ] : * ? / \\ は不可）: {new_name}"
            )
        if kind == "取引単体":
            ws = wb[source_name]
        else:
            source = wb[source_name]
            ws = wb.copy_worksheet(source)
            copy_data_validations(source, ws)
            ws.title = new_name
            ws.freeze_panes = source.freeze_panes
            if source.auto_filter.ref:
                ws.auto_filter.ref = source.auto_filter.ref
        used_templates.add(source_name)
        built.append((ws, sheet_def))

    # 2. 使わないテンプレートのシートを消し、作ったシートを並べる(クラス単体 → リクエスト単体 → 取引単体)
    for name in TEMPLATE_SHEETS.values():
        if name in wb.sheetnames and not (
            name == "取引単体" and name in used_templates
        ):
            del wb[name]
    order = {"クラス単体": 0, "リクエスト単体": 1, "取引単体": 2}
    fixed = [wb[n] for n in ("表紙", "変更履歴")]
    sheets = sorted(
        (ws for ws, _ in built),
        key=lambda s: order[[k for k in order if s.title.endswith(k)][0]],
    )
    wb._sheets = fixed + sheets

    # 3. シートごとに中身を組み立てる
    summary = []
    all_cases = {}
    for ws, sheet_def in built:
        header = dict(definition.get("header", {}))
        header.update(sheet_def.get("header", {}))
        _fill_header(ws, header)
        layout, last_row, missing = build_sheet(ws, sheet_def, allow_unaddressed)
        cases = compute_case_numbers(ws, layout, last_row)
        all_cases[ws.title] = cases
        written = [c for c in cases if c.get("テスト内容")]
        excluded = [c for c in cases if c.get("対象外理由")]
        summary.append((ws.title, len(written), len(excluded), missing))

    _write_document(wb, definition["document"])
    save_with_shapes(wb, template, output)

    if cases_json:
        Path(cases_json).write_text(
            json.dumps(all_cases, ensure_ascii=False, indent=2, default=str),
            encoding="utf-8",
        )
    return summary


def main():
    ensure_utf8_stdio()
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("template")
    parser.add_argument("definition")
    parser.add_argument("output")
    parser.add_argument("--allow-unaddressed", action="store_true")
    parser.add_argument("--cases-json")
    args = parser.parse_args()

    summary = build(
        args.template,
        args.definition,
        args.output,
        args.allow_unaddressed,
        args.cases_json,
    )
    for title, written, excluded, missing in summary:
        print(f"{title}: テストケース {written} 件 / 対象外 {excluded} 件")
        if missing:
            print(f"  記入の無い観点: {', '.join(missing)}", file=sys.stderr)


if __name__ == "__main__":
    main()
