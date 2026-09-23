# -*- coding: utf-8 -*-
"""
外部インタフェース設計書(JSON/CSV)テンプレートの『目次』シートを、
実際のレコード詳細シート名に機械的に同期するスクリプト。

【背景】
duplicate_sheet.py・rename_sheet.pyでレコード詳細シート（「3.1. [レコード名]」等）を
作成・改名しても、『目次』シート（C12, C13, ...）への反映は自動化されておらず、
両スクリプトの実行後メッセージでも「目次シートは自動更新されません。手動で追記してください」
と案内されるのみだった。この手動フォローが漏れると、目次だけがテンプレートの
プレースホルダー（「3.1. 【レコード名】」）のまま取り残される（2026-09-20に4ファイルで発生）。
このスクリプトはその同期を機械的に行い、手動での更新漏れを防ぐ。

【前提】
- 『目次』『2. レコード構成』『データ』の3シートが存在すること。
- レコード詳細シートは、シート順で「2. レコード構成」の直後から「データ」の直前までの
  連続した範囲にあること。duplicate_sheet.pyを`--before-sheet データ`で実行していれば
  自然にこの並びになる。rename_sheet.py単独でシート名だけを変えた場合も、
  シートの並び順自体は変わらないため問題ない。
- 目次のC12を1件目として、以降1行ずつ（C13, C14, ...）に対応させる。
  レコードが2件以上ある場合の目次側の行の書式（フォント等）はテンプレート側に
  あらかじめ用意されているため、このスクリプト側で罫線・書式のコピーは行わない
  （値の書き込みのみ）。想定を超える件数（既定8件）を検出した場合はエラーで止める。

使い方:
    python sync_toc_record_sheets.py <元xlsxパス> <出力xlsxパス> [--max-records N]

例:
    python sync_toc_record_sheets.py 外部インタフェース設計書_EV0109_CSV出力ダウンロード応答_(CSV).xlsx 出力.xlsx
"""

import sys
import argparse
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent))
from xlsx_common import save_with_shapes, ensure_utf8_stdio

TOC_SHEET = "目次"
RECORD_LIST_SHEET = "2. レコード構成"
DATA_SHEET = "データ"
FIRST_RECORD_ROW = 12
RECORD_COL = 3  # C列
DEFAULT_MAX_RECORDS = 8


def sync_toc_record_sheets(xlsx_path, output_path, max_records=DEFAULT_MAX_RECORDS):
    wb = openpyxl.load_workbook(xlsx_path)
    for required in (TOC_SHEET, RECORD_LIST_SHEET, DATA_SHEET):
        if required not in wb.sheetnames:
            raise ValueError(
                f"シートが見つかりません: {required}（現在のシート構成: {wb.sheetnames}）\n"
                f"このスクリプトは外部インタフェース設計書(JSON/CSV)テンプレート専用です。"
            )

    names = wb.sheetnames
    start = names.index(RECORD_LIST_SHEET) + 1
    end = names.index(DATA_SHEET)
    if start >= end:
        raise ValueError(
            f"『{RECORD_LIST_SHEET}』と『{DATA_SHEET}』の間にレコード詳細シートが見つかりません"
            f"（シート順: {names}）。duplicate_sheet.pyの`--before-sheet データ`の慣習どおりに"
            f"シートが並んでいるか確認してください。"
        )
    record_sheets = names[start:end]
    if len(record_sheets) > max_records:
        raise ValueError(
            f"レコード詳細シートが{len(record_sheets)}件検出されましたが、想定上限（{max_records}件）を"
            f"超えています。『{RECORD_LIST_SHEET}』と『{DATA_SHEET}』の間に想定外のシートが紛れていないか"
            f"確認してください。意図的にこの件数を扱いたい場合は--max-recordsで上限を引き上げてください。"
        )

    ws = wb[TOC_SHEET]
    written = []
    for i, sheet_name in enumerate(record_sheets):
        row = FIRST_RECORD_ROW + i
        cell = ws.cell(row=row, column=RECORD_COL)
        before = cell.value
        cell.value = sheet_name
        written.append((cell.coordinate, before, sheet_name))

    save_with_shapes(wb, xlsx_path, output_path)
    return {"record_sheets": record_sheets, "written": written}


def main():
    ensure_utf8_stdio()
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("xlsx_path")
    parser.add_argument("output_path")
    parser.add_argument("--max-records", type=int, default=DEFAULT_MAX_RECORDS)
    args = parser.parse_args()

    result = sync_toc_record_sheets(args.xlsx_path, args.output_path, args.max_records)
    print(f"=== 目次シート同期結果: {args.output_path} ===")
    print(f"検出したレコード詳細シート: {result['record_sheets']}")
    for coord, before, after in result["written"]:
        print(f"  目次!{coord} : {before!r} -> {after!r}")


if __name__ == "__main__":
    main()
