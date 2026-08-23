# -*- coding: utf-8 -*-
"""
Mermaidのフローチャート記法から画像を生成し、既存の図形(凡例等)を保持したまま
指定シートに埋め込むスクリプト。あわせてMermaidのソースコードを指定セルからテキストとして
書き込む(後から図の内容を追跡・修正できるようにするため)。

前提: `mmdc`(@mermaid-js/mermaid-cli)がインストール済みであること。
    npm install -g @mermaid-js/mermaid-cli

【方針】
OOXMLの仕様上、1シートにつき<drawing>要素(=drawingパーツへの参照)は1つしか持てない。
そのため、画像はopenpyxlのadd_image()ではなく、既存の図形(凡例等)が入っているdrawingパーツの
XMLに直接<xdr:oneCellAnchor>(画像)を追記する形で挿入する。openpyxlのadd_image()を使うと
別のdrawingパーツが作られてしまい、restore_shapes.pyによる図形復元と衝突する。

使い方:
    python insert_flow_diagram.py <元xlsxパス> <シート名> <mermaidファイルパス> <画像アンカーセル> \\
        <ソース貼り付け開始セル> <出力xlsxパス>

例:
    python insert_flow_diagram.py システム機能設計書.xlsx "1.2. 処理フロー" flow.mmd P5 AM5 出力.xlsx
"""
import sys
import re
import os
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from restore_shapes import _sheet_name_to_target, _find_drawing_for_sheet, NS_R
from xlsx_common import save_with_shapes, ensure_utf8_stdio

import openpyxl
from openpyxl.utils.cell import coordinate_from_string
from openpyxl.utils import column_index_from_string
from PIL import Image as PILImage

EMU_PER_PIXEL = 9525
MAX_WIDTH_EMU = 6000000  # 約16.7cm程度に収める


def _find_mmdc():
    mmdc_path = shutil.which("mmdc")
    if mmdc_path is None:
        raise RuntimeError(
            "mmdc(@mermaid-js/mermaid-cli)が見つかりません。"
            "`npm install -g @mermaid-js/mermaid-cli`でインストールしてください。"
        )
    return mmdc_path


def render_mermaid_to_png(mermaid_source, out_png):
    with tempfile.NamedTemporaryFile(suffix=".mmd", delete=False, mode="w", encoding="utf-8") as f:
        f.write(mermaid_source)
        mmd_path = f.name
    try:
        mmdc_path = _find_mmdc()
        result = subprocess.run(
            [mmdc_path, "-i", mmd_path, "-o", out_png, "-b", "white"],
            capture_output=True, text=True, timeout=60, shell=(os.name == "nt"),
        )
        if result.returncode != 0:
            raise RuntimeError(f"mmdc(Mermaid CLI)の実行に失敗しました。Mermaid記法を確認してください。\n{result.stderr}")
    finally:
        os.remove(mmd_path)


def write_mermaid_source_text(xlsx_path, sheet_name, mermaid_source, source_anchor_cell, output_path):
    """Mermaidソースをセルに書き込む(セルの値のみ変更。図形はsave_with_shapesで保つ)"""
    wb = openpyxl.load_workbook(xlsx_path)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"シートが存在しません: {sheet_name}")
    ws = wb[sheet_name]
    col_str, row = coordinate_from_string(source_anchor_cell)
    col = column_index_from_string(col_str)
    ws.cell(row=row, column=col, value="【Mermaidソース】")
    for i, line in enumerate(mermaid_source.splitlines(), start=1):
        ws.cell(row=row + i, column=col, value=line)
    save_with_shapes(wb, xlsx_path, output_path)


def get_image_size_emu(png_path):
    with PILImage.open(png_path) as img:
        w, h = img.size
    width_emu = w * EMU_PER_PIXEL
    height_emu = h * EMU_PER_PIXEL
    if width_emu > MAX_WIDTH_EMU:
        scale = MAX_WIDTH_EMU / width_emu
        width_emu = int(width_emu * scale)
        height_emu = int(height_emu * scale)
    return width_emu, height_emu


def estimate_image_bottom_right(ws, anchor_row0, anchor_col0, width_emu, height_emu,
                                 default_row_height_pt=15, default_col_width_chars=8.43):
    """
    画像の挿入位置(0-based行・列)とサイズ(EMU)から、画像が占めるおおよその
    行範囲・列範囲(0-based、下端を含む)を見積もる。行高・列幅は実際の値があればそれを使い、
    無ければExcelの既定値で近似する(あくまで概算であり、正確なピクセル計算ではない)。
    """
    POINTS_TO_PIXELS = 96 / 72
    CHAR_TO_PIXELS = 7  # Excel列幅の「文字数」から概算px換算する簡易係数

    height_px = height_emu / EMU_PER_PIXEL
    cum = 0.0
    row1 = anchor_row0
    r = anchor_row0 + 1
    while cum < height_px and r < anchor_row0 + 2000:
        h_pt = ws.row_dimensions[r].height or default_row_height_pt
        cum += h_pt * POINTS_TO_PIXELS
        row1 = r
        r += 1

    width_px = width_emu / EMU_PER_PIXEL
    cum = 0.0
    col1 = anchor_col0
    c = anchor_col0 + 1
    while cum < width_px and c < anchor_col0 + 200:
        w_chars = ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width or default_col_width_chars
        cum += w_chars * CHAR_TO_PIXELS + 5
        col1 = c
        c += 1

    return row1, col1


def _get_existing_shape_bboxes_with_names(xlsx_path, sheet_name):
    """_get_existing_shape_bboxesに図形名も添えて返す(重なり検出時の報告用)"""
    with zipfile.ZipFile(xlsx_path) as z:
        sheet_map = _sheet_name_to_target(z)
        if sheet_name not in sheet_map:
            return []
        drawing_path = _find_drawing_for_sheet(z, sheet_map[sheet_name])
        if drawing_path is None:
            return []
        data = z.read(drawing_path).decode("utf-8")

    bboxes = []
    for tag in ("twoCellAnchor", "oneCellAnchor"):
        for m in re.finditer(rf'<xdr:{tag}[^>]*>(.*?)</xdr:{tag}>', data, re.S):
            body = m.group(1)
            from_m = re.search(r'<xdr:from>\s*<xdr:col>(\d+)</xdr:col>.*?<xdr:row>(\d+)</xdr:row>', body, re.S)
            if not from_m:
                continue
            col0, row0 = int(from_m.group(1)), int(from_m.group(2))
            to_m = re.search(r'<xdr:to>\s*<xdr:col>(\d+)</xdr:col>.*?<xdr:row>(\d+)</xdr:row>', body, re.S)
            col1, row1 = (int(to_m.group(1)), int(to_m.group(2))) if to_m else (col0, row0)
            name_m = re.search(r'name="([^"]*)"', body)
            bboxes.append((row0, row1, col0, col1, name_m.group(1) if name_m else "(名前不明の図形)"))
    return bboxes


def check_overlap_with_existing_shapes(xlsx_path, sheet_name, anchor_cell, width_emu, height_emu):
    """
    挿入しようとしている画像が、シートに既にある図形(凡例等)と行・列範囲で重なるかを
    概算チェックする。重なりがあれば、重なった図形の名前のリストを返す(無ければ空リスト)。
    このチェックは行高・列幅の概算に基づく目安であり、確実な判定には
    `export_sheet_preview.py`でPDF化して目視確認すること。
    """
    wb = openpyxl.load_workbook(xlsx_path)
    if sheet_name not in wb.sheetnames:
        return []
    ws = wb[sheet_name]

    col_str, row = coordinate_from_string(anchor_cell)
    col0 = column_index_from_string(col_str) - 1
    row0 = row - 1
    row1, col1 = estimate_image_bottom_right(ws, row0, col0, width_emu, height_emu)

    overlapping = []
    for (s_row0, s_row1, s_col0, s_col1, s_name) in _get_existing_shape_bboxes_with_names(xlsx_path, sheet_name):
        rows_overlap = row0 <= s_row1 and row1 >= s_row0
        cols_overlap = col0 <= s_col1 and col1 >= s_col0
        if rows_overlap and cols_overlap:
            overlapping.append(s_name)

    return overlapping


def insert_image_into_drawing(xlsx_path, sheet_name, png_path, anchor_cell, output_path):
    col_str, row = coordinate_from_string(anchor_cell)
    col0 = column_index_from_string(col_str) - 1
    row0 = row - 1
    width_emu, height_emu = get_image_size_emu(png_path)

    with zipfile.ZipFile(xlsx_path) as zin:
        names = zin.namelist()
        sheet_map = _sheet_name_to_target(zin)
        if sheet_name not in sheet_map:
            raise ValueError(f"シートが見つかりません: {sheet_name}")
        sheet_target = sheet_map[sheet_name]
        drawing_path = _find_drawing_for_sheet(zin, sheet_target)
        if drawing_path is None:
            raise ValueError(
                f"シート『{sheet_name}』に既存の図形(drawing)が見つかりません。"
                f"先にsave_with_shapes等で凡例等の図形を復元してから実行してください。"
            )

        drawing_xml = zin.read(drawing_path).decode("utf-8")
        drawing_dir = "/".join(drawing_path.split("/")[:-1])
        drawing_file = drawing_path.split("/")[-1]
        drawing_rels_path = f"{drawing_dir}/_rels/{drawing_file}.rels"
        if drawing_rels_path in names:
            drawing_rels_xml = zin.read(drawing_rels_path).decode("utf-8")
        else:
            drawing_rels_xml = (
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '</Relationships>'
            )

        content_types_xml = zin.read("[Content_Types].xml").decode("utf-8")

        existing_media = [
            int(m.group(1)) for n in names
            for m in [re.match(r"xl/media/image(\d+)\.\w+$", n)] if m
        ]
        next_media_num = max(existing_media, default=0) + 1
        media_name = f"image{next_media_num}.png"
        media_path = f"xl/media/{media_name}"

        with open(png_path, "rb") as f:
            image_bytes = f.read()

        existing_rids = re.findall(r'Id="rId(\d+)"', drawing_rels_xml)
        new_rid_num = max([int(x) for x in existing_rids], default=0) + 1
        new_rid = f"rId{new_rid_num}"
        new_rel = (
            f'<Relationship Id="{new_rid}" '
            f'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" '
            f'Target="../media/{media_name}"/>'
        )
        drawing_rels_xml = drawing_rels_xml.replace("</Relationships>", new_rel + "</Relationships>")

        if 'Extension="png"' not in content_types_xml:
            content_types_xml = content_types_xml.replace(
                "</Types>", '<Default Extension="png" ContentType="image/png"/></Types>'
            )

        existing_ids = [int(m) for m in re.findall(r'<xdr:cNvPr id="(\d+)"', drawing_xml)]
        new_shape_id = max(existing_ids, default=0) + 1

        pic_xml = (
            f'<xdr:oneCellAnchor>'
            f'<xdr:from><xdr:col>{col0}</xdr:col><xdr:colOff>0</xdr:colOff>'
            f'<xdr:row>{row0}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>'
            f'<xdr:ext cx="{width_emu}" cy="{height_emu}"/>'
            f'<xdr:pic>'
            f'<xdr:nvPicPr>'
            f'<xdr:cNvPr id="{new_shape_id}" name="処理フロー図"/>'
            f'<xdr:cNvPicPr><a:picLocks noChangeAspect="1"/></xdr:cNvPicPr>'
            f'</xdr:nvPicPr>'
            f'<xdr:blipFill>'
            f'<a:blip xmlns:r="{NS_R}" r:embed="{new_rid}"/>'
            f'<a:stretch><a:fillRect/></a:stretch>'
            f'</xdr:blipFill>'
            f'<xdr:spPr>'
            f'<a:xfrm><a:off x="0" y="0"/><a:ext cx="{width_emu}" cy="{height_emu}"/></a:xfrm>'
            f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
            f'</xdr:spPr>'
            f'</xdr:pic>'
            f'<xdr:clientData/>'
            f'</xdr:oneCellAnchor>'
        )

        if "</xdr:wsDr>" not in drawing_xml:
            raise ValueError("drawing XMLのルート要素(</xdr:wsDr>)が見つかりません。想定外の形式です。")
        drawing_xml = drawing_xml.replace("</xdr:wsDr>", pic_xml + "</xdr:wsDr>")

        updated_files = {
            drawing_path: drawing_xml.encode("utf-8"),
            drawing_rels_path: drawing_rels_xml.encode("utf-8"),
            "[Content_Types].xml": content_types_xml.encode("utf-8"),
            media_path: image_bytes,
        }

        with zipfile.ZipFile(output_path, "w", zipfile.ZIP_DEFLATED) as zout:
            written = set()
            for item in zin.infolist():
                data = updated_files.get(item.filename, zin.read(item.filename))
                zout.writestr(item, data)
                written.add(item.filename)
            for path, data in updated_files.items():
                if path not in written:
                    zout.writestr(path, data)
                    written.add(path)

    return {"media": media_path, "shape_id": new_shape_id, "size_emu": (width_emu, height_emu)}


def insert_flow_diagram(xlsx_path, sheet_name, mermaid_source, image_anchor_cell, source_anchor_cell, output_path):
    """
    mmdcが使える場合: 画像を生成してシートに挿入し、Mermaidソースもテキストとして貼り付ける。
    mmdcが使えない場合: 画像の挿入はスキップし、Mermaidソースのテキスト貼り付けのみ行う
    (この場合、戻り値の"image_inserted"がFalseになる。呼び出し側はユーザーに、
    ソースから画像を作成して手動で貼り付けるよう案内すること)。
    """
    if shutil.which("mmdc") is None:
        write_mermaid_source_text(xlsx_path, sheet_name, mermaid_source, source_anchor_cell, output_path)
        return {
            "image_inserted": False,
            "media": None,
            "shape_id": None,
            "size_emu": None,
            "source_anchor_cell": source_anchor_cell,
            "overlap_warning": None,
        }

    tmp_dir = tempfile.mkdtemp()
    try:
        png_path = os.path.join(tmp_dir, "flow.png")
        render_mermaid_to_png(mermaid_source, png_path)

        tmp_with_text = os.path.join(tmp_dir, "with_text.xlsx")
        write_mermaid_source_text(xlsx_path, sheet_name, mermaid_source, source_anchor_cell, tmp_with_text)

        width_emu, height_emu = get_image_size_emu(png_path)
        overlapping = check_overlap_with_existing_shapes(
            tmp_with_text, sheet_name, image_anchor_cell, width_emu, height_emu
        )

        result = insert_image_into_drawing(tmp_with_text, sheet_name, png_path, image_anchor_cell, output_path)
        result["image_inserted"] = True
        result["source_anchor_cell"] = source_anchor_cell
        result["overlap_warning"] = overlapping or None
        return result
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)


def main():
    ensure_utf8_stdio()
    if len(sys.argv) != 7:
        print(__doc__)
        sys.exit(2)

    xlsx_path, sheet_name, mermaid_file, image_anchor_cell, source_anchor_cell, output_path = sys.argv[1:7]

    with open(mermaid_file, encoding="utf-8") as f:
        mermaid_source = f.read()

    result = insert_flow_diagram(
        xlsx_path, sheet_name, mermaid_source, image_anchor_cell, source_anchor_cell, output_path
    )
    print(f"=== 処理フロー図の挿入結果: {output_path} ===")
    if result["image_inserted"]:
        print(f"画像パーツ: {result['media']}")
        width_px, height_px = result['size_emu'][0] // 9525, result['size_emu'][1] // 9525
        print(f"アンカーセル: {image_anchor_cell}（サイズ: 幅約{width_px}px × 高さ約{height_px}px）")
        print(f"Mermaidソースの貼り付け開始セル: {source_anchor_cell}")
        if result.get("overlap_warning"):
            print("")
            print(f"【注意】ノード数が多いため画像の縦幅が大きくなっています。"
                  f"既存の図形と重なっている可能性があります: {result['overlap_warning']}")
            print(f"この判定は行高・列幅の概算に基づく目安です。"
                  f"`export_sheet_preview.py`で『{sheet_name}』シートをPDF化し、"
                  f"実際に重なっていないか必ず目視確認してください。"
                  f"重なっている場合は、アンカーセルを図形の無い位置にずらすか、"
                  f"Mermaid記法の向き(例: `flowchart LR`で横方向に展開する等)を変えて縦幅を抑えること。")
    else:
        print("【注意】mmdc(@mermaid-js/mermaid-cli)が見つからなかったため、画像は挿入していません。")
        print(f"Mermaidソースのみ『{sheet_name}』シートの{source_anchor_cell}セル以降に貼り付けました。")
        print("")
        print("【ユーザーへの案内が必要です】以下をそのまま伝えてください：")
        print(
            f"「処理フロー図の画像化ツール(mermaid-cli)がローカル環境に無かったため、"
            f"『{sheet_name}』シートの{source_anchor_cell}セル以降にMermaid記法のソースコードのみ貼り付けました。"
            f"お手数ですが、以下のいずれかの方法で画像を作成し、{image_anchor_cell}セル付近に貼り付けてください。\n"
            f"  1. https://mermaid.live などのMermaidエディタにソースを貼り付けて画像として保存し、Excelに挿入する\n"
            f"  2. `npm install -g @mermaid-js/mermaid-cli` でmermaid-cliをインストールした上で、"
            f"このスキルの`insert_flow_diagram.py`を再実行する\n"
            f"  3. 図形オブジェクト（四角・ひし形・矢印等）で直接描く」"
        )


if __name__ == "__main__":
    main()
