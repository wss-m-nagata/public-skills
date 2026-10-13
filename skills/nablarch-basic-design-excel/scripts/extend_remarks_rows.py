# -*- coding: utf-8 -*-
"""
外部インタフェース設計書の『1. 外部インタフェース仕様』シートの、特記事項の行を増やす。

【背景】特記事項は、テンプレートでは E28〜E31 の4行分しか罫線が引かれていない。
これは、テンプレートがそう配布されているだけで、4行までと決まっているわけではない。
IFを使う側が知っておくべき振る舞いが4行に収まらないときは、行を増やして書く。
行の増やし方は、extend_table_rows.py を使うが、このシートでは、それだけだと次の3点が足りない。
このスクリプトは、それを含めて行う。

  1. 増やした行は、最終行（31行目）の書式の複製のため、各行に下線が付く。
     途中の行の下線を消し、特記事項の枠の最後の行だけに下線を残す（元の4行と同じ見た目にする）。
  2. 印刷範囲（A1:AI32 など）の最終行が、増やした分だけ足りなくなる。増やした行数だけ広げる。
  3. 印刷は「横を1ページに収め、縦は自動」のため、行が増えると2ページになる。縦も1ページに収める。

特記事項の1行は、全角70文字ほどに収める（折り返さない1行のセルで、超えると切れる）。
収まらない内容は、行を分けて書く。

使い方:
    python extend_remarks_rows.py <入力xlsx> <出力xlsx> --count <増やす行数>

    --count : 増やす行数。例：11行にしたいときは 7（元の4行＋7行）。
    --sheet : 対象のシート名（省略時は「1. 外部インタフェース仕様」）。

実行したあと、E32〜 に特記事項を書く（apply_mapping.py）。構造の検証（verify_layout.py）は、
このスクリプトの出力を基準にする（行が増えているため、入力とは比べない）。
"""

import argparse
import sys
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.styles import Border, Side

sys.path.insert(0, str(Path(__file__).parent))
from extend_table_rows import extend_table_rows  # noqa: E402
from xlsx_common import ensure_utf8_stdio, save_with_shapes  # noqa: E402

DEFAULT_SHEET = "1. 外部インタフェース仕様"
# 特記事項の最終行（テンプレートの E28〜E31 の最後）と、複製する列の範囲
LAST_REMARKS_ROW = 31
FIRST_COLUMN = "A"
LAST_COLUMN = "AI"


def extend_remarks(input_path, output_path, count, sheet_name=DEFAULT_SHEET):
    """特記事項の行を増やし、印刷範囲・ページ・途中の行の下線を整える。"""
    tmp = str(output_path) + ".extend.xlsx"
    extend_table_rows(
        input_path, sheet_name, LAST_REMARKS_ROW, count, FIRST_COLUMN, LAST_COLUMN, tmp
    )

    wb = openpyxl.load_workbook(tmp)
    ws = wb[sheet_name]

    # 2. 印刷範囲を、増やした行数だけ広げる
    first, last = ws.print_area.split("!")[-1].replace("$", "").split(":")
    column = "".join(ch for ch in last if ch.isalpha())
    row = int("".join(ch for ch in last if ch.isdigit()))
    ws.print_area = f"{first}:{column}{row + count}"

    # 3. 縦も1ページに収める
    ws.page_setup.fitToHeight = 1

    # 1. 途中の行の下線を消す（最後の行だけに残す）
    for r in range(LAST_REMARKS_ROW, LAST_REMARKS_ROW + count):
        for c in range(1, ws.max_column + 1):
            border = ws.cell(row=r, column=c).border
            ws.cell(row=r, column=c).border = Border(
                left=copy(border.left),
                right=copy(border.right),
                top=copy(border.top),
                bottom=Side(),
            )

    save_with_shapes(wb, input_path, output_path)
    Path(tmp).unlink()


def main():
    ensure_utf8_stdio()
    parser = argparse.ArgumentParser(
        description="外部インタフェース設計書の特記事項の行を増やす"
    )
    parser.add_argument("input")
    parser.add_argument("output")
    parser.add_argument("--count", type=int, required=True)
    parser.add_argument("--sheet", default=DEFAULT_SHEET)
    args = parser.parse_args()
    if args.count < 1:
        parser.error("--count は1以上にしてください")

    extend_remarks(args.input, args.output, args.count, args.sheet)
    print(f"OK: 特記事項の行を {args.count} 行増やしました（E28〜E{LAST_REMARKS_ROW + args.count}）。")
    print(f"=== 出力: {args.output} ===")


if __name__ == "__main__":
    main()
