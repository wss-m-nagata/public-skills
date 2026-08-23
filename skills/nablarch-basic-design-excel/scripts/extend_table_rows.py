# -*- coding: utf-8 -*-
"""
一覧系テンプレート(システム機能一覧・外部インタフェース一覧・WebサービスAPI一覧・
テーブル一覧等)で、テンプレートに用意された記入可能行数(罫線が引かれた範囲)を
超える件数を記入する必要がある場合に、最終データ行の書式を下方向に複製して
記入可能な行を増やすスクリプト。

行の挿入(insert_rows)は行わない。既存のシート上の未使用領域(罫線が無い空白行)に、
最終データ行と同じ書式(罫線・背景色・フォント・結合セル)を複製するだけなので、
これより下や他のシートの内容には一切影響しない。

入力規則(プルダウン)が最終データ行に設定されている場合は、その適用範囲を
新しく増やした行にも拡張する。

**安全確認**：追加先の範囲内に、既に値が入っているセル(次のセクションの見出し等)が
無いかを事前に確認し、あれば拒否する(後続の見出し行を上書きしないための対策)。
複数行から成るセクションごと後ろにずらしたい場合は、このスクリプトではなく
`insert_row_block.py`を使うこと。

使い方:
    python extend_table_rows.py <元xlsxパス> <シート名> <最終データ行番号> <追加する行数> \\
        <対象列(開始)> <対象列(終了)> <出力xlsxパス> [--allow-overwrite]

出力パスは必須(元ファイルを直接上書きする事故を防ぐため)。

例(外部インタフェース一覧シート「1」で、最終データ行が18、対象列がB〜AX、
   10行追加する場合):
    python extend_table_rows.py 外部インタフェース一覧_xxx.xlsx 1 18 10 B AX 出力.xlsx
"""
import sys
import argparse
from pathlib import Path
from copy import copy

import openpyxl
from openpyxl.comments import Comment
from openpyxl.utils import column_index_from_string, get_column_letter

sys.path.insert(0, str(Path(__file__).parent))
from xlsx_common import (
    get_single_row_merges, extend_validation_ranges, get_bordered_bbox,
    find_next_populated_row, save_with_shapes, ensure_utf8_stdio,
)


def extend_table_rows(xlsx_path, sheet_name, source_row, num_new_rows, min_col_letter, max_col_letter,
                       output_path, allow_overwrite=False):
    wb = openpyxl.load_workbook(xlsx_path)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"シートが存在しません: {sheet_name}")
    ws = wb[sheet_name]

    min_col = column_index_from_string(min_col_letter)
    max_col = column_index_from_string(max_col_letter)

    bbox_row, _ = get_bordered_bbox(ws)
    if source_row > bbox_row:
        raise ValueError(
            f"指定された最終データ行({source_row})が、現在の罫線設定範囲の最大行({bbox_row})を"
            f"超えています。既に書式が無い行を複製元にはできません。"
        )

    if not allow_overwrite:
        boundary_row = find_next_populated_row(ws, source_row, min_col, max_col)
        if boundary_row is not None:
            safe_new_rows = boundary_row - source_row - 1
            if num_new_rows > safe_new_rows:
                raise ValueError(
                    f"{boundary_row}行目に既に値が入っているセルがあります"
                    f"(次のセクションの見出し等の可能性が高い)。"
                    f"{source_row}行目から安全に追加できるのは最大{max(safe_new_rows, 0)}行までです"
                    f"(指定された{num_new_rows}行では{boundary_row}行目以降を上書きします)。\n"
                    f"複数行から成るセクションごと後ろにずらして拡張したい場合は、"
                    f"このスクリプトではなく`insert_row_block.py`を使ってください。"
                    f"(意図的にこのまま拡張したい場合のみ --allow-overwrite を指定)"
                )

    row_merges = get_single_row_merges(ws, source_row, min_col, max_col)
    row_height = ws.row_dimensions[source_row].height

    new_row_numbers = [source_row + i for i in range(1, num_new_rows + 1)]

    for new_row in new_row_numbers:
        for col in range(min_col, max_col + 1):
            src_cell = ws.cell(row=source_row, column=col)
            dst_cell = ws.cell(row=new_row, column=col)
            dst_cell.font = copy(src_cell.font)
            dst_cell.fill = copy(src_cell.fill)
            dst_cell.border = copy(src_cell.border)
            dst_cell.alignment = copy(src_cell.alignment)
            dst_cell.number_format = src_cell.number_format
            dst_cell.protection = copy(src_cell.protection)
            if src_cell.comment is not None:
                dst_cell.comment = Comment(src_cell.comment.text, src_cell.comment.author or "")

        if row_height is not None:
            ws.row_dimensions[new_row].height = row_height

        for (mcol_start, mcol_end) in row_merges:
            ws.merge_cells(start_row=new_row, start_column=mcol_start,
                            end_row=new_row, end_column=mcol_end)

    extended = extend_validation_ranges(ws, source_row, new_row_numbers, min_col, max_col)

    save_with_shapes(wb, xlsx_path, output_path)

    return {
        "sheet": sheet_name,
        "new_rows": new_row_numbers,
        "merges_replicated_per_row": len(row_merges),
        "validations_extended": extended,
    }


def main():
    ensure_utf8_stdio()
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("xlsx_path")
    parser.add_argument("sheet_name")
    parser.add_argument("source_row", type=int)
    parser.add_argument("num_new_rows", type=int)
    parser.add_argument("min_col")
    parser.add_argument("max_col")
    parser.add_argument("output_path")
    parser.add_argument("--allow-overwrite", action="store_true")
    args = parser.parse_args()

    result = extend_table_rows(args.xlsx_path, args.sheet_name, args.source_row,
                                args.num_new_rows, args.min_col, args.max_col, args.output_path,
                                args.allow_overwrite)

    print(f"=== 行拡張結果: {args.output_path} ===")
    print(f"追加した行: {result['new_rows']}")
    print(f"複製した結合セルパターン数(1行あたり): {result['merges_replicated_per_row']}")
    if result["validations_extended"]:
        print(f"拡張した入力規則: {result['validations_extended']}")
    print("\n※新しい行にはまだ値が入っていません。apply_mapping.pyで値を反映してください。")


if __name__ == "__main__":
    main()
