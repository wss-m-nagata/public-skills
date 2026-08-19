#!/usr/bin/env python3
"""
Excel出力スクリプト。
既存テンプレート(assets/AI駆動開発_実践力カタログ.xlsx)を読み込み、評価結果を
反映した新規Excelファイルとして保存する。テンプレート自体は変更しない。

「評価根拠」「到達段階」「次の段階」は check-results.json / criteria.json から
機械的に転記する。「推奨対応」だけは check-results.json の構造(check単位の notes)
では criteria単位の文章を保持できないため、--narratives-file で
{criteria_id: 推奨対応の文章} 形式のJSONを別途受け取り、その値をそのまま転記する。
この文章自体の作成(意味判断)はAIが export 実行前に行う。

Revision History
# 2026/08/19: 新規作成。
"""

import argparse
import json
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
import _common
from calculate_score import calculate_scores

TEMPLATE_XLSX = _common.SKILL_DIR / "assets" / "AI駆動開発_実践力カタログ.xlsx"


def _find_header_row(ws):
    """ヘッダ行('No.'が入っている行)を探す。"""
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        if row[0].value == "No.":
            return row[0].row
    raise ValueError("ヘッダ行('No.')がテンプレートExcel内に見つかりません。")


def _build_column_index(ws, header_row):
    """ヘッダ名(列名文字列)から列番号(1始まり)への対応表を作る。"""
    index = {}
    for cell in ws[header_row]:
        if cell.value:
            index[cell.value] = cell.column
    return index


def _write_label_value(ws, header_row, label, value):
    """ヘッダ行より上にあるラベルセルを探し、その右隣のセルへ値を書き込む。"""
    if value is None:
        return
    for row in ws.iter_rows(min_row=1, max_row=header_row):
        for cell in row:
            if cell.value == label:
                ws.cell(row=cell.row, column=cell.column + 1, value=value)
                return


def _level_state(levels, score):
    for level in levels:
        if level["score"] == score:
            return level["state"]
    return None


def _evidence_note(item_score, item):
    """採用された到達条件(0点の場合は25点条件)の evidence + notes を連結する。"""
    target_threshold = item_score if item_score else 25
    for check in item["checks"]:
        if check["score_threshold"] == target_threshold:
            parts = []
            if check.get("evidence"):
                parts.append("／".join(check["evidence"]))
            if check.get("notes"):
                parts.append(check["notes"])
            return "\n".join(parts) if parts else None
    return None


def export_assessment(output_path, narratives_file=None, force=False):
    """
    評価結果をExcelへ出力する。

    引数:
        output_path: 出力先Excelファイルパス
        narratives_file: 「推奨対応」列に転記する文章を持つJSONファイルのパス
            ({criteria_id: 推奨対応の文章})。省略時は「推奨対応」列を空欄のままにする。
        force: True の場合、出力先が既に存在していても上書きする。

    戻り値:
        Path: 実際に書き込んだ出力先パス
    """
    output_path = Path(output_path)
    if output_path.resolve() == TEMPLATE_XLSX.resolve():
        raise ValueError(
            "テンプレートExcelを直接上書きすることはできません。--output に別のパスを指定してください。"
        )
    if output_path.exists() and not force:
        raise FileExistsError(
            f"出力先が既に存在します: {output_path} (--force を指定すると上書きします)"
        )

    state, _source = _common.load_state()
    criteria = _common.load_criteria()
    criteria_by_name = {c["name"]: c for c in criteria["criteria"]}

    scored = calculate_scores(save=True)
    score_by_id = {s["criteria_id"]: s for s in scored["items"]}
    item_by_id = {item["criteria_id"]: item for item in state["items"]}

    narratives = {}
    if narratives_file:
        with open(narratives_file, "r", encoding="utf-8") as f:
            narratives = json.load(f)

    wb = openpyxl.load_workbook(TEMPLATE_XLSX)
    ws = wb.active

    header_row = _find_header_row(ws)
    col = _build_column_index(ws, header_row)

    project = state.get("project", {})
    _write_label_value(ws, header_row, "プロジェクト名", project.get("name"))
    _write_label_value(ws, header_row, "評価日", project.get("evaluated_at"))

    for row_idx in range(header_row + 1, ws.max_row + 1):
        name_cell = ws.cell(row=row_idx, column=col["評価項目"])
        if not name_cell.value:
            continue

        criteria_def = criteria_by_name.get(name_cell.value)
        if criteria_def is None:
            raise ValueError(
                f"評価項目 '{name_cell.value}' に対応する criteria.json の定義が見つかりません。"
            )

        criteria_id = criteria_def["id"]
        score_info = score_by_id.get(criteria_id)
        item = item_by_id.get(criteria_id)
        if score_info is None or item is None:
            raise ValueError(
                f"criteria_id '{criteria_id}' の確認結果が check-results.json に存在しません。"
            )

        applicable = score_info["applicable"]
        score = score_info["score"]
        has_progress = any(c["status"] != "unconfirmed" for c in item["checks"])

        if not applicable:
            status_label = "対象外"
        elif has_progress:
            status_label = "評価済み"
        else:
            status_label = "未評価"
        ws.cell(row=row_idx, column=col["適用区分"], value=status_label)

        if applicable:
            ws.cell(row=row_idx, column=col["点数"], value=score)
            ws.cell(
                row=row_idx,
                column=col["到達段階"],
                value=_level_state(criteria_def["levels"], score),
            )
            if score < 100:
                ws.cell(
                    row=row_idx,
                    column=col["次の段階"],
                    value=_level_state(criteria_def["levels"], score + 25),
                )
            ws.cell(
                row=row_idx,
                column=col["評価根拠"],
                value=_evidence_note(score, item),
            )
            ws.cell(
                row=row_idx,
                column=col["推奨対応"],
                value=narratives.get(criteria_id),
            )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def main():
    parser = argparse.ArgumentParser(description="評価結果をExcelへ出力する。")
    parser.add_argument("--output", required=True, help="出力先Excelファイルパス")
    parser.add_argument(
        "--narratives-file",
        default=None,
        help="推奨対応の文章を持つJSONファイル({criteria_id: text})",
    )
    parser.add_argument(
        "--force", action="store_true", help="出力先が既存でも上書きする"
    )
    args = parser.parse_args()

    try:
        path = export_assessment(args.output, args.narratives_file, args.force)
    except (ValueError, FileExistsError) as e:
        print(json.dumps({"error": str(e)}, ensure_ascii=False), file=sys.stderr)
        sys.exit(1)

    print(json.dumps({"output": str(path)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
